"""Offline checks for the downloadable reconciliation snapshot."""

from __future__ import annotations

import io
import json
import unittest
from datetime import datetime

from openpyxl import load_workbook

from credinomina_reconciliation.control_export import (
    build_control_workbook, _control_state, _collection_pending, monthly_rows,
)


class TestControlExport(unittest.TestCase):
    def test_historical_missing_evidence_and_many_to_many_links(self):
        historical = {
            "name": "H-2025-04", "employer": "EMP-1", "employer_name": "Empresa Uno",
            "month": "2025-04", "reconciliation_mode": "Historica",
            "historical_scope": "Fecha exacta", "historical_application_date": "2025-04-15",
            "status": "Conciliado", "control_state": "historico_conciliado",
            "expected_usd": 0, "deducted_usd": 0, "applied_usd": 46.52,
            "complementary_usd": 0, "rounding_adjustment_usd": 0.01,
            "remitted_usd": 46.53, "historical_pending_usd": 0,
            "worker_gap_usd": 0, "employer_gap_usd": 0, "unclassified_deposit_usd": 0,
            "rows": [], "historical_rows": [{
                "name": "APP-1", "client_number": "001", "client_name": "=HYPERLINK(\"x\")",
                "loan_number": "L-1", "event_date": "2025-04-15", "reference": "AP-1",
                "amount": 46.52, "historical_remitted_usd": 46.53,
                "historical_balance_usd": 0, "deposit_match_status": "Depósito conciliado",
                "historical_detail": json.dumps([
                    {"referencia": "DEP-1", "fecha": "2025-05-03", "importe_usd": 20},
                    {"referencia": "DEP-2", "fecha": "2025-05-18", "importe_usd": 26.52},
                    {"referencia": "DEP-2", "fecha": "2025-05-18", "importe_usd": 0.01,
                     "diferencia_usd": 0.01, "origen": "Tolerancia automática"},
                ]),
            }],
            "exceptions": [], "surpluses": [],
        }
        operational = {
            "name": "O-2026-09", "employer": "EMP-1", "employer_name": "Empresa Uno",
            "month": "2026-09", "reconciliation_mode": "Operativa",
            "remittance_due_date": "2026-10-15", "control_state": "diferencia",
            "control_cut_on": "2026-09-30 18:00:00",
            "control_cut_note": "Solicitar pago parcial y detalle actualizado",
            "control_cut_summary": "Cobranza US$ 100.00; deducido US$ 90.00",
            "remark": "Deducción parcial por ingreso insuficiente",
            "expected_usd": 100, "deducted_usd": 90, "applied_usd": 90,
            "complementary_usd": 0, "rounding_adjustment_usd": 0,
            "remitted_usd": 90, "historical_pending_usd": 0,
            "worker_gap_usd": 10, "employer_gap_usd": 0, "unclassified_deposit_usd": 0,
            "rows": [{
                "row_key": "ROW-1", "client_number": "002", "client_name": "Ana López",
                "loan_number": "L-2", "expected_usd": 100, "deducted_usd": 90,
                "applied_usd": 90, "complementary_usd": 0,
                "rounding_adjustment_usd": 0, "remitted_usd": 90,
                "employee_receivable_usd": 10, "application_reference": "AP-2",
                "remittance_detail": json.dumps([
                    {"referencia": "DEP-2", "fecha": "2026-10-16", "importe_usd": 90},
                ]),
            }], "historical_rows": [], "exceptions": [], "surpluses": [],
        }
        data = {
            "year": 2025, "periods": [historical, operational],
            "totals": {
                "expected_usd": 100, "deducted_usd": 90,
                "applied_usd": 136.52, "remitted_usd": 136.53,
                "complementary_usd": 0, "rounding_adjustment_usd": 0.01,
                "historical_pending_usd": 0, "worker_gap_usd": 10,
                "employer_gap_usd": 0, "unclassified_deposit_usd": 0,
            },
            "open_deposits": [], "unassigned_historical_applications": [],
            "cash_deposits": [{
                "name": "DEP-5-2025-0001", "employer": "EMP-1", "date": "2025-05-18", "month": "2025-05",
                "reference": "000123", "bank_account": "Banco Demo USD", "currency": "USD", "original_amount": 100,
                "total_usd": 100, "credits_usd": 46.53, "other_usd": 10, "adjustments_usd": 0,
                "credit_balance_usd": 30, "undetailed_credit_usd": 5, "unclassified_usd": 13.47, "review_usd": 0,
                "result": "Parcial con saldo a favor", "needs_review": True, "payroll_months": ["2025-04"],
                "destinations": [
                    {"type": "Créditos", "employer": "EMP-1", "label": "H-2025-04", "month": "2025-04", "amount_usd": 46.53,
                     "people": [{"client_name": "Cliente de prueba", "client_number": "001", "loan_number": "L-1", "amount_usd": 46.53}]},
                    {"type": "Partida complementaria", "employer": "EMP-1", "label": "Administración", "amount_usd": 10},
                    {"type": "Saldo a favor del cliente", "employer": "EMP-1", "label": "Reembolso pendiente", "amount_usd": 25,
                     "people": [{"client_name": "Cliente de prueba", "client_number": "001", "amount_usd": 25}]},
                ],
            }],
        }
        exceptions = [{
            "name": "EX-1", "period": "O-2026-09", "status": "Resuelta",
            "exception_type": "Deduccion parcial", "cause_category": "Ingreso insuficiente",
            "assigned_to": "operador@example.com", "next_action": "Confirmar saldo",
            "commitment_date": "2026-10-20", "description": "Faltan 10",
            "resolution": "Cliente notificado", "amount_usd": 10,
        }]
        actions = [{
            "parent": "EX-1", "action_at": "2026-10-18 09:00:00",
            "action_by": "operador@example.com", "action_type": "Contacto con empleado",
            "details": "Se llamó al cliente", "evidence_file": None,
            "external_reference": "TICKET-1",
        }]
        payload = build_control_workbook(
            data, exceptions=exceptions, actions=actions,
            employer_label="Empresa Uno", generated_at=datetime(2026, 9, 27, 12),
        )
        book = load_workbook(io.BytesIO(payload), data_only=False)
        self.assertEqual(book.sheetnames, [
            "Resumen mensual", "Períodos", "Detalle cliente", "Cruces", "Depósitos", "Distribución depósitos",
            "Aplicaciones sin período", "Partidas y excepciones", "Gestiones", "Guía",
        ])
        summary = book["Períodos"]
        self.assertEqual(summary["J5"].value, "N/D")
        self.assertEqual(summary["K5"].value, "N/D")
        self.assertEqual(summary["D5"].value, "Histórica")
        self.assertEqual(summary["I5"].value, "Conciliado")
        self.assertEqual(summary["I6"].value, "Con diferencias")
        self.assertEqual(summary["P5"].value, 0)
        self.assertEqual(summary["P6"].value, 0)
        self.assertEqual(summary["N5"].value, 0.01)
        self.assertEqual(summary["H6"].value.year, 2026)
        self.assertEqual(summary["U6"].value.year, 2026)
        self.assertEqual(summary["V6"].value, "Solicitar pago parcial y detalle actualizado")
        self.assertEqual(summary["W6"].value, "Deducción parcial por ingreso insuficiente")
        self.assertEqual(summary["X6"].value, "Abierto")
        detail = book["Detalle cliente"]
        self.assertEqual(detail["N5"].value, "N/D")
        self.assertEqual(detail["O5"].value, "N/D")
        self.assertEqual(detail["F5"].data_type, "s")
        self.assertTrue(detail["F5"].value.startswith("'="))
        self.assertEqual(detail["P5"].data_type, "n")
        links = book["Cruces"]
        self.assertEqual(links.max_row, 8)
        self.assertEqual([links[f"M{row}"].value for row in range(5, 9)],
                         [20, 26.52, 0.01, 90])
        self.assertEqual(book["Partidas y excepciones"]["I6"].value, "Resuelta")
        self.assertEqual(book["Gestiones"]["A5"].value, "EX-1")
        self.assertEqual(sum(row[10].value for row in book["Distribución depósitos"].iter_rows(min_row=5)), 100)
        self.assertEqual(book["Depósitos"]["D5"].value, "000123")

    def test_same_status_labels_in_both_modes(self):
        for historical, operational, expected in (
            ("historico_excepcion", "diferencia", "Con diferencias"),
            ("historico_excedente", "excedente", "Con excedente"),
            ("historico_conciliado", "conciliado", "Conciliado"),
            ("historico_parcial", "parcial", "Parcial"),
            ("historico_pendiente", "en_transito", "Pendiente"),
        ):
            with self.subTest(expected=expected):
                self.assertEqual(_control_state(historical), expected)
                self.assertEqual(_control_state(operational), expected)

    def test_pending_includes_both_modes_without_netting_unrelated_balances(self):
        periods = [{
            "name": "H", "reconciliation_mode": "Historica",
            "control_state": "historico_parcial", "historical_pending_usd": 0,
            "historical_rows": [{"historical_balance_usd": 30}, {"historical_balance_usd": 0}],
        }, {
            "name": "O", "reconciliation_mode": "Operativa", "control_state": "parcial",
            "collection_cycle": "Mensual",
            "rows": [{
                "applied_usd": 90, "remitted_usd": 80,
                "remittance_detail": [{"importe_usd": 70},
                                      {"importe_usd": 10, "destino": "Partida complementaria"}],
            }, {"applied_usd": 10, "remittance_detail": [{"importe_usd": 50}]}],
        }]
        book = load_workbook(io.BytesIO(build_control_workbook(
            {"year": 2026, "periods": periods, "totals": {}},
            exceptions=[], actions=[], employer_label="Todas", generated_at=datetime(2026, 9, 30),
        )))
        summary = book["Períodos"]
        self.assertEqual(sum(row[7] for row in monthly_rows({"periods": periods})), 50)
        self.assertEqual([summary[f"P{r}"].value for r in (5, 6)], [30, 20])
        self.assertEqual(summary["E6"].value, "Mensual")
        detail = book["Detalle cliente"]
        self.assertEqual([detail[f"U{r}"].value for r in range(5, 9)], [30, 0, 20, 0])
        for sheet, header_row in ((summary, 4), (detail, 4), (book["Cruces"], 4)):
            for cell in sheet[header_row]:
                self.assertNotRegex(cell.value, r"(?i)históric|operativ")

    def test_rounding_adjustment_is_signed_and_fx_is_not_cash(self):
        for applied, paid, adjustment in ((46.52, 46.53, 0.01), (46.53, 46.52, -0.01)):
            self.assertEqual(_collection_pending({
                "applied_usd": applied, "rounding_adjustment_usd": adjustment,
                "remittance_detail": [{"importe_usd": paid}],
            }), 0)
        self.assertEqual(_collection_pending({"applied_usd": 10, "fx_variance_usd": 10}), 10)

    def test_empty_and_single_mode_exports_keep_common_headers_and_missing_data(self):
        for mode in (None, "Historica", "Operativa"):
            with self.subTest(mode=mode):
                book = load_workbook(io.BytesIO(build_control_workbook(
                    {"year": 2026, "periods": [{"name": "P", "reconciliation_mode": mode}] if mode else [], "totals": {}},
                    exceptions=[], actions=[], employer_label="Todas", generated_at=datetime(2026, 9, 30),
                )))
                summary = book["Períodos"]
                if mode:
                    self.assertEqual(summary["J5"].value, 0 if mode == "Operativa" else "N/D")
                    self.assertEqual(summary["P5"].value, 0)
                self.assertEqual(summary["P4"].value, "Aplicado pendiente de depósito USD")
                self.assertEqual(summary.freeze_panes, "C5")

    def test_monthly_groups_cuts_keeps_cash_separate_and_does_not_net_customers(self):
        periods = [{"name": "P1", "employer": "A", "employer_name": "Empresa A", "month": "2025-04",
                    "status": "Cerrado", "reconciliation_mode": "Historica", "applied_usd": 100,
                    "historical_rows": [{"historical_remitted_usd": 120, "historical_balance_usd": 0}],
                    "exceptions": [{"name": "E1"}]},
                   {"name": "P2", "employer": "A", "employer_name": "Empresa A", "month": "2025-04",
                    "reconciliation_mode": "Operativa", "applied_usd": 50,
                    "rows": [{"applied_usd": 50, "remittance_detail": [{"importe_usd": 10},
                              {"importe_usd": 20, "destino": "Partida complementaria"}]}],
                    "exceptions": [{"name": "E1"}, {"name": "E2"}]}]
        deposit = {"name": "D1", "employer": "B", "month": "2025-05", "total_usd": 500,
                   "unclassified_usd": 20, "needs_review": True}
        rows = monthly_rows({"periods": periods, "cash_deposits": [deposit, deposit],
                             "unassigned_operational_applications": [{"name": "A1", "employer": "A",
                                "event_date": "2025-04-09", "amount": 60, "application_adjustment_usd": 10}]})
        april, may = rows
        self.assertEqual(april[0], "Empresa A")
        self.assertEqual(april[2:10], (2, 1, 150, 130, 0, 40, 2, 50))
        self.assertEqual(april[10:14], (0, 0, 0, 0))
        self.assertEqual(may[0], "B")
        self.assertEqual(may[2:10], (0, 0, 0, 0, 0, 0, 0, 0))
        self.assertEqual(may[10:14], (1, 500, 20, 1))

    def test_no_period_exception_and_both_application_modes_are_visible(self):
        data = {"year": None, "periods": [], "totals": {},
                "unassigned_historical_applications": [{"name": "H1", "parent": "CONTA-H", "employer": "A", "amount": 100}],
                "unassigned_operational_applications": [{"name": "O1", "parent": "CONTA-O", "employer": "B", "amount": 50}]}
        book = load_workbook(io.BytesIO(build_control_workbook(data,
            exceptions=[{"name": "E1", "employer": "B", "amount_usd": 1, "status": "Abierta",
                         "commitment_date": "2026-09-01"}],
            actions=[{"parent": "E1", "action_type": "Registrar ajuste en el core"}],
            employer_label="Todas", generated_at=datetime(2026, 10, 2), date_format="mm-dd-yyyy")))
        self.assertEqual(book["Aplicaciones sin período"]["A5"].value, "CONTA-H")
        self.assertEqual(book["Aplicaciones sin período"]["A6"].value, "CONTA-O")
        self.assertEqual(book["Partidas y excepciones"]["D5"].value, "B")
        self.assertEqual(book["Partidas y excepciones"]["S5"].value, 31)
        self.assertEqual(book["Partidas y excepciones"]["M5"].number_format, "mm-dd-yyyy")
        self.assertEqual(book["Gestiones"]["C5"].value, "B")

    def test_monthly_does_not_report_missing_deductions_as_zero(self):
        period = {"employer": "A", "month": "2026-09", "reconciliation_mode": "Operativa",
                  "expected_usd": 100, "deducted_usd": 0,
                  "rows": [{"expected_usd": 100, "employee_receivable_usd": None}]}
        row, = monthly_rows({"periods": [period]})
        self.assertEqual(row[14:], (100, "N/D", "N/D", 100))
        period["rows"].append({"expected_usd": 50, "deducted_usd": 40, "employee_receivable_usd": 10})
        period["expected_usd"] = 150
        row, = monthly_rows({"periods": [period]})
        self.assertEqual(row[14:], (150, 40, 10, 100))


if __name__ == "__main__":
    unittest.main()
