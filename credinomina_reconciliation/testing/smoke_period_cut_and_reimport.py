"""Rollback-only Frappe smoke: repeat imports without losing a managed exception.

Run with: bench --site cn-reconciliation-test.local execute
    credinomina_reconciliation.testing.smoke_period_cut_and_reimport.run
"""

from __future__ import annotations

import io
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4

import frappe
from openpyxl import load_workbook

from credinomina_reconciliation.conciliacion_credinomina.doctype.cn_reconciliation_period import (
    cn_reconciliation_period as period_module,
)
from credinomina_reconciliation.conciliacion_credinomina.doctype.cn_accounting_import import (
    cn_accounting_import as source_module,
)
from credinomina_reconciliation.templates import build_template_xlsx


def _collection(amount, first_name="Ana Pérez"):
    workbook = load_workbook(io.BytesIO(build_template_xlsx("cobranza")))
    sheet = workbook.active
    sheet.append(["C-101", "E-101", first_name, "001101010001A", "L-101", "1", "12", amount, 0, "", "", "", ""])
    sheet.append(["C-102", "E-102", "Luis Gómez", "001102020002B", "L-102", "1", "12", 40, 0, "", "", "", ""])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _response(rows, first_deduction):
    workbook = load_workbook(io.BytesIO(build_template_xlsx(
        "empresa", (row.as_dict() for row in rows),
    )))
    sheet = workbook.active
    sheet.cell(2, 14, first_deduction)
    sheet.cell(3, 14, 40)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _alias_only_response(rows, alias, first_deduction):
    workbook = load_workbook(io.BytesIO(_response(rows, first_deduction)))
    sheet = workbook.active
    for column in (1, 2, 4, 5, 15):
        sheet.cell(2, column).value = None
    sheet.cell(2, 3, alias)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def run():
    label = uuid4().hex[:10]
    frappe.set_user("Administrator")
    try:
        employer = frappe.get_doc({
            "doctype": "CN Employer", "employer_name": f"Smoke Cut {label}",
            "employer_code": f"SC{label}", "payroll_frequency": "Mensual",
        }).insert(ignore_permissions=True)
        period = frappe.get_doc({
            "doctype": "CN Reconciliation Period", "employer": employer.name,
            "payroll_month": "2026-09-01", "reconciliation_mode": "Operativa",
            "collection_cycle": "Mensual", "collection_file": "smoke_collection.xlsx",
        }).insert(ignore_permissions=True)
        files = {"collection": _collection(50), "response": b""}

        def attached(document, url):
            if url == "smoke_collection.xlsx":
                return SimpleNamespace(file_name="collection.xlsx"), files["collection"]
            return SimpleNamespace(file_name="response.xlsx"), files["response"]

        with patch.object(period_module, "_attached_file", side_effect=attached), \
             patch.object(period_module, "_reconcile_if_sources", return_value=None) as reconcile:
            period_module.import_collection(period.name)
            assert reconcile.call_count == 1
            period.reload()
            original_names = [row.name for row in period.collection_rows]
            assert len(original_names) == 2

            unchanged = period_module.import_collection(period.name)
            assert unchanged["unchanged"]
            assert reconcile.call_count == 2
            period.reload()
            assert [row.name for row in period.collection_rows] == original_names

            files["response"] = _response(period.collection_rows, 30)
            period.employer_response_file = "smoke_response.xlsx"
            period.deduction_evidence_date = "2026-09-30"
            period.save()
            period_module.import_employer_response(period.name)
            period.reload()
            assert period.status == "Pendiente", period.status
            exceptions = frappe.get_all(
                "CN Reconciliation Exception", filters={"period": period.name},
                fields=["name", "exception_type", "exception_key"],
            )
            assert len(exceptions) == 1 and exceptions[0].exception_type == "Deduccion parcial", exceptions
            exception = frappe.get_doc("CN Reconciliation Exception", exceptions[0].name)
            exception.status = "En revision"
            exception.assigned_to = "Administrator"
            exception.next_action = "Solicitar aclaración de retención"
            exception.commitment_date = "2026-10-10"
            exception.append("follow_up_actions", {
                "action_type": "Contacto con empresa", "details": "Empresa confirmó retención parcial",
            })
            exception.save(ignore_permissions=True)

            no_change = period_module.import_employer_response(period.name)
            assert no_change["unchanged"]
            reloaded = frappe.get_doc("CN Reconciliation Exception", exception.name)
            assert reloaded.status == "En revision"
            assert len(reloaded.follow_up_actions) == 1
            assert frappe.db.count("CN Reconciliation Exception", {"period": period.name}) == 1

            files["response"] = _response(period.collection_rows, 35)
            period_module.import_employer_response(period.name)
            reloaded = frappe.get_doc("CN Reconciliation Exception", exception.name)
            assert reloaded.status == "En revision"
            assert len(reloaded.follow_up_actions) == 1
            assert round(reloaded.amount_usd, 2) == 35

            files["collection"] = _collection(55)
            period_module.import_collection(period.name)
            period.reload()
            assert [row.name for row in period.collection_rows] == original_names
            reloaded = frappe.get_doc("CN Reconciliation Exception", exception.name)
            assert reloaded.status == "En revision"
            assert len(reloaded.follow_up_actions) == 1

            detail_key_before = period.employer_response_import_key
            files["collection"] = _collection(55, first_name="Ana P. Pérez")
            period_module.import_collection(period.name)
            period.reload()
            assert [row.name for row in period.collection_rows] == original_names
            assert period.employer_response_import_key != detail_key_before
            reloaded = frappe.get_doc("CN Reconciliation Exception", exception.name)
            assert reloaded.status == "En revision"
            assert len(reloaded.follow_up_actions) == 1

            cut = period_module.record_control_cut(
                period.name, "Seguir saldo parcial con la empresa y el empleado"
            )
            assert cut["status"] == "Pendiente"
            period.reload()
            assert period.control_cut_on and "Cobranza no deducida" in period.control_cut_summary
            assert len([row for row in period.collection_rows if row.deduction_status == "Deduccion parcial"]) == 1

            resolved = frappe.get_doc("CN Reconciliation Exception", exception.name)
            resolved.status = "Resuelta"
            resolved.cause_category = "Ingreso insuficiente"
            resolved.resolution = "La empresa confirmó el subsidio; gestionar al retorno"
            resolved.save(ignore_permissions=True)

            files["response"] = _response(period.collection_rows, 55)
            period_module.import_employer_response(period.name)
            period.reload()
            assert period.status == "Pendiente"
            obsolete = frappe.get_doc("CN Reconciliation Exception", exception.name)
            assert obsolete.status == "Descartada"
            assert len(obsolete.follow_up_actions) == 2
            assert "gestionar al retorno" in obsolete.follow_up_actions[-1].details

            files["response"] = _response(period.collection_rows, 35)
            period_module.import_employer_response(period.name)
            reopened = frappe.get_doc("CN Reconciliation Exception", exception.name)
            assert reopened.status == "Abierta"
            assert reopened.resolution == ""
            assert len(reopened.follow_up_actions) == 3
            assert frappe.db.count("CN Reconciliation Exception", {"period": period.name}) == 1

            alias = "Ana Planilla Alterada"
            files["response"] = _alias_only_response(period.collection_rows, alias, 35)
            first_attempt = period_module.import_employer_response(period.name)
            assert first_attempt["unmatched"] == 1
            assert first_attempt["missing"] == 1
            client = frappe.get_doc("CN Client", period.collection_rows[0].client)
            client.append("aliases", {"alias_name": alias})
            client.save(ignore_permissions=True)
            second_attempt = period_module.import_employer_response(period.name)
            assert not second_attempt.get("unchanged")
            assert second_attempt["unmatched"] == second_attempt["missing"] == 0
            period.reload()
            assert period.collection_rows[0].deduction_status == "Deduccion parcial"
            assert frappe.get_doc("CN Reconciliation Exception", exception.name).name == exception.name

            managed = frappe.get_doc("CN Reconciliation Exception", exception.name)
            managed.status = "En revision"
            managed.assigned_to = "Administrator"
            managed.next_action = "Confirmar por qué no se retuvo la cuota"
            managed.commitment_date = "2026-10-20"
            # Emulate a case created by the older type-specific key scheme.
            managed.exception_key = period_module._deduction_exception_key(
                period.name, period.collection_rows[0].row_key, "Deduccion parcial"
            )
            managed.save(ignore_permissions=True)
            files["response"] = _response(period.collection_rows, 0)
            period_module.import_employer_response(period.name)
            reclassified = frappe.get_doc("CN Reconciliation Exception", exception.name)
            assert reclassified.exception_type == "No deducido"
            assert reclassified.exception_key == period_module._deduction_exception_key(
                period.name, period.collection_rows[0].row_key
            )
            assert reclassified.status == "En revision"
            assert reclassified.assigned_to == "Administrator"
            assert len(reclassified.follow_up_actions) >= 3
            assert frappe.db.count("CN Reconciliation Exception", {
                "period": period.name, "status": ["in", ["Abierta", "En revision"]],
            }) == 1
            print({
                "period": period.name, "row_names_preserved": True,
                "exception": exception.name,
                "reclassified_without_losing_actions": True,
                "alias_retried_without_file_change": True,
                "control_cut": str(period.control_cut_on), "rolled_back": True,
            })
    finally:
        frappe.db.rollback()


def run_remittance_before_files():
    """A submitted remittance must be reconsidered when collection/detail arrive."""
    label = uuid4().hex[:10]
    frappe.set_user("Administrator")
    try:
        employer = frappe.get_doc({
            "doctype": "CN Employer", "employer_name": f"Smoke Late {label}",
            "employer_code": f"SL{label}", "payroll_frequency": "Mensual",
        }).insert(ignore_permissions=True)
        period = frappe.get_doc({
            "doctype": "CN Reconciliation Period", "employer": employer.name,
            "payroll_month": "2026-09-01", "reconciliation_mode": "Operativa",
            "collection_cycle": "Mensual", "collection_file": "late_collection.xlsx",
        }).insert(ignore_permissions=True)
        remittance = frappe.get_doc({
            "doctype": "CN Remittance Allocation", "employer": employer.name,
            "deposit_reference": f"DEP-{label}", "deposit_voucher": f"V-{label}",
            "deposit_date": "2026-10-15", "deposit_currency": "USD",
            "deposit_amount": 30, "detail_periods": [{"period": period.name}],
            "detail_file": "late_detail.xlsx", "detail_source_file": "late_detail.xlsx",
            "detail_hash": f"smoke-{label}",
            "notes": "Depósito y detalle recibidos antes de cargar la cobranza",
            "detail_rows": [{
                "source_row": 2, "client_number": "C-101", "client_name": "Ana Pérez",
                "loan_number": "L-101", "installment_number": "1",
                "deducted_usd": 30,
            }],
        }).insert(ignore_permissions=True)
        remittance.submit()
        remittance.reload()
        before_status = remittance.detail_rows[0].match_status
        assert before_status != "Conciliada", before_status

        files = {"collection": _collection(50), "response": b""}

        def attached(document, url):
            if url == "late_collection.xlsx":
                return SimpleNamespace(file_name="collection.xlsx"), files["collection"]
            return SimpleNamespace(file_name="response.xlsx"), files["response"]

        original_reconcile = source_module._reconcile_sources
        with patch.object(period_module, "_attached_file", side_effect=attached), \
             patch.object(source_module, "_reconcile_sources", wraps=original_reconcile) as reconcile:
            collection_result = period_module.import_collection(period.name)
            assert collection_result["source_reconciliation"] is not None
            period.reload()
            files["response"] = _response(period.collection_rows, 30)
            period.employer_response_file = "late_response.xlsx"
            period.deduction_evidence_date = "2026-09-30"
            period.save()
            response_result = period_module.import_employer_response(period.name)
            assert response_result["source_reconciliation"] is not None
            assert reconcile.call_count >= 2

        remittance.reload()
        period.reload()
        assert remittance.detail_rows[0].match_status == "Conciliada", (
            remittance.detail_rows[0].match_status,
            remittance.detail_rows[0].match_reason,
        )
        assert round(period.collection_rows[0].remitted_usd, 2) == 30
        print({
            "period": period.name, "remittance_first": remittance.name,
            "detail_after_reconcile": remittance.detail_rows[0].match_status,
            "remitted_usd": period.collection_rows[0].remitted_usd,
            "rolled_back": True,
        })
    finally:
        frappe.db.rollback()
