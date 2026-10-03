"""Read-only demo-site smoke for the control-page Excel download."""

from __future__ import annotations

import io
from collections import defaultdict

import frappe
from openpyxl import load_workbook
from credinomina_reconciliation.rounding import money, sum_money

from credinomina_reconciliation.conciliacion_credinomina.page.control_credinomina.control_credinomina import (
    export_control_excel,
)


def run(year=2025, employer=None):
    frappe.local.response = frappe._dict()
    export_control_excel(year=year, employer=employer)
    response = frappe.local.response
    assert response.type == "download"
    assert response.filename.endswith(".xlsx")
    book = load_workbook(io.BytesIO(response.filecontent), read_only=True)
    assert book.sheetnames == [
        "Resumen mensual", "Períodos", "Detalle cliente", "Cruces", "Depósitos", "Distribución depósitos",
        "Aplicaciones sin período", "Partidas y excepciones", "Gestiones", "Guía",
    ]
    summary = book["Períodos"]
    period_rows = list(summary.iter_rows(min_row=5, values_only=True))
    assert period_rows, "No se encontraron períodos de prueba."
    if employer:
        names = {row[1] for row in period_rows}
        assert len(names) == 1, "El filtro de empresa mostró otros convenios."
    historical = [row for row in period_rows if row[3] == "Histórica"]
    assert historical, "Falta cobertura histórica."
    assert all(row[9] == "N/D" and row[10] == "N/D" for row in historical)
    detail_count = max(book["Detalle cliente"].max_row - 4, 0)
    link_count = max(book["Cruces"].max_row - 4, 0)
    assert detail_count > 0
    filters = {"docstatus": 1}
    if year not in (None, "Todos"):
        filters["deposit_date"] = ["between", [f"{year}-01-01", f"{year}-12-31"]]
    if employer:
        filters["employer"] = employer
    deposits = frappe.get_list("CN Remittance Allocation", filters=filters,
        fields=["name", "amount_usd"], limit_page_length=0)
    cash = {row[0]: money(row[7]) for row in book["Depósitos"].iter_rows(min_row=5, values_only=True)}
    assert cash == {d.name: money(d.amount_usd) for d in deposits}, "Depósitos omitidos o duplicados."
    distribution = defaultdict(lambda: money(0))
    for row in book["Distribución depósitos"].iter_rows(min_row=5, values_only=True):
        distribution[row[0]] += money(row[10])
    assert all(distribution[name] == amount for name, amount in cash.items()), "La distribución no explica el depósito completo."
    monthly = [row for row in book["Resumen mensual"].iter_rows(min_row=5, values_only=True)
               if row[1] is not None]  # Ignore explanatory notes below the table.
    assert len({(row[0], row[1]) for row in monthly}) == len(monthly)
    assert sum_money(row[11] for row in monthly) == sum_money(cash.values())
    assert sum_money(row[4] for row in monthly) == sum_money(row[11] for row in period_rows)
    return {
        "filename": response.filename,
        "periods": len(period_rows),
        "historical_periods": len(historical),
        "detail_rows": detail_count,
        "linked_rows": link_count,
        "issue_rows": max(book["Partidas y excepciones"].max_row - 4, 0),
        "action_rows": max(book["Gestiones"].max_row - 4, 0),
        "bytes": len(response.filecontent),
        "monthly_rows": len(monthly), "deposits": len(cash), "cash_tie_out": "OK",
    }


def run_scoped(year=2025):
    period = frappe.get_all(
        "CN Reconciliation Period",
        filters={"payroll_month": ["between", [f"{year}-01-01", f"{year}-12-31"]]},
        fields=["employer"], limit_page_length=1,
    )[0]
    return run(year=year, employer=period.employer)
