"""Rollback-only company credit, reservation and reconciliation audit acceptance."""
import json
import hashlib
import io
from openpyxl import load_workbook
from unittest.mock import patch
import frappe
from credinomina_reconciliation import client_credit
from credinomina_reconciliation.control_deposits import get_cash_deposits
from credinomina_reconciliation.deposit_reconciliation import reconcile_deposit
from credinomina_reconciliation.reconciliation_audit import get_history
from credinomina_reconciliation.conciliacion_credinomina.doctype.cn_accounting_import.cn_accounting_import import _reconcile_sources
from credinomina_reconciliation.conciliacion_credinomina.doctype.cn_remittance_allocation.cn_remittance_allocation import create_complementary_item


def run():
    if frappe.local.site != "cn-reconciliation-test.local":
        raise RuntimeError("Solo para cn-reconciliation-test.local")
    frappe.set_user("Administrator")
    marker = "COMPANY-CREDIT-" + frappe.generate_hash(length=8)
    try:
        with patch.object(frappe, "enqueue"), patch.object(frappe, "publish_realtime"):
            employer = frappe.get_doc({"doctype": "CN Employer", "employer_name": marker, "employer_code": marker}).insert()
            period = frappe.get_doc({"doctype": "CN Reconciliation Period", "employer": employer.name,
                "payroll_month": "2025-04-01", "reconciliation_mode": "Historica", "historical_scope": "Fecha exacta",
                "historical_application_date": "2025-04-30"}).insert()
            source = frappe.get_doc({"doctype": "CN Accounting Import", "employer": employer.name,
                "currency": "USD", "source_file": f"/private/files/{marker}.csv", "status": "Importado"})
            source.append("rows", {"event_type": "Aplicacion", "event_date": "2025-04-30", "source_key": marker,
                "client_name": "Cliente prueba empresa", "client_number": marker, "loan_number": "987650001-1",
                "amount": 80, "amount_usd": 80, "currency": "USD", "effective": 1,
                "historical_period": period.name, "processing_route": "Historica"})
            source.insert()
            _reconcile_sources(employer.name)
            deposit = frappe.get_doc({"doctype": "CN Remittance Allocation", "employer": employer.name,
                "deposit_reference": marker, "deposit_date": "2025-05-10", "deposit_currency": "USD", "deposit_amount": 100,
                "targets": [{"historical_application": source.rows[0].name, "amount_usd": 80}],
                "detail_periods": [{"period": period.name}]}).insert()
            deposit.submit()
            reconcile_deposit(deposit)
            deposit.reload()
            created = create_complementary_item(deposit.name, str(deposit.modified), {
                "category": "Saldo a favor de la empresa", "amount": 20, "currency": "USD", "posting_date": "2025-05-10",
                "reason_type": "Error de la empresa", "description": "Excedente para devolución a la empresa",
                "credit_assigned_to": "Administrator", "credit_commitment_date": "2026-10-10", "credit_treatment": "Devolución"})
            item = frappe.get_doc("CN Complementary Item", created["name"])
            deposit.reload()
            assert item.result == "Saldo a favor documentado", item.result
            from credinomina_reconciliation.complementary_balances import get_balance
            from credinomina_reconciliation.follow_up_queue import load_follow_up
            position = get_balance(item.name)
            assert position["pending_usd"] == 0 and position["management_pending_usd"] == 20
            tasks = load_follow_up(None, employer.name)
            assert any(task["target_name"] == item.name and task["kind"] == "credit_management" for task in tasks)
            assert (deposit.allocated_usd, deposit.justified_surplus_usd, deposit.unclassified_usd) == (80, 20, 0)
            initial = get_history(deposit.name)
            assert len(initial["rows"]) >= 2
            reconcile_deposit(deposit)
            assert len(get_history(deposit.name)["rows"]) == len(initial["rows"]), "Idempotent run added audit noise"
            from credinomina_reconciliation.conciliacion_credinomina.doctype.cn_reconciliation_period.cn_reconciliation_period import record_control_cut
            from credinomina_reconciliation.control_cuts import get_cuts
            cut = record_control_cut(period.name, "Corte antes de recibir una aplicación posterior")
            saved_json = frappe.get_doc("File", cut["files"]["json"]["file"]).get_content()
            if isinstance(saved_json, str):
                saved_json = saved_json.encode("utf-8")
            assert hashlib.sha256(saved_json).hexdigest() == cut["files"]["json"]["sha256"]
            saved_cut = json.loads(saved_json)
            assert saved_cut["period_document"]["applied_usd"] == 80
            workbook = load_workbook(io.BytesIO(frappe.get_doc("File", cut["files"]["xlsx"]["file"]).get_content()), read_only=True)
            assert "Antigüedad guardada" in workbook.sheetnames
            workbook.close()
            # Later claims must not take money already promised to the company.
            source.reload()
            source.append("rows", {"event_type": "Aplicacion", "event_date": "2025-04-30", "source_key": marker + "-LATE",
                "client_name": "Otro cliente", "client_number": marker + "2", "loan_number": "987650002-1",
                "amount": 20, "amount_usd": 20, "currency": "USD", "effective": 1,
                "historical_period": period.name, "processing_route": "Historica"})
            source.save()
            _reconcile_sources(employer.name)
            source.reload(); deposit.reload(); item.reload()
            assert source.rows[1].historical_remitted_usd == 0
            assert str(source.rows[0].payment_due_date) == "2025-05-10"
            cut_history = get_cuts(period.name)
            assert cut_history["rows"][0]["files"] == cut["files"]
            unchanged = frappe.get_doc("File", cut["files"]["json"]["file"]).get_content()
            assert json.loads(unchanged)["period_document"]["applied_usd"] == 80
            assert deposit.allocated_usd == 80 and item.result == "Saldo a favor documentado"
            support = frappe.get_doc({"doctype": "File", "file_name": marker + ".pdf", "file_url": "https://example.invalid/" + marker,
                "attached_to_doctype": item.doctype, "attached_to_name": item.name}).insert()
            result = client_credit.record_management(item.name, str(item.modified), "Devolución", 5, "2025-05-20", "REFUND", support.file_url)
            assert result["pending_usd"] == 15
            item.reload()
            client_credit.record_management(item.name, str(item.modified), "Devolución", 15, "2025-05-21", "REFUND-2", support.file_url)
            reconcile_deposit(deposit)
            item.reload(); deposit.reload()
            assert item.credit_pending_usd == 0 and len(json.loads(item.credit_history)) == 2
            assert not any(task["target_name"] == item.name and task["kind"] == "credit_management"
                           for task in load_follow_up(None, employer.name))
            assert deposit.allocated_usd == 80 and deposit.justified_surplus_usd == 20
            overview, = get_cash_deposits(None, deposit_name=deposit.name)
            assert overview["settled"] and overview["company_credit_pending_usd"] == 0
            assert sum(part["amount_usd"] for part in overview["destinations"]) == 100
            return {"company_reservation": True, "late_application_not_funded": True, "partial_and_full_refund": True,
                "cash_not_released": True, "audit_before_after": True, "idempotent_audit": True,
                "cut_unchanged_after_late_data": True, "private_excel_and_json": True, "rolled_back": True}
    finally:
        frappe.db.rollback()
