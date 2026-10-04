"""Snapshot workbook for the Credinómina control page.

The source of each monetary amount is the reconciliation state already stored
on a period, claim or application. A deposit-to-claim link is emitted once per
actual allocation, so neither many-to-many deposits nor quincenas are reduced
to a first reference match.
"""

from __future__ import annotations

import io
import json
import math
from datetime import date, datetime
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from credinomina_reconciliation.reconciliation import net_application_amount
from credinomina_reconciliation.rounding import CASH_EPSILON, money, money_float, sum_money


NA = "N/D"
MAX_EXCEL_ROWS = 1_048_576
NAVY = "1F4E78"
TEXT = "263547"
MONEY_FORMAT = '#,##0.00;[Red](#,##0.00);"-"'
DATE_FORMAT = "dd/mm/yyyy"


def build_control_workbook(
    data: dict, *, exceptions: list, actions: list, employer_label: str,
    generated_at: datetime, date_format: str = DATE_FORMAT,
) -> bytes:
    """Return a read-only-in-practice, value-based audit snapshot as XLSX bytes."""
    book = Workbook()
    summary = book.active
    summary.title = "Resumen mensual"
    periods = data["periods"]
    display_format = date_format.replace("yyyy", "%Y").replace("mm", "%m").replace("dd", "%d")
    scope = f"Año: {data.get('year') or 'Todos'}    Empresa: {employer_label}    Generado: {generated_at.strftime(display_format + ' %H:%M')}"
    _write_monthly(summary, data, scope)
    _write_summary(book.create_sheet("Períodos"), data, employer_label, generated_at)

    detail_rows = []
    link_rows = []
    issue_rows = []
    period_by_name = {period["name"]: period for period in periods}
    for period in periods:
        historical = period.get("reconciliation_mode") == "Historica"
        for claim in period.get("rows") or []:
            detail_rows.append(_collection_detail(period, claim))
            link_rows.extend(_claim_links(period, claim, historical=False))
            due = claim.get("collection_shortfall_usd")
            if due is None and _money(claim.get("expected_usd")) > CASH_EPSILON:
                issue_rows.append(_issue(
                    "Detalle de deducción pendiente", period,
                    claim.get("client_number"), claim.get("loan_number"),
                    claim.get("application_reference"), _money(claim.get("expected_usd")),
                    "Pendiente", detail="Falta confirmar la deducción; la cobranza no es una cuenta por cobrar.",
                ))
            elif due is not None and _money(due) > CASH_EPSILON:
                issue_rows.append(_issue(
                    "Cobranza no deducida", period, claim.get("client_number"),
                    claim.get("loan_number"), claim.get("application_reference"),
                    _money(due), "Pendiente", detail=claim.get("deduction_match_note"),
                ))
            pending = _collection_pending(claim)
            if pending > CASH_EPSILON:
                issue_rows.append(_issue(
                    "Aplicación sin depósito", period, claim.get("client_number"),
                    claim.get("loan_number"), claim.get("application_reference"), pending,
                    "Pendiente", detail=claim.get("application_status"),
                ))
        for application in period.get("historical_rows") or []:
            detail_rows.append(_historical_detail(period, application))
            link_rows.extend(_claim_links(period, application, historical=True))
            balance = _money(application.get("historical_balance_usd"))
            if balance > CASH_EPSILON:
                issue_rows.append(_issue(
                    "Aplicación sin depósito", period,
                    application.get("client_number"), application.get("loan_number"),
                    application.get("reference"), balance,
                    application.get("deposit_match_status") or "Pendiente",
                    detail=application.get("deposit_match_reason"),
                ))
        assignment_gap = _money(period.get("employer_gap_usd"))
        if not historical and assignment_gap > CASH_EPSILON:
            issue_rows.append(_issue(
                "Deducido sin depósito asignado (no CxC confirmada)", period,
                None, None, None, assignment_gap, "Pendiente de asignación",
                detail="Puede existir un depósito recibido sin detalle suficiente para asignarlo.",
            ))
        for surplus in period.get("surpluses") or []:
            issue_rows.append(_issue(
                "Excedente documentado", period, None, None,
                surplus.get("deposit_reference"), _money(surplus.get("amount_usd")),
                surplus.get("result"), cause=surplus.get("reason_type"),
                detail=surplus.get("explanation"),
            ))

    for exception in exceptions:
        period = period_by_name.get(exception.get("period"))
        issue_rows.append(_issue(
            "Excepción", period, exception.get("client_number"),
            exception.get("loan_number"), exception.get("name"),
            _money(exception.get("amount_usd")), exception.get("status"),
            cause=exception.get("cause_category") or exception.get("exception_type"),
            owner=exception.get("assigned_to"), action=exception.get("next_action"),
            commitment=exception.get("commitment_date"),
            evidence=exception.get("evidence_file"),
            external_reference=exception.get("external_reference"),
            detail=exception.get("description"), resolution=exception.get("resolution"),
            employer=exception.get("employer"), event_date=exception.get("creation"),
        ))
    for deposit in data.get("open_deposits") or []:
        unclassified = _money(deposit.get("unclassified_usd"))
        if unclassified <= CASH_EPSILON:
            continue
        issue_rows.append(_issue(
            "Depósito sin clasificar", None, None, None,
            deposit.get("reference"), unclassified, "Pendiente de clasificación",
            employer=deposit.get("employer_text"),
            detail=deposit.get("allocation_reason"),
            event_date=deposit.get("event_date"),
        ))
    for surplus in data.get("unassigned_surpluses") or []:
        issue_rows.append(_issue(
            "Saldo a favor documentado sin período", None, None, None,
            surplus.get("deposit_reference"), _money(surplus.get("amount_usd")),
            surplus.get("result"), employer=surplus.get("employer"),
            cause=surplus.get("reason_type"), detail=surplus.get("explanation"),
        ))
    for application in _unassigned(data):
        issue_rows.append(_issue(
            "Aplicación sin período", None, application.get("client_number"),
            application.get("loan_number"), application.get("reference"),
            net_application_amount(application), "Sin período",
            detail=application.get("match_reason"),
            event_date=application.get("event_date"),
            employer=application.get("employer"),
        ))
    issue_rows = [
        (*row, _commitment_delay(row, generated_at.date()), None, None, None, None) for row in issue_rows
    ]
    for item in data.get('core_complementary_items') or []:
        balance = item['_balance']
        row = _issue('Partida complementaria del core', None, item.get('client_number'),
            item.get('loan_number'), item.get('name'), item.get('_signed_pending_usd'),
            balance.get('financial_status'), employer=item.get('employer'),
            cause=item.get('category'), event_date=item.get('source_date') or item.get('posting_date'),
            external_reference=item.get('source_voucher') or item.get('voucher'),
            evidence=item.get('source_file'),
            detail=item.get('source_description') or item.get('description'))
        issue_rows.append((*row, None, item.get('source_debit'), item.get('source_credit'),
                           item.get('source_currency'), balance.get('used_usd')))

    _write_table(
        book.create_sheet("Detalle cliente"), "Detalle por cliente y aplicación", scope,
        DETAIL_HEADERS, detail_rows,
        money_columns={14, 15, 16, 17, 18, 19, 20, 21, 26, 27},
        date_columns={10},
    )
    _write_table(
        book.create_sheet("Cruces"), "Cruces aplicación y depósito", scope,
        LINK_HEADERS, link_rows, money_columns={13, 14}, date_columns={11},
    )
    _write_deposits(book, data, scope)
    _write_table(
        book.create_sheet("Aplicaciones sin período"), "Aplicaciones pendientes de vincular a un período", scope,
        ("Importación", "Fila", "Empresa", "Fecha aplicación", "N.º cliente", "Nombre del cliente",
         "N.º crédito", "Asiento contable", "Recibo", "Referencia", "Aplicado neto USD", "Motivo"),
        [(r.get("parent"), r.get("source_row"), r.get("employer"), _date(r.get("event_date")),
          r.get("client_number"), r.get("client_name"), r.get("loan_number"), r.get("accounting_entry"),
          r.get("receipt"), r.get("reference"), net_application_amount(r), r.get("match_reason"))
         for r in _unassigned(data)], money_columns={11}, date_columns={4},
        total_columns={11},
    )
    _write_table(
        book.create_sheet("Partidas y excepciones"), "Partidas y excepciones", scope,
        ISSUE_HEADERS, issue_rows, money_columns={8, 20, 21, 23}, date_columns={2, 13, 18},
    )
    _write_table(
        book.create_sheet("Gestiones"), "Gestiones de excepciones", scope,
        ACTION_HEADERS, _action_rows(actions, exceptions, period_by_name),
        date_columns={4},
    )
    _write_guide(book, scope)
    if "aging_rows" in data:
        _write_table(book.create_sheet("Antigüedad guardada"), "Saldos guardados de este período al registrar el corte",
            scope + " · No reconstruye fechas anteriores ni representa todo el saldo de la empresa.",
            ("Importación", "Cliente", "Nro. cliente", "Crédito", "Fecha aplicación", "Vencimiento conservado",
             "Origen del plazo", "Aplicado US$", "Depositado US$", "Ajuste US$", "Pendiente US$", "Días de atraso", "Observación"),
            [(r.get("source_import"), r.get("client_name"), r.get("client_number"), r.get("loan_number"),
              _date(r.get("application_date")), _date(r.get("due_date")), r.get("payment_term_origin"),
              r.get("applied_usd"), r.get("paid_usd"), r.get("adjustment_usd"), r.get("amount_usd"), r.get("age_days"),
              r.get("observation")) for r in data["aging_rows"]], money_columns={8, 9, 10, 11}, date_columns={5, 6})
    # System Settings uses the same numeric day/month/year tokens as Excel.
    # Keep native dates (sortable/filterable), changing only their display style.
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, datetime):
                    cell.number_format = f"{date_format} hh:mm"
                elif isinstance(cell.value, date) and cell.number_format != "yyyy-mm":
                    cell.number_format = date_format
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


SUMMARY_HEADERS = (
    "Período", "Empresa", "Mes cobranza", "Modalidad", "Alcance del período",
    "Fecha inicio aplicación", "Fecha fin aplicación", "Vencimiento pago",
    "Resultado", "Cobranza USD", "Deducido USD", "Aplicado neto USD",
    "Complementario USD", "Ajuste de redondeo USD", "Depósito asignado total USD",
    "Aplicado pendiente de depósito USD", "Cobranza no deducida USD", "Deducido sin depósito asignado USD",
    "Depósito asignado a créditos USD", "Excepciones abiertas", "Último corte de control",
    "Motivo y próxima gestión", "Observaciones", "Cierre", "Cobranza sin detalle USD",
)

DETAIL_HEADERS = (
    "Período", "Empresa", "Modalidad", "N.º cliente", "N.º empleado",
    "Nombre del cliente", "Cédula", "N.º crédito", "N.º cuota",
    "Fecha aplicación", "Asiento contable", "Recibo", "Referencia aplicación",
    "Cobranza USD", "Deducido USD", "Aplicado neto USD", "Complementario USD",
    "Ajuste de redondeo USD", "Depósito asignado a créditos USD", "Cobranza no deducida USD",
    "Aplicado pendiente de depósito USD", "Estado", "Comentarios",
    "Importación contable", "Fila de origen", "Aplicación bruta USD", "Reducción confirmada USD",
)

LINK_HEADERS = (
    "Período", "Empresa", "Modalidad", "N.º cliente", "Nombre del cliente",
    "N.º crédito", "ID partida", "Referencia aplicación",
    "Referencia depósito", "Comprobante depósito", "Fecha depósito", "Destino",
    "Monto vinculado USD", "Ajuste USD", "Origen", "Comentario heredado",
)

ISSUE_HEADERS = (
    "Tipo", "Fecha", "Período", "Empresa", "N.º cliente", "N.º crédito",
    "Referencia / ID", "Monto USD", "Estado", "Causa", "Responsable",
    "Próxima acción", "Compromiso", "Evidencia", "Referencia externa",
    "Detalle", "Resolución", "Vencimiento contractual", "Días atraso compromiso",
    "Débito original core", "Crédito original core", "Moneda original core", "Utilizado partida USD",
)

ACTION_HEADERS = (
    "Excepción", "Período", "Empresa", "Fecha gestión", "Usuario",
    "Tipo de gestión", "Detalle", "Evidencia", "Referencia externa",
)


MONTHLY_HEADERS = (
    "Empresa", "Mes", "Períodos", "Períodos cerrados", "Aplicado neto en períodos USD",
    "Depósito asignado a créditos USD", "Ajuste de redondeo USD", "Pendiente USD",
    "Excepciones abiertas de períodos", "Aplicado sin período USD", "Depósitos recibidos",
    "Depositado completo USD", "Depósito sin clasificar USD", "Depósitos por revisar",
    "Cobranza USD", "Deducido USD", "Cobranza no deducida conocida USD", "Cobranza sin detalle USD",
)


def _sum(values):
    return money_float(sum_money(values))


def _month(value):
    return _date(str(value)[:7] + "-01") if value else None


def _unassigned(data):
    rows = [*(data.get("unassigned_historical_applications") or []),
            *(data.get("unassigned_operational_applications") or [])]
    # Stable row IDs distinguish legitimate equal accounting lines.
    result, seen = [], set()
    for row in rows:
        if row.get("name") and row["name"] in seen:
            continue
        if row.get("name"):
            seen.add(row["name"])
        result.append(row)
    return result


def _employee_due(period):
    known = [r["collection_shortfall_usd"] for r in period.get("rows") or []
             if r.get("collection_shortfall_usd") is not None]
    return _sum(known) if known else NA


def _missing_detail(period):
    return _sum(r.get("expected_usd") for r in period.get("rows") or []
                if r.get("collection_shortfall_usd") is None)


def _known_deductions(period):
    known = [r.get("deducted_usd") for r in period.get("rows") or []
             if r.get("collection_shortfall_usd") is not None]
    return _sum(known) if known else NA


def monthly_rows(data):
    """One employer/month, not one row per deposit or payroll cut.

    Payroll assignment and cash receipt have independent dates and employers.
    Never infer a receivable by subtracting the two monthly populations.
    """
    groups = defaultdict(lambda: {"periods": [], "deposits": [], "unassigned": []})
    labels = {}
    for p in data["periods"]:
        key = p.get("employer") or p.get("employer_name") or NA
        labels[key] = p.get("employer_name") or key
        groups[key, str(p.get("month") or p.get("payroll_month") or "")[:7]]["periods"].append(p)
    for d in {d["name"]: d for d in data.get("cash_deposits") or []}.values():
        groups[d.get("employer") or NA, str(d.get("month") or d.get("date") or "")[:7]]["deposits"].append(d)
    for a in _unassigned(data):
        groups[a.get("employer") or NA, str(a.get("event_date") or "")[:7]]["unassigned"].append(a)
    result = []
    for (employer, month), group in sorted(groups.items()):
        periods, deposits, unassigned = group["periods"], group["deposits"], group["unassigned"]
        collections = [p for p in periods if p.get("reconciliation_mode") != "Historica"]
        known_due = [_employee_due(p) for p in collections if _employee_due(p) != NA]
        known_deductions = [_known_deductions(p) for p in collections if _known_deductions(p) != NA]
        result.append((
            labels.get(employer, employer), _month(month), len(periods),
            sum(p.get("status") == "Cerrado" for p in periods),
            _sum(p.get("applied_usd") for p in periods),
            _sum(_period_credit_cash(p) for p in periods),
            _sum(p.get("rounding_adjustment_usd") for p in periods),
            _sum(_period_pending(p) for p in periods),
            len({e["name"] for p in periods for e in p.get("exceptions") or [] if e.get("name")}),
            _sum(net_application_amount(a) for a in unassigned),
            len(deposits), _sum(d.get("total_usd") for d in deposits),
            _sum(max(money(d.get("unclassified_usd")), 0) for d in deposits),
            sum(bool(d.get("needs_review")) for d in deposits),
            _sum(p.get("expected_usd") for p in collections) if collections else NA,
            _sum(known_deductions) if known_deductions else NA,
            _sum(known_due) if known_due else NA,
            _sum(_missing_detail(p) for p in collections) if collections else NA,
        ))
    return result


def _write_monthly(sheet, data, scope):
    _write_table(sheet, "Resumen por empresa y mes", scope, MONTHLY_HEADERS, monthly_rows(data),
                 money_columns={5, 6, 7, 8, 10, 12, 13, 15, 16, 17, 18}, date_columns={2}, total_columns=set(range(3, 19)))
    sheet.freeze_panes = "C5"
    sheet.sheet_properties.tabColor = NAVY
    sheet.column_dimensions["A"].width = 30
    for cell in sheet["B"][4:]:
        cell.number_format = "yyyy-mm"
    # Scope notes remain outside the filter area, with enough height to read.
    last = sheet.max_row + 3
    notes = [
        "Aplicado, asignado y pendiente: mes de cobranza del período. El pendiente suma saldos por cliente, sin compensar excedentes ajenos.",
        "Depositado completo: fecha real del depósito y empresa pagadora. Puede pagar otros meses o empresas. No se resta del aplicado de este mes.",
        "Aplicado sin período: fecha de aplicación. Se muestra aparte y no se incluye en Aplicado neto ni Pendiente de los períodos.",
        "Estado a la fecha de exportación, no reconstrucción del cierre de cada mes. Los depósitos en borrador no son efectivo confirmado.",
    ]
    for offset, note in enumerate(notes):
        sheet.cell(last + offset, 1, note).font = Font(name="Arial", size=10, italic=True, color=TEXT)


DEPOSIT_HEADERS = (
    "Depósito", "Empresa pagadora", "Fecha depósito", "Referencia", "Cuenta bancaria", "Moneda",
    "Importe original", "Depositado completo USD", "A créditos USD", "A complementarias USD",
    "Efectivo de ajustes USD", "Saldo a favor documentado USD", "Sin clasificar USD",
    "Asignación por revisar USD", "Resultado", "Meses de cobranza destino", "Tasa C$/USD",
)


def _write_deposits(book, data, scope):
    deposits = list({d["name"]: d for d in data.get("cash_deposits") or []}.values())
    _write_table(book.create_sheet("Depósitos"), "Depósitos confirmados por fecha de recepción", scope,
        DEPOSIT_HEADERS,
        [(d["name"], d.get("employer"), _date(d.get("date")), d.get("reference"), d.get("bank_account"),
          d.get("currency"), d.get("original_amount"), d.get("total_usd"), d.get("credits_usd"),
          d.get("other_usd"), d.get("adjustments_usd"), d.get("credit_balance_usd"),
          d.get("unclassified_usd"), d.get("review_usd"), d.get("result"),
          ", ".join(d.get("payroll_months") or []), d.get("fx_rate") or None) for d in deposits],
        money_columns=set(range(7, 15)), date_columns={3}, total_columns=set(range(8, 15)))
    for cell in book["Depósitos"]["Q"][4:]:
        cell.number_format = "0.00000000"
    rows = []
    for d in deposits:
        for target in d.get("destinations") or []:
            # Emit either people or the aggregate, never both as additive rows.
            people = target.get("people") or [{}]
            for person in people:
                rows.append((d["name"], d.get("employer"), _date(d.get("date")),
                    target.get("type"), target.get("employer"), target.get("label"), _month(target.get("month")),
                    person.get("client_number"), person.get("client_name"), person.get("loan_number"),
                    person.get("amount_usd") if person else target.get("amount_usd")))
        # Residual classification is visible even without a client/period link.
        for field, label in (("undetailed_credit_usd", "Saldo a favor sin detalle de cliente"),
                             ("review_usd", "Asignación por revisar"), ("unclassified_usd", "Sin clasificar")):
            if money(d.get(field)):
                rows.append((d["name"], d.get("employer"), _date(d.get("date")), label,
                             None, None, None, None, None, None, d[field]))
    _write_table(book.create_sheet("Distribución depósitos"), "Destino del dinero recibido", scope,
        ("Depósito", "Empresa pagadora", "Fecha depósito", "Tipo de destino", "Empresa destino",
         "Período o partida", "Mes de cobranza", "N.º cliente", "Nombre del cliente", "N.º crédito", "Distribuido USD"),
        rows, money_columns={11}, date_columns={3, 7}, total_columns={11})
    for cell in book["Distribución depósitos"]["G"][4:]:
        cell.number_format = "yyyy-mm"


def _write_guide(book, scope):
    notes = [
        ("Naturaleza del informe", "Fotografía de la conciliación a la fecha de exportación. Descargue nuevamente para actualizar. No sustituye la balanza contable."),
        ("Resumen mensual", "Una fila por empresa y mes. Agrupa todos los períodos del mes, aunque sean quincenales o por fecha exacta. Incluye meses con depósitos y sin períodos."),
        ("Dos bases de fecha", "Aplicaciones y asignaciones: mes de cobranza. Depósitos completos: fecha de recepción. No se calcula una diferencia entre ambos totales mensuales."),
        ("Empresa pagadora", "El depósito completo se cuenta una sola vez para quien paga. Distribución depósitos muestra la empresa destino, aunque sea diferente."),
        ("Pendiente USD", "Suma saldos positivos por cliente/crédito. No compensa una deuda con el excedente de otra partida. Incluye ajustes de redondeo con su signo."),
        ("Aplicado neto USD", "Aplicación después de reducciones confirmadas mediante partidas complementarias. No se debe volver a descontar la misma reducción."),
        ("Complementario USD", "Concepto distinto al pago del crédito, por ejemplo cobranza administrativa. Asignado total puede incluirlo. Asignado a créditos lo excluye."),
        ("Efectivo de ajustes", "Parte del depósito consumida por ajustes. No equivale al signo del ajuste de redondeo que cambia el saldo de la aplicación."),
        ("Saldo a favor documentado", "Clasificación del depósito original. No representa necesariamente lo pendiente de devolver hoy: puede haber gestiones posteriores."),
        ("Importe original", "Moneda indicada en cada depósito. No sumar importes NIO y USD en una misma cifra. Las demás columnas monetarias se concilian en USD."),
        ("Modalidad y cierre", "Histórica y operativa usan iguales conceptos de resultado. Cierre indica si el período está bloqueado, no que sus diferencias estén resueltas."),
        ("N/D", "No disponible o no aplicable, no equivale a cero. En histórico no se inventan cobranza ni deducción."),
        ("Cuenta por cobrar", "Aplicado neto menos depósitos asignados, considerando tolerancias. Las compensaciones confirmadas a aplicaciones ya reducen el aplicado neto; no se descuentan dos veces. Cobranza y deducción son informativas, no generan CxC."),
        ("Detalle de deducción", "Deducido y Cobranza no deducida suman solo filas con evidencia suficiente. Cobranza sin detalle muestra el importe solicitado todavía sin esa evidencia."),
        ("Partidas y excepciones", "Reúne alertas detectadas y casos documentados. Un caso puede explicar una alerta. Sus importes NO se suman para obtener una deuda."),
        ("Partidas complementarias del core", "Incluye importadas no canceladas, aun en borrador, por fecha contable y empresa. Monto USD es el remanente: débito positivo y crédito negativo. Conserva débito/crédito en moneda original y utilizado en USD. Un importe sin signo contable identificable queda sin determinar. Las partidas ya utilizadas no se descuentan otra vez."),
        ("Excepciones sin período", "Se incluyen por año de creación y empresa. Excepciones vinculadas: por mes de cobranza del período. Compromiso es fecha de gestión, no vencimiento del pago."),
        ("Cruces y distribución", "Cruces parte de los períodos seleccionados, incluso con depósitos de otro año. Distribución parte de los depósitos recibidos en el año, incluso hacia períodos de otro año."),
        ("Trazabilidad", "Datos de Períodos, Importaciones contables, Distribuciones de depósito y Excepciones de Credinómina, respetando permisos del usuario."),
        ("Totales filtrados", "La fila superior recalcula los importes visibles al usar los filtros de Excel. No se totaliza Importe original porque puede mezclar monedas."),
    ]
    sheet = book.create_sheet("Guía")
    _write_table(sheet, "Definiciones y alcance", scope, ("Concepto", "Interpretación"), notes)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 108
    for row in sheet.iter_rows(min_row=5):
        row[1].alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row[0].row].height = 44


def _write_summary(sheet, data, employer_label, generated_at):
    periods = data["periods"]
    summary_rows = []
    for period in periods:
        historical_mode = period.get("reconciliation_mode") == "Historica"
        first_date, last_date = _historical_dates(period) if historical_mode else (None, None)
        summary_rows.append((
            period.get("name"), period.get("employer_name"), _month(period.get("month")),
            "Histórica" if historical_mode else "Operativa",
            period.get("historical_scope") if historical_mode else period.get("collection_cycle"),
            first_date, last_date, _date(period.get("remittance_due_date")),
            _control_state(period.get("control_state")),
            NA if historical_mode else _money(period.get("expected_usd")),
            NA if historical_mode else _known_deductions(period),
            _money(period.get("applied_usd")), _money(period.get("complementary_usd")),
            _money(period.get("rounding_adjustment_usd")), _money(period.get("remitted_usd")),
            _period_pending(period),
            NA if historical_mode else _employee_due(period),
            NA if historical_mode else _money(period.get("employer_gap_usd")),
            _period_credit_cash(period),
            len(period.get("exceptions") or []),
            _date(period.get("control_cut_on")),
            period.get("control_cut_note"),
            period.get("remark"), "Cerrado" if period.get("status") == "Cerrado" else "Abierto",
            NA if historical_mode else _missing_detail(period),
        ))
    _write_table(sheet, "Conciliación por período",
                 f"Año de cobranza: {data.get('year') or 'Todos'}    Empresa: {employer_label}",
                 SUMMARY_HEADERS, summary_rows,
                 money_columns=set(range(10, 20)) | {25}, date_columns={6, 7, 8, 21},
                 total_columns=set(range(12, 17)) | {19, 20})
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["E"].width = 27
    sheet.column_dimensions["U"].width = 23
    sheet.column_dimensions["V"].width = 40
    sheet.column_dimensions["W"].width = 65
    sheet.freeze_panes = "C5"
    for cell in sheet["C"][4:]:
        cell.number_format = "yyyy-mm"
    for column in ("P", "R"):
        sheet.column_dimensions[column].width = 25
    sheet.sheet_properties.tabColor = NAVY


def _historical_dates(period):
    scope = period.get("historical_scope")
    if scope == "Fecha exacta":
        day = _date(period.get("historical_application_date"))
        return day, day
    if scope == "Rango de fechas":
        return _date(period.get("historical_start_date")), _date(period.get("historical_end_date"))
    return None, None


def _control_state(value):
    return {
        "historico_excepcion": "Con diferencias",
        "historico_excedente": "Con excedente",
        "historico_conciliado": "Conciliado",
        "historico_parcial": "Parcial",
        "historico_pendiente": "Pendiente",
        "diferencia": "Con diferencias",
        "pendiente_detalle": "Pendiente de detalle",
        "conciliado": "Conciliado",
        "parcial": "Parcial",
        "en_transito": "Pendiente",
        "excedente": "Con excedente",
    }.get(value, value)


def _collection_pending(claim):
    """Same loan-cash basis as aging; complementary cash cannot pay a loan."""
    paid = _collection_cash(claim)
    return money_float(max(
        money(claim.get("applied_usd")) + money(claim.get("rounding_adjustment_usd")) - money(paid), 0,
    ))


def _collection_cash(claim):
    return _sum(
        entry.get("importe_usd") for entry in _json_list(claim.get("remittance_detail"))
        if isinstance(entry, dict) and entry.get("destino") != "Partida complementaria"
    )


def _period_credit_cash(period):
    if period.get("reconciliation_mode") == "Historica":
        return _sum(r.get("historical_remitted_usd") for r in period.get("historical_rows") or [])
    return _sum(_collection_cash(r) for r in period.get("rows") or [])


def _period_pending(period):
    # Sum positive balances per claim, never offset another client's debt with
    # an excess. Reuse the historical balance already computed by reconciliation.
    if period.get("reconciliation_mode") == "Historica":
        values = (max(money(row.get("historical_balance_usd")), 0)
                  for row in period.get("historical_rows") or [])
    else:
        values = (_collection_pending(row) for row in period.get("rows") or [])
    return money_float(sum_money(values))


def _collection_detail(period, claim):
    comments = " | ".join(str(item) for item in (
        claim.get("comments"), claim.get("application_comment"),
        claim.get("inherited_exception_comment"),
    ) if item)
    return (
        period.get("name"), period.get("employer_name"), "Operativa",
        claim.get("client_number"), claim.get("employee_number"), claim.get("client_name"),
        claim.get("national_id"), claim.get("loan_number"), claim.get("installment_number"),
        None, None, None, claim.get("application_reference"),
        _money(claim.get("expected_usd")),
        _money(claim.get("deducted_usd")) if claim.get("collection_shortfall_usd") is not None else NA,
        _money(claim.get("applied_usd")), _money(claim.get("complementary_usd")),
        _money(claim.get("rounding_adjustment_usd")), _collection_cash(claim),
        _money(claim.get("collection_shortfall_usd")) if claim.get("collection_shortfall_usd") is not None else NA,
        _collection_pending(claim),
        " / ".join(str(item) for item in (claim.get("deduction_status"), claim.get("application_status")) if item),
        comments,
        None, None, None, None,
    )


def _historical_detail(period, application):
    return (
        period.get("name"), period.get("employer_name"), "Histórica",
        application.get("client_number"), application.get("employee_number"), application.get("client_name"),
        application.get("national_id"), application.get("loan_number"), application.get("installment_number"),
        _date(application.get("event_date")), application.get("accounting_entry"), application.get("receipt"),
        application.get("reference"), NA, NA, net_application_amount(application),
        0, _sum(e.get("diferencia_usd") for e in _json_list(application.get("historical_detail")) if isinstance(e, dict)),
        _money(application.get("historical_remitted_usd")), NA,
        _money(application.get("historical_balance_usd")),
        application.get("deposit_match_status"), application.get("deposit_match_reason"),
        application.get("parent"), application.get("source_row"), _money(application.get("amount")),
        _money(application.get("application_adjustment_usd")),
    )


def _claim_links(period, claim, *, historical):
    if historical:
        details = _json_list(claim.get("historical_detail"))
        claim_id = claim.get("name")
        application_reference = claim.get("reference")
    else:
        details = _json_list(claim.get("remittance_detail"))
        claim_id = claim.get("row_key")
        application_reference = claim.get("application_reference")
    rows = []
    for detail in details:
        if not isinstance(detail, dict):
            continue
        inherited = detail.get("excepciones_heredadas") or []
        inherited_text = " | ".join(
            str(note.get("comment") or "") for note in inherited if isinstance(note, dict)
        ) if isinstance(inherited, list) else ""
        rows.append((
            period.get("name"), period.get("employer_name"),
            "Histórica" if historical else "Operativa",
            claim.get("client_number"), claim.get("client_name"), claim.get("loan_number"),
            claim_id, application_reference, detail.get("referencia"), detail.get("comprobante"),
            _date(detail.get("fecha")),
            ("Aplicación" if detail.get("destino") == "Aplicación histórica"
             else detail.get("destino") or ("Aplicación" if historical else "Cobranza")),
            _money(detail.get("importe_usd")),
            _money(detail.get("diferencia_usd")) if detail.get("diferencia_usd") is not None else None,
            detail.get("origen"), inherited_text or claim.get("inherited_exception_comment"),
        ))
    return rows


def _issue(issue_type, period, client, loan, reference, amount, status, *,
           employer=None, cause=None, owner=None, action=None, commitment=None,
           evidence=None, external_reference=None, detail=None, resolution=None,
           event_date=None):
    return (
        issue_type, _date(event_date), period.get("name") if period else None,
        period.get("employer_name") if period else employer,
        client, loan, reference, amount, status, cause, owner, action,
        _date(commitment), evidence, external_reference, detail, resolution,
        _date(period.get("remittance_due_date")) if period else None,
    )


def _action_rows(actions, exceptions, period_by_name):
    exception_periods = {item.get("name"): item.get("period") for item in exceptions}
    exception_employers = {item.get("name"): item.get("employer") for item in exceptions}
    result = []
    for action in actions:
        period = period_by_name.get(exception_periods.get(action.get("parent")))
        result.append((
            action.get("parent"), period.get("name") if period else None,
            period.get("employer_name") if period else exception_employers.get(action.get("parent")),
            _date(action.get("action_at")), action.get("action_by"),
            action.get("action_type"), action.get("details"),
            action.get("evidence_file"), action.get("external_reference"),
        ))
    return result


def _commitment_delay(issue_row, as_of):
    # This is a follow-up deadline, not the employer's contractual payment date.
    if issue_row[0] != "Excepción" or issue_row[8] not in {"Abierta", "En revision"}:
        return None
    commitment = _date(issue_row[12])
    if not isinstance(commitment, date):
        return None
    day = commitment.date() if isinstance(commitment, datetime) else commitment
    return max((as_of - day).days, 0)


def _json_list(value):
    if isinstance(value, list):
        return value
    if not value:
        return []
    try:
        parsed = json.loads(value)
    except (TypeError, ValueError):
        return []
    return parsed if isinstance(parsed, list) else []


def _money(value):
    return money_float(value)


def _date(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return value
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value)) if len(str(value)) > 10 else date.fromisoformat(str(value))
    except ValueError:
        return value


def _safe_text(value):
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def _base_sheet(sheet, title):
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Arial", size=15, bold=True, color=NAVY)
    sheet.row_dimensions[1].height = 27


def _write_table(sheet, title, scope, headers, rows, *, money_columns=(), date_columns=(), total_columns=()):
    _base_sheet(sheet, title)
    sheet["A2"] = _safe_text(scope)
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="5A6673")
    _table(sheet, headers, rows, header_row=4,
           money_columns=money_columns, date_columns=date_columns)
    sheet.freeze_panes = "E5" if len(headers) > 12 else "A5"
    if total_columns:
        sheet["A3"] = "Totales filtrados"
        for column in total_columns:
            letter = get_column_letter(column)
            value = f"=SUBTOTAL(109,{letter}5:{letter}{4 + len(rows)})" if rows else 0
            if rows and not any(isinstance(row[column - 1], (int, float)) for row in rows):
                value = NA
            cell = sheet.cell(3, column, value)
            cell.number_format = MONEY_FORMAT if column in money_columns else "#,##0"
            cell.font = Font(name="Arial", size=10, bold=True, color=NAVY)
            cell.alignment = Alignment(horizontal="right")


def _table(sheet, headers, rows, *, header_row, money_columns=(), date_columns=()):
    if header_row + len(rows) > MAX_EXCEL_ROWS:
        raise ValueError("El informe supera el máximo de filas permitido por Excel.")
    for column, heading in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, heading)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        width = 17
        if any(text in heading.lower() for text in ("nombre", "detalle", "comentario", "acción", "resolución", "observaciones", "motivo")):
            width = 35
        elif "empresa" in heading.lower() or "estado" in heading.lower() or "resultado" in heading.lower():
            width = 28
        elif any(text in heading.lower() for text in ("referencia", "evidencia", "asiento", "comprobante")):
            width = 25
        elif len(heading) > 24:
            width = 26
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[header_row].height = 45
    for row_number, values in enumerate(rows, start=header_row + 1):
        line_count = 1
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, _safe_text(value))
            cell.font = Font(name="Arial", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="center")
            if isinstance(value, str):
                cell.number_format = "@"
                width = sheet.column_dimensions[get_column_letter(column)].width
                line_count = max(line_count, sum(max(math.ceil(len(part) / max(width - 2, 1)), 1)
                                                 for part in value.split("\n")))
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if column in money_columns and isinstance(value, (int, float)):
                cell.number_format = MONEY_FORMAT
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif column in money_columns and value == NA:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif column in date_columns and isinstance(value, (date, datetime)):
                cell.number_format = DATE_FORMAT if isinstance(value, date) and not isinstance(value, datetime) else "dd/mm/yyyy hh:mm"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7F9FC")
        sheet.row_dimensions[row_number].height = min(409, max(20, line_count * 14 + 4))
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(header_row, header_row + len(rows))}"
