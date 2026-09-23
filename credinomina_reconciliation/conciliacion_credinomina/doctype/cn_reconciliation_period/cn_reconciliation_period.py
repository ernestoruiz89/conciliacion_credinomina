from __future__ import annotations

import io
import json
from collections import defaultdict

import frappe
from frappe import _
from frappe.model.document import Document
from frappe.utils import add_days, add_months, flt, getdate, now_datetime

from credinomina_reconciliation.cadence import (
    MONTHLY,
    cycle_code,
    cycle_cutoff,
    cycle_for_frequency,
    cycles_conflict,
)
from credinomina_reconciliation.client_registry import ClientIndex, load_client_index, names_for_claim
from credinomina_reconciliation.deduction_recognition import recognition_reason
from credinomina_reconciliation.employer_naming import employer_alias_index
from credinomina_reconciliation.historical import (
    HISTORICAL_MONTHLY,
    OPERATIVE_START,
    historical_scope_interval,
    historical_scopes_conflict,
    is_historical_date,
)
from credinomina_reconciliation.parsers import (
    SOURCE_ACCOUNTING,
    SOURCE_DEPOSITS,
    SourceFileError,
    clean_text,
    file_sha256,
    parse_collection_file,
    source_key,
)
from credinomina_reconciliation.period_lock import current_period_write_action, period_write_action
from credinomina_reconciliation.reconciliation import (
    classify_deduction,
    converted_amount,
    deposit_pair_result,
    match_collection_record,
    narrow_deposit_candidates_by_date,
    remittance_fx_basis,
)
from credinomina_reconciliation.rounding import CASH_EPSILON


EXCEPTION_STATES = {
    "Deduccion parcial",
    "No deducido",
    "Deduccion en exceso",
    "Moneda no coincide",
    "Importes inconsistentes",
    "Importe invalido",
}


class CNReconciliationPeriod(Document):
    def on_trash(self):
        if self.status == "Cerrado":
            frappe.throw(_("Reabra el período antes de eliminarlo."))

    def validate(self):
        self._validate_closed_transition()
        if self.payroll_month:
            self.payroll_month = getdate(self.payroll_month).replace(day=1)
        self._validate_mode()
        self._set_due_date()
        self._validate_unique_period()
        self.recalculate_totals()

    def _validate_closed_transition(self):
        previous = self.get_doc_before_save()
        action = current_period_write_action()
        if previous and previous.status == "Cerrado":
            if action == "reopen" and self.status != "Cerrado":
                return
            if action == "reconcile" and self.status == "Cerrado":
                return
            frappe.throw(_("El período está cerrado. Use Reabrir período antes de modificarlo."))
        if self.status == "Cerrado" and action != "close":
            frappe.throw(_("Use Cerrar período para completar el cierre y dejarlo bloqueado."))

    def _validate_mode(self):
        if not self.payroll_month:
            return
        previous = self.get_doc_before_save()
        linked_applications = bool(
            previous and previous.reconciliation_mode == "Historica"
            and frappe.db.exists(
                "CN Source Row", {"historical_period": self.name, "event_type": "Aplicacion"}
            )
        )
        if linked_applications and self.reconciliation_mode != "Historica":
            frappe.throw(_("Reasigne primero las aplicaciones antes de cambiar la modalidad histórica."))
        if self.reconciliation_mode == "Historica":
            self.collection_cycle = ""
            if not is_historical_date(self.payroll_month):
                frappe.throw(_("El período histórico debe estar entre abril de 2025 y agosto de 2026."))
            self.historical_scope = self.historical_scope or HISTORICAL_MONTHLY
            try:
                historical_scope_interval(
                    self.historical_scope, self.historical_application_date,
                    self.historical_start_date, self.historical_end_date,
                )
            except ValueError as exc:
                frappe.throw(_(str(exc)))
            if self.collection_rows or self.collection_file or self.employer_response_file:
                frappe.throw(_("Un período histórico no admite cobranza ni detalle de deducción."))
            if previous and previous.reconciliation_mode == "Historica":
                before = (
                    previous.employer, str(previous.payroll_month or "")[:10],
                    previous.historical_scope or HISTORICAL_MONTHLY,
                    str(previous.historical_application_date or "")[:10],
                    str(previous.historical_start_date or "")[:10],
                    str(previous.historical_end_date or "")[:10],
                )
                after = (
                    self.employer, str(self.payroll_month or "")[:10], self.historical_scope,
                    str(self.historical_application_date or "")[:10],
                    str(self.historical_start_date or "")[:10],
                    str(self.historical_end_date or "")[:10],
                )
                if before != after and linked_applications:
                    frappe.throw(_(
                        "Reasigne primero las aplicaciones antes de cambiar el corte histórico."
                    ))
        else:
            self.historical_scope = ""
            self.historical_application_date = None
            self.historical_start_date = None
            self.historical_end_date = None
            if self.is_new() and getdate(self.payroll_month) < OPERATIVE_START:
                frappe.throw(_("Para abril de 2025 a agosto de 2026 seleccione la modalidad Histórica."))
            previous = self.get_doc_before_save()
            previous_cycle = (previous.collection_cycle or MONTHLY) if previous else None
            if (
                previous and previous_cycle != (self.collection_cycle or MONTHLY)
                and (previous.collection_rows or previous.status != "Borrador")
            ):
                frappe.throw(_("No cambie el ciclo de una cobranza ya cargada."))
            if self.is_new() or (previous and previous_cycle != (self.collection_cycle or MONTHLY)):
                frequency = frappe.db.get_value(
                    "CN Employer", self.employer, "payroll_frequency"
                ) or MONTHLY
                try:
                    self.collection_cycle = cycle_for_frequency(
                        frequency, self.collection_cycle
                    )
                except ValueError as exc:
                    frappe.throw(_(str(exc)))
            elif not self.collection_cycle:
                self.collection_cycle = MONTHLY  # Existing monthly periods before this field.

    def on_update(self):
        if self.flags.get("skip_comment_reconciliation"):
            return
        previous = self.get_doc_before_save()
        if not previous:
            return
        previous_rows = {row.name: row for row in previous.collection_rows or []}
        changed = any(
            row.name in previous_rows
            and any(
                (row.get(field) or "") != (previous_rows[row.name].get(field) or "")
                for field in ("first_exception_comment", "application_comment")
            )
            for row in self.collection_rows or []
        )
        if not changed or not frappe.db.exists(
            "CN Source Import",
            {"status": ["in", ["Importado", "Importado con excepciones"]]},
        ):
            return
        from credinomina_reconciliation.conciliacion_credinomina.doctype.cn_source_import.cn_source_import import (
            reconcile_all_sources,
        )

        reconcile_all_sources()

    def _set_due_date(self):
        if self.reconciliation_mode == "Historica":
            self.cutoff_date = None
            self.remittance_due_date = None
            return
        if not self.payroll_month or not self.employer:
            return
        self.cutoff_date = cycle_cutoff(getdate(self.payroll_month), self.collection_cycle)
        if self.collection_cycle != MONTHLY:
            if self.remittance_due_date and getdate(self.remittance_due_date) < self.cutoff_date:
                frappe.throw(_("El vencimiento de la remesa no puede preceder el cierre de la quincena."))
            return
        grace_days = frappe.db.get_value("CN Employer", self.employer, "grace_days") or 10
        first_next_month = getdate(add_months(getdate(self.payroll_month).replace(day=1), 1))
        self.remittance_due_date = add_days(first_next_month, max(int(grace_days), 1) - 1)

    def _validate_unique_period(self):
        if not self.employer or not self.payroll_month:
            return
        candidates = frappe.get_all(
            self.doctype,
            filters={
                "employer": self.employer,
                "payroll_month": getdate(self.payroll_month).replace(day=1),
            },
            fields=[
                "name", "collection_cycle", "reconciliation_mode", "historical_scope",
                "historical_application_date", "historical_start_date", "historical_end_date",
            ],
            limit_page_length=1000,
        )
        for candidate in candidates:
            if candidate.name == self.name:
                continue
            if candidate.reconciliation_mode != self.reconciliation_mode:
                frappe.throw(_("Ya existe otra modalidad para esta empresa y mes."))
            if self.reconciliation_mode == "Historica":
                existing = historical_scope_interval(
                    candidate.historical_scope or HISTORICAL_MONTHLY,
                    candidate.historical_application_date,
                    candidate.historical_start_date, candidate.historical_end_date,
                )
                proposed = historical_scope_interval(
                    self.historical_scope, self.historical_application_date,
                    self.historical_start_date, self.historical_end_date,
                )
                conflict = historical_scopes_conflict(existing, proposed)
            else:
                conflict = cycles_conflict(
                    candidate.collection_cycle or MONTHLY, self.collection_cycle or MONTHLY
                )
            if conflict:
                frappe.throw(
                    _("Ya existe el período {0} para esta empresa y corte.").format(candidate.name)
                )

    def recalculate_totals(self):
        if self.reconciliation_mode == "Historica":
            return  # Rebuilt from linked core applications, never from payroll rows.
        rows = list(self.collection_rows or [])
        for fieldname in (
            "expected_usd",
            "expected_nio",
            "deducted_usd",
            "deducted_nio",
            "applied_usd",
            "applied_nio",
            "complementary_usd",
            "remitted_usd",
            "remitted_nio",
            "fx_variance_usd",
            "rounding_adjustment_usd",
        ):
            self.set(fieldname, sum(flt(row.get(fieldname)) for row in rows))
        if self.name and not self.is_new():
            self.exception_count = frappe.db.count(
                "CN Reconciliation Exception",
                {"period": self.name, "status": ["in", ["Abierta", "En revision"]]},
            )
        else:
            self.exception_count = 0


def _attached_file(document, file_url):
    if not file_url:
        frappe.throw(_("Adjunte el archivo antes de ejecutar la importacion."))
    file_doc = frappe.get_doc("File", {"file_url": file_url})
    if (
        file_doc.attached_to_doctype != document.doctype
        or file_doc.attached_to_name != document.name
    ):
        frappe.throw(_("El archivo debe estar adjunto a este periodo."))
    content = file_doc.get_content()
    if isinstance(content, str):
        content = content.encode("utf-8")
    return file_doc, content


def _assert_editable(period):
    if period.reconciliation_mode == "Historica":
        frappe.throw(_("En un período histórico importe aplicaciones y depósitos; no la cobranza."))
    if period.status == "Cerrado":
        frappe.throw(_("Un periodo cerrado no se puede volver a importar."))


def _recognition_pairs():
    """Find legacy matched deposits and supported deposits entered in the app."""
    imports = frappe.get_all(
        "CN Source Import",
        filters={"status": ["in", ["Importado", "Importado con excepciones"]]},
        fields=["name", "source_type"],
        limit_page_length=100000,
    )
    source_type = {item.name: item.source_type for item in imports}
    relevant = [
        name for name, kind in source_type.items()
        if kind in {SOURCE_ACCOUNTING, SOURCE_DEPOSITS}
    ]
    deposits = frappe.get_all(
        "CN Source Row",
        filters={"parent": ["in", relevant], "event_type": "Deposito", "effective": 1},
        fields=[
            "name", "parent", "reference", "voucher", "event_date", "employer_text",
            "currency", "amount", "equivalent_currency", "equivalent_amount",
            "fx_rate", "fx_basis", "manual_fx_rate", "manual_fx_evidence",
            "allocated_usd", "justified_surplus_usd",
        ],
        limit_page_length=100000,
    ) if relevant else []
    accounting = [row for row in deposits if source_type.get(row.parent) == SOURCE_ACCOUNTING]
    bank = [row for row in deposits if source_type.get(row.parent) == SOURCE_DEPOSITS]
    accounting_by_reference = defaultdict(list)
    bank_by_reference = defaultdict(list)
    for row in accounting:
        accounting_by_reference[clean_text(row.reference)].append(row)
    for row in bank:
        bank_by_reference[clean_text(row.reference)].append(row)
    paired = []
    for account in accounting:
        matches = narrow_deposit_candidates_by_date(
            account, [
                item for item in bank_by_reference[clean_text(account.reference)]
                if deposit_pair_result(account, item)[0]
            ]
        )
        if len(matches) != 1:
            continue
        reverse = narrow_deposit_candidates_by_date(
            matches[0],
            [
                item for item in accounting_by_reference[clean_text(matches[0].reference)]
                if deposit_pair_result(item, matches[0])[0]
            ],
        )
        if len(reverse) == 1 and reverse[0].name == account.name:
            paired.append((account, matches[0]))
    for item in frappe.get_all(
        "CN Remittance Allocation",
        filters={"docstatus": 1},
        fields=[
            "name", "employer", "deposit_reference", "deposit_voucher",
            "deposit_date", "deposit_currency", "deposit_amount", "amount_usd",
            "fx_rate", "notes", "support_file", "allocated_usd", "unallocated_usd",
        ],
        limit_page_length=100000,
    ):
        fx_basis = remittance_fx_basis(item) if item.deposit_currency == "NIO" else "Moneda original USD"
        deposit = frappe._dict(
            name=item.name, reference=item.deposit_reference,
            voucher=item.deposit_voucher, event_date=item.deposit_date,
            employer_text=item.employer, currency=item.deposit_currency,
            amount=item.deposit_amount, equivalent_currency="USD",
            equivalent_amount=item.amount_usd,
            fx_basis=fx_basis,
            manual_fx_rate=item.fx_rate, manual_fx_evidence=fx_basis,
            allocated_usd=item.allocated_usd, justified_surplus_usd=0,
        )
        paired = [
            pair for pair in paired
            if not deposit_pair_result(pair[0], deposit)[0]
        ]
        paired.append((deposit, deposit))
    return paired


def _recognition_candidates(period):
    _assert_editable(period)
    if not period.collection_rows or period.employer_response_file:
        return []
    if period.deduction_basis:
        return []
    used = {
        item.deduction_recognition_deposit
        for item in frappe.get_all(
            "CN Reconciliation Period",
            fields=["name", "deduction_recognition_deposit"],
            limit_page_length=100000,
        )
        if item.name != period.name and item.deduction_recognition_deposit
    }
    known_employers = frappe.get_all(
        "CN Employer", fields=["name", "employer_name", "employer_code"],
        limit_page_length=100000,
    )
    employer_by_label, ambiguous_labels = employer_alias_index(known_employers)
    rows = [row.as_dict() for row in period.collection_rows]
    candidates = []
    for account, bank in _recognition_pairs():
        if account.name in used:
            continue
        source_labels = {
            clean_text(item.employer_text).casefold()
            for item in (account, bank) if clean_text(item.employer_text)
        }
        if source_labels & ambiguous_labels:
            continue
        labels = {
            employer_by_label.get(label) for label in source_labels
        }
        if any(label and label != period.employer for label in labels):
            continue
        if not account.event_date and not bank.event_date:
            continue
        event_date = getdate(account.event_date or bank.event_date)
        if period.cutoff_date and event_date < getdate(period.cutoff_date):
            continue
        if recognition_reason(rows, account, bank):
            continue
        usd = converted_amount(account, "USD")
        if usd is None:
            usd = converted_amount(bank, "USD")
        candidates.append({
            "source_row_id": account.name,
            "bank_row_id": bank.name,
            "reference": account.reference,
            "voucher": account.voucher,
            "event_date": str(event_date),
            "currency": account.currency,
            "amount": flt(account.amount),
            "amount_usd": round(usd, 4),
            "account": account,
            "bank": bank,
        })
    return candidates


@frappe.whitelist()
def get_recognizable_deposits(period_name: str):
    period = frappe.get_doc("CN Reconciliation Period", period_name)
    period.check_permission("read")
    if not (frappe.has_permission("CN Source Import", "read") or frappe.has_permission("CN Remittance Allocation", "read")):
        frappe.throw(_("No tiene permiso para consultar depósitos."))
    return [
        {key: value for key, value in candidate.items() if key not in {"account", "bank"}}
        for candidate in _recognition_candidates(period)
    ]


@frappe.whitelist(methods=["POST"])
def recognize_collection_from_deposit(period_name: str, source_row_id: str, justification: str):
    # Serialize competing recognitions of the same cash row, then recheck all
    # conditions inside the current transaction.
    lock_table = "tabCN Remittance Allocation" if frappe.db.exists(
        "CN Remittance Allocation", source_row_id
    ) else "tabCN Source Row"
    frappe.db.sql(
        f"SELECT name FROM `{lock_table}` WHERE name = %s FOR UPDATE",
        (source_row_id,),
    )
    period = frappe.get_doc("CN Reconciliation Period", period_name)
    period.check_permission("write")
    if not (frappe.has_permission("CN Source Import", "read") or frappe.has_permission("CN Remittance Allocation", "read")):
        frappe.throw(_("No tiene permiso para consultar depósitos."))
    justification = clean_text(justification)
    if len(justification) < 12:
        frappe.throw(_("Explique por qué el depósito permite reconocer la cobranza completa (mínimo 12 caracteres)."))
    matches = [
        candidate for candidate in _recognition_candidates(period)
        if candidate["source_row_id"] == source_row_id
    ]
    if len(matches) != 1:
        frappe.throw(_("El depósito no coincide de forma única con esta cobranza o ya fue usado."))
    chosen = matches[0]
    evidence_date = getdate(chosen["event_date"])
    if period.deduction_evidence_date:
        period.notes = _append_note(
            period.notes,
            _("Fecha de evidencia de planilla previamente indicada ({0}); no usada para la inferencia por depósito.").format(
                period.deduction_evidence_date
            ),
        )
    period.deduction_evidence_date = None
    for row in period.collection_rows:
        row.deducted_usd = flt(row.expected_usd)
        row.deducted_nio = flt(row.expected_nio)
        row.deduction_currency = ""  # The payment currency does not prove payroll currency.
        row.deduction_status = "Inferida por depósito"
        row.deduction_evidence_date = evidence_date
        row.deduction_match_note = _(
            "Deducción inferida, no confirmada por planilla. Depósito {0}; respaldo: {1}."
        ).format(chosen["reference"], chosen["source_row_id"])
    period.deduction_basis = "Depósito coincidente"
    period.deduction_recognition_deposit = chosen["source_row_id"]
    period.deduction_recognition_reference = chosen["reference"]
    period.deduction_recognition_note = justification
    period.deduction_recognition_on = now_datetime()
    period.deduction_recognition_by = frappe.session.user
    period.status = "Deduccion conciliada"
    period.notes = _append_note(
        period.notes,
        _("Cobranza reconocida provisionalmente por depósito {0} ({1}) el {2} por {3}: {4}").format(
            chosen["reference"], chosen["source_row_id"],
            period.deduction_recognition_on, frappe.session.user, justification,
        ),
    )
    period.flags.skip_comment_reconciliation = True
    period.save()
    from credinomina_reconciliation.conciliacion_credinomina.doctype.cn_source_import.cn_source_import import (
        reconcile_all_sources,
    )

    summary = reconcile_all_sources()
    return {"period": period.name, "source_reconciliation": summary}


@frappe.whitelist(methods=["POST"])
def revert_deposit_recognition(period_name: str):
    period = frappe.get_doc("CN Reconciliation Period", period_name)
    period.check_permission("write")
    _assert_editable(period)
    if period.deduction_basis != "Depósito coincidente":
        frappe.throw(_("Este período no usa un depósito como detalle de deducción."))
    old_reference = period.deduction_recognition_reference
    _clear_deposit_recognition(period)
    for row in period.collection_rows:
        row.deducted_usd = 0
        row.deducted_nio = 0
        row.deduction_currency = ""
        row.deduction_status = "Pendiente de detalle"
        row.deduction_evidence_date = None
        row.deduction_match_note = ""
    period.status = "Cobranza cargada"
    period.notes = _append_note(
        period.notes,
        _("Reconocimiento provisional por depósito {0} revertido el {1} por {2}.").format(
            old_reference, now_datetime(), frappe.session.user,
        ),
    )
    period.flags.skip_comment_reconciliation = True
    period.save()
    from credinomina_reconciliation.conciliacion_credinomina.doctype.cn_source_import.cn_source_import import (
        reconcile_all_sources,
    )

    summary = reconcile_all_sources()
    return {"period": period.name, "source_reconciliation": summary}


def _clear_deposit_recognition(period):
    period.deduction_recognition_deposit = ""
    period.deduction_recognition_reference = ""
    period.deduction_recognition_note = ""
    period.deduction_recognition_on = None
    period.deduction_recognition_by = ""
    period.deduction_basis = ""


@frappe.whitelist(methods=["POST"])
def import_collection(period_name: str):
    period = frappe.get_doc("CN Reconciliation Period", period_name)
    period.check_permission("write")
    _assert_editable(period)
    if period.deduction_basis == "Depósito coincidente":
        frappe.throw(_("Revierta primero el reconocimiento por depósito antes de reemplazar la cobranza."))
    file_doc, content = _attached_file(period, period.collection_file)
    try:
        parsed = parse_collection_file(file_doc.file_name, content, require_name=True)
    except SourceFileError as exc:
        frappe.throw(str(exc), title=_("Archivo de cobranza invalido"))

    _supersede_open_period_exceptions(period.name)
    period.set("collection_rows", [])
    clients = ClientIndex()
    seen = set()
    for record in parsed:
        if not record.get("loan_number"):
            frappe.throw(_("La fila {0} de cobranza no tiene número de crédito.").format(record["source_row"]))
        client = clients.ensure_from_collection(record, period.employer)
        row_key = record.get("row_key") or source_key(
            period.employer,
            getdate(period.payroll_month).replace(day=1),
            period.collection_cycle,
            record.get("client_number") or record.get("national_id") or record.get("employee_number") or record.get("client_name"),
            record.get("loan_number"),
            record.get("installment_number"),
        )[:24]
        if row_key in seen:
            frappe.throw(
                _("La fila {0} duplica cliente, credito y cuota.").format(
                    record["source_row"]
                )
            )
        seen.add(row_key)
        period.append(
            "collection_rows",
            {
                **record,
                "client": client,
                "row_key": row_key,
                "deduction_status": "Pendiente de detalle",
                "application_status": "Pendiente",
            },
        )
    period.status = "Cobranza cargada"
    period.notes = _append_note(
        period.notes,
        _("Cobranza importada: {0} filas; SHA-256 {1}.").format(
            len(parsed), file_sha256(content)
        ),
    )
    period.save()
    return {"period": period.name, "rows": len(parsed), "status": period.status}


@frappe.whitelist(methods=["POST"])
def import_employer_response(period_name: str):
    period = frappe.get_doc("CN Reconciliation Period", period_name)
    period.check_permission("write")
    _assert_editable(period)
    if not period.collection_rows:
        frappe.throw(_("Importe primero el archivo de cobranza."))
    if not period.deduction_evidence_date:
        frappe.throw(
            _(
                "Indique la fecha de evidencia de deduccion antes de importar el detalle de la empresa."
            )
        )
    file_doc, content = _attached_file(period, period.employer_response_file)
    try:
        responses = parse_collection_file(
            file_doc.file_name, content,
            require_deduction=True, require_name=True, keep_zero_rows=True,
        )
    except SourceFileError as exc:
        frappe.throw(str(exc), title=_("Detalle de empresa invalido"))

    if period.deduction_basis == "Depósito coincidente":
        period.notes = _append_note(
            period.notes,
            _("El detalle real de la empresa reemplazó el reconocimiento provisional por depósito {0}.").format(
                period.deduction_recognition_reference
            ),
        )
        _clear_deposit_recognition(period)
        for row in period.collection_rows:
            if row.deduction_status == "Inferida por depósito":
                row.deducted_usd = 0
                row.deducted_nio = 0
                row.deduction_currency = ""
                row.deduction_status = "Pendiente de detalle"
                row.deduction_evidence_date = None
                row.deduction_match_note = ""
    period.deduction_basis = "Detalle de empresa"
    _supersede_open_period_exceptions(period.name)

    rows_by_name = {row.name: row for row in period.collection_rows}
    client_catalog = load_client_index()
    by_client = {client["name"]: client for client in client_catalog}
    candidates = []
    for row in period.collection_rows:
        candidate = {**row.as_dict(), "name": row.name}
        linked = by_client.get(row.get("client"))
        candidate["client_aliases"] = (
            [linked["client_name"], *(linked.get("client_aliases") or ())]
            if linked else names_for_claim(candidate, client_catalog, period.employer)
        )
        candidates.append(candidate)
    matched_names = set()
    unmatched = 0
    for response in responses:
        match, reason = match_collection_record(response, candidates)
        if not match or match["name"] in matched_names:
            unmatched += 1
            _create_exception(
                period,
                exception_type="Detalle de empresa sin coincidencia",
                source_row=response.get("source_row"),
                client_number=response.get("client_number"),
                loan_number=response.get("loan_number"),
                amount_usd=response.get("deducted_usd"),
                amount_nio=response.get("deducted_nio"),
                description=(
                    "Fila duplicada en la respuesta"
                    if match and match["name"] in matched_names
                    else reason
                ),
            )
            continue
        row = rows_by_name[match["name"]]
        matched_names.add(row.name)
        deducted_usd, deducted_nio, conversion_note = _deduction_equivalents(
            row, response
        )
        row.deducted_usd = deducted_usd
        row.deducted_nio = deducted_nio
        if flt(response.get("deducted_usd")) and flt(response.get("deducted_nio")):
            row.deduction_currency = "Ambas"
        elif flt(response.get("deducted_nio")):
            row.deduction_currency = "NIO"
        elif flt(response.get("deducted_usd")):
            row.deduction_currency = "USD"
        else:
            row.deduction_currency = ""
        row.deduction_evidence_date = period.deduction_evidence_date
        row.deduction_match_note = " ".join(
            part
            for part in (
                _("Coincidencia exacta por {0}.").format(reason),
                conversion_note,
            )
            if part
        )
        row.deduction_status = classify_deduction(
            expected_usd=row.expected_usd,
            expected_nio=row.expected_nio,
            deducted_usd=row.deducted_usd,
            deducted_nio=row.deducted_nio,
        )
        if response.get("application_reference"):
            row.application_reference = response["application_reference"]
        if response.get("application_comment"):
            row.application_comment = response["application_comment"]
        if row.deduction_status in EXCEPTION_STATES:
            _create_exception(
                period,
                exception_type=row.deduction_status,
                source_row=response.get("source_row"),
                client_number=row.client_number,
                loan_number=row.loan_number,
                amount_usd=row.deducted_usd,
                amount_nio=row.deducted_nio,
                description=row.application_comment or row.comments,
                collection_row_id=row.name,
            )

    missing = 0
    for row in period.collection_rows:
        if row.name not in matched_names:
            missing += 1
            row.deduction_status = "Pendiente de detalle"
            _create_exception(
                period,
                exception_type="Pendiente de detalle de empresa",
                client_number=row.client_number,
                loan_number=row.loan_number,
                amount_usd=row.expected_usd,
                amount_nio=row.expected_nio,
                description=_(
                    "La respuesta de la empresa no incluyo esta fila de cobranza."
                ),
                collection_row_id=row.name,
            )

    period.status = (
        "Deduccion conciliada"
        if unmatched == 0 and missing == 0
        else "Detalle empresa cargado"
    )
    period.notes = _append_note(
        period.notes,
        _(
            "Detalle de empresa importado: {0} coincidencias, {1} filas sin enlace y {2} cuotas ausentes."
        ).format(
            len(matched_names), unmatched, missing
        ),
    )
    period.flags.skip_comment_reconciliation = True
    period.save()
    source_summary = None
    if frappe.db.exists(
        "CN Source Import",
        {"status": ["in", ["Importado", "Importado con excepciones"]]},
    ):
        from credinomina_reconciliation.conciliacion_credinomina.doctype.cn_source_import.cn_source_import import (
            reconcile_all_sources,
        )

        source_summary = reconcile_all_sources()
        period.reload()
    return {
        "period": period.name,
        "matched": len(matched_names),
        "unmatched": unmatched,
        "missing": missing,
        "status": period.status,
        "source_reconciliation": source_summary,
    }


def _pending_registered_targets(target_filters):
    submitted = frappe.get_all(
        "CN Remittance Allocation", filters={"docstatus": 1}, pluck="name",
        limit_page_length=100000,
    )
    if not submitted:
        return False
    return bool(frappe.db.count(
        "CN Remittance Target",
        {**target_filters, "parent": ["in", submitted], "result": ["!=", "Aplicada"]},
    ))


@frappe.whitelist(methods=["POST"])
def close_period(period_name: str):
    period = frappe.get_doc("CN Reconciliation Period", period_name)
    period.check_permission("write")
    if period.status == "Cerrado":
        frappe.throw(_("Este período ya está cerrado."))
    if frappe.db.count(
        "CN Remittance Allocation",
        {
            "docstatus": 1, "detail_period": period.name,
            "detail_status": ["in", [
                "Revisar filas", "Detalle supera depósito", "Detalle pendiente",
                "Importar detalle actualizado",
            ]],
        },
    ):
        frappe.throw(_("Hay detalles de depósito por cliente pendientes de revisión para este período."))
    if period.reconciliation_mode == "Historica":
        if period.status != "Historico conciliado":
            frappe.throw(_("Todas las aplicaciones históricas deben estar cubiertas por depósitos antes del cierre."))
        application_ids = frappe.get_all(
            "CN Source Row",
            filters={
                "historical_period": period.name,
                "event_type": "Aplicacion", "effective": 1,
            },
            pluck="name", limit_page_length=100000,
        )
        if application_ids and _pending_registered_targets(
            {"historical_application": ["in", application_ids]}
        ):
            frappe.throw(_("Hay destinos de depósitos históricos pendientes o inválidos."))
        if frappe.db.count(
            "CN Deposit Surplus",
            {"docstatus": 1, "period": period.name, "result": ["!=", "Saldo a favor documentado"]},
        ):
            frappe.throw(_("Hay excedentes históricos pendientes de validar."))
        _mark_period_closed(period)
        return {"period": period.name, "status": period.status}
    period.recalculate_totals()
    if _pending_registered_targets({"period": period.name}):
        frappe.throw(_("Hay destinos de depósitos pendientes o inválidos para este período."))
    if frappe.db.count(
        "CN Deposit Surplus",
        {
            "docstatus": 1, "period": period.name,
            "result": ["!=", "Saldo a favor documentado"],
        },
    ):
        frappe.throw(_("Hay excedentes de deposito pendientes de validar."))
    references = {
        row.application_reference
        for row in period.collection_rows if row.application_reference
    }
    references.update(
        frappe.get_all(
            "CN Source Row",
            filters={"event_type": "Aplicacion", "collection_period": period.name, "effective": 1},
            pluck="reference",
        )
    )
    target_deposits = frappe.get_all(
        "CN Remittance Target", filters={"period": period.name},
        pluck="parent", limit_page_length=100000,
    )
    if target_deposits:
        references.update(frappe.get_all(
            "CN Remittance Allocation",
            filters={"name": ["in", target_deposits], "docstatus": 1},
            pluck="deposit_reference", limit_page_length=100000,
        ))
    references.discard("")
    references.discard(None)
    if references:
        import_names = frappe.get_all(
            "CN Source Import",
            filters={"source_type": SOURCE_ACCOUNTING},
            pluck="name",
        )
        if import_names and frappe.db.count(
            "CN Source Row",
            {
                "parent": ["in", import_names],
                "event_type": "Deposito",
                "reference": ["in", list(references)],
                "unclassified_usd": [">", CASH_EPSILON],
            },
        ):
            frappe.throw(
                _("Hay depositos relacionados con este periodo sin distribuir ni justificar.")
            )
    for deposit in frappe.get_all(
        "CN Remittance Allocation",
        filters={"docstatus": 1, "unclassified_usd": [">", CASH_EPSILON]},
        fields=["name", "allocation_detail"],
        limit_page_length=100000,
    ):
        if any(
            entry.get("periodo") == period.name
            for entry in json.loads(deposit.allocation_detail or "[]")
        ):
            frappe.throw(_("El depósito {0} tiene un saldo sin clasificar relacionado con este período.").format(deposit.name))
    open_exceptions = frappe.db.count(
        "CN Reconciliation Exception",
        {"period": period.name, "status": ["in", ["Abierta", "En revision"]]},
    )
    if open_exceptions:
        frappe.throw(_("Resuelva las excepciones antes de cerrar el periodo."))
    if any(abs(flt(row.fx_variance_usd)) > 0.01 for row in period.collection_rows):
        frappe.throw(
            _("Hay diferencias cambiarias pendientes de revisar y aplicar en el core.")
        )
    if any(row.application_status != "Aplicado y remitido" for row in period.collection_rows):
        frappe.throw(_("Todas las filas deben estar aplicadas y remitidas antes del cierre."))
    _mark_period_closed(period)
    return {"period": period.name, "status": period.status}


def _mark_period_closed(period):
    period.status_before_close = period.status
    period.closed_on = now_datetime()
    period.closed_by = frappe.session.user
    period.status = "Cerrado"
    with period_write_action("close"):
        period.save()


@frappe.whitelist(methods=["POST"])
def reopen_period(period_name: str, reason: str):
    frappe.only_for(["System Manager", "Supervisor Credinomina"])
    period = frappe.get_doc("CN Reconciliation Period", period_name)
    period.check_permission("write")
    if period.status != "Cerrado":
        frappe.throw(_("Solo se puede reabrir un período cerrado."))
    reason = clean_text(reason)
    if not reason:
        frappe.throw(_("Indique el motivo de la reapertura."))
    previous_status = clean_text(period.status_before_close)
    period.status = (
        previous_status if previous_status and previous_status != "Cerrado"
        else "Historico conciliado" if period.reconciliation_mode == "Historica"
        else "Deposito conciliado"
    )
    period.reopened_on = now_datetime()
    period.reopened_by = frappe.session.user
    period.reopen_reason = reason
    with period_write_action("reopen"):
        period.save()
    return {"period": period.name, "status": period.status}


@frappe.whitelist(methods=["POST"])
def export_collection(period_name: str):
    period = frappe.get_doc("CN Reconciliation Period", period_name)
    period.check_permission("read")
    if period.reconciliation_mode == "Historica":
        frappe.throw(_("La modalidad histórica no genera archivo de cobranza."))
    if not period.collection_rows:
        frappe.throw(_("El periodo no tiene detalle para exportar."))
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        frappe.throw(_("Se requiere openpyxl para exportar el archivo."))

    headers = [
        "Nro. Cliente",
        "Nro. Empleado",
        "Nombre y Apellidos del Cliente",
        "Nro Cédula",
        "Nro. Crédito",
        "Nro. cuota",
        "Nro. de cuotas totales",
        "Monto de la cuota en US$",
        "Monto de la cuota en C$",
        "Comentarios",
        "Referencia de Aplicación",
        "Comentario de Aplicación",
        "Deducido C$",
        "Deducido US$",
        "Fila ID",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cobranza"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in period.collection_rows:
        sheet.append(
            [
                row.client_number,
                row.employee_number,
                row.client_name,
                row.national_id,
                row.loan_number,
                row.installment_number,
                row.total_installments,
                flt(row.expected_usd),
                flt(row.expected_nio),
                row.comments,
                row.application_reference,
                row.application_comment,
                None,
                None,
                row.row_key,
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [15, 17, 38, 20, 16, 12, 20, 22, 22, 32, 26, 32, 18, 18, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    stream = io.BytesIO()
    workbook.save(stream)
    content = stream.getvalue()
    file_name = f"cobranza_{period.name}_{cycle_code(period.collection_cycle or MONTHLY)}.xlsx"
    file_doc = frappe.get_doc(
        {
            "doctype": "File",
            "file_name": file_name,
            "attached_to_doctype": period.doctype,
            "attached_to_name": period.name,
            "is_private": 1,
            "content": content,
        }
    ).insert(ignore_permissions=True)
    return {"file_url": file_doc.file_url, "file_name": file_name}


def _create_exception(
    period,
    *,
    exception_type,
    source_row=None,
    client_number=None,
    loan_number=None,
    amount_usd=0,
    amount_nio=0,
    description=None,
    collection_row_id=None,
):
    exception = frappe.get_doc(
        {
            "doctype": "CN Reconciliation Exception",
            "exception_type": exception_type,
            "period": period.name,
            "employer": period.employer,
            "source_row": source_row,
            "client_number": client_number,
            "loan_number": loan_number,
            "amount_usd": amount_usd,
            "amount_nio": amount_nio,
            "collection_row_id": collection_row_id,
            "description": description,
            "status": "Abierta",
        }
    )
    exception.flags.skip_comment_reconciliation = True
    exception.insert(ignore_permissions=True)


def _append_note(existing, line):
    return "\n".join(part for part in (existing, line) if part)


def _deduction_equivalents(row, response):
    deducted_usd = flt(response.get("deducted_usd"))
    deducted_nio = flt(response.get("deducted_nio"))
    note = ""
    if (
        deducted_nio > 0
        and deducted_usd == 0
        and flt(row.expected_nio) > 0
        and flt(row.expected_usd) > 0
    ):
        deducted_usd = deducted_nio * flt(row.expected_usd) / flt(row.expected_nio)
        note = _(
            "El equivalente US$ se calculo con la relacion congelada del archivo de cobranza."
        )
    elif (
        deducted_usd > 0
        and deducted_nio == 0
        and flt(row.expected_usd) > 0
        and flt(row.expected_nio) > 0
    ):
        deducted_nio = deducted_usd * flt(row.expected_nio) / flt(row.expected_usd)
        note = _(
            "El equivalente C$ se calculo con la relacion congelada del archivo de cobranza."
        )
    return round(deducted_usd, 4), round(deducted_nio, 4), note


def _supersede_open_period_exceptions(period_name):
    names = frappe.get_all(
        "CN Reconciliation Exception",
        filters={
            "period": period_name,
            "source_import": ["is", "not set"],
            "status": ["in", ["Abierta", "En revision"]],
        },
        pluck="name",
    )
    for name in names:
        frappe.db.set_value(
            "CN Reconciliation Exception",
            name,
            {
                "status": "Descartada",
                "resolution": _("Reemplazada por una nueva importacion del detalle de empresa."),
            },
            update_modified=False,
        )
