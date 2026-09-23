"""XLSX templates matching the collection and deduction importers."""

from __future__ import annotations

import io
from collections.abc import Iterable, Mapping
from typing import Any


COLLECTION_HEADERS = (
    "Nro. Cliente",
    "Nro. Empleado",
    "Nombre y Apellidos del Cliente",
    "Nro Cédula",
    "Nro. Crédito",
    "Nro. cuota",
    "Nro. de cuotas totales",
    "Monto de la cuota en US$",
    "Monto de la cuota en C$",
    "Comentarios",
    "Referencia de Aplicación",
    "Comentario de Aplicación",
    "Fila ID",
)

DETAIL_HEADERS = (
    *COLLECTION_HEADERS[:-1],
    "Deducido C$",
    "Deducido US$",
    COLLECTION_HEADERS[-1],
)

TEMPLATE_TYPES = {
    "cobranza": ("Cobranza", COLLECTION_HEADERS, "plantilla_cobranza.xlsx"),
    "empresa": ("Detalle empresa", DETAIL_HEADERS, "plantilla_detalle_empresa.xlsx"),
    "deposito": ("Detalle depósito", DETAIL_HEADERS, "plantilla_detalle_deposito.xlsx"),
}

_ROW_FIELDS = (
    "client_number", "employee_number", "client_name", "national_id",
    "loan_number", "installment_number", "total_installments", "expected_usd",
    "expected_nio", "comments", "application_reference", "application_comment",
)

_HEADER_NOTES = {
    "Nro. Cliente": "Identificador del cliente en el core. Conserve los ceros a la izquierda.",
    "Nro. Empleado": "Número interno de la empresa; distinto del número de cliente.",
    "Nombre y Apellidos del Cliente": "Obligatorio en todas las filas. Debe identificar al cliente o uno de sus alias.",
    "Nro Cédula": "Opcional si el cliente puede identificarse por otro dato; conserve los ceros a la izquierda.",
    "Nro. Crédito": "Obligatorio en cobranza. En los detalles facilita la conciliación.",
    "Nro. cuota": "Número de cuota del crédito, si se conoce.",
    "Nro. de cuotas totales": "Total de cuotas pactadas, si se conoce.",
    "Monto de la cuota en US$": "Importe enviado a cobrar en US$. Informe este o el importe en C$.",
    "Monto de la cuota en C$": "Importe enviado a cobrar en C$. Informe este o el importe en US$.",
    "Comentarios": "Observaciones de la cobranza o de la deducción.",
    "Referencia de Aplicación": "Referencia que puede vincular la aplicación del core con la cuota.",
    "Comentario de Aplicación": "Observaciones de la aplicación, si existen.",
    "Deducido C$": "Importe retenido al empleado en C$. Informe este o Deducido US$; use 0 si no se dedujo.",
    "Deducido US$": "Importe retenido al empleado en US$. Informe este o Deducido C$; use 0 si no se dedujo.",
    "Fila ID": "Opcional en cobranza; se genera si queda vacío. En los detalles conserve el ID exportado para identificar la cuota.",
}

_COLUMN_WIDTHS = (17, 17, 38, 21, 17, 13, 22, 24, 24, 32, 27, 32, 19, 19, 28)
_TEXT_COLUMNS = (1, 2, 4, 5, 6, 7, 11)


def build_template_xlsx(
    template_type: str, collection_rows: Iterable[Mapping[str, Any]] = (),
) -> bytes:
    """Build a header-only workbook or prefill existing collection identities."""
    if template_type not in TEMPLATE_TYPES:
        raise ValueError(f"Tipo de plantilla no admitido: {template_type}")

    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet_name, headers, _filename = TEMPLATE_TYPES[template_type]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.comment = Comment(_HEADER_NOTES[cell.value], "Credinómina")
        sheet.column_dimensions[cell.column_letter].width = _COLUMN_WIDTHS[cell.column - 1]
    sheet.row_dimensions[1].height = 34

    if template_type != "cobranza":
        for row in collection_rows:
            values = [row.get(field) for field in _ROW_FIELDS]
            sheet.append([*values, None, None, row.get("row_key")])

    for column in (*_TEXT_COLUMNS, len(headers)):
        sheet.column_dimensions[get_column_letter(column)].number_format = "@"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, sheet.max_row)}"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
