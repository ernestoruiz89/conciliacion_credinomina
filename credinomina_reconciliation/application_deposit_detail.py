"""Explicit, auditable reuse of pending applications as a deposit detail."""

from collections import defaultdict
import hashlib
import io
import json
import re

import frappe
from frappe import _

from credinomina_reconciliation.remittance_selection import target_key
from credinomina_reconciliation.rounding import money, money_float, sum_money
from credinomina_reconciliation.templates import build_template_xlsx, DEPOSIT_HEADERS
from credinomina_reconciliation.tolerance_items import CATEGORY as TOLERANCE_CATEGORY


def pending_application_rows(candidates, deposits, movements, current_name):
    paid = defaultdict(lambda: money(0))
    confirmed = set()
    for deposit in deposits:
        if deposit.get("docstatus") != 1 or deposit["name"] == current_name:
            continue
        confirmed.add(deposit["name"])
        for entry in json.loads(deposit.get("allocation_detail") or "[]"):
            key = target_key(entry)
            if key:
                paid[key] += money(entry.get("importe_usd"))
    adjustments = defaultdict(lambda: money(0))
    for movement in movements:
        if movement.get("status") == "Vigente" and movement.get("deposit_source_row") in confirmed:
            adjustments[movement["claim_id"]] += min(money(movement.get("signed_amount_usd")), 0)
    result = []
    for candidate in candidates:
        key = target_key(candidate)
        pending = max(money(candidate["applied_usd"]) + adjustments[candidate["claim_id"]] - paid[key], 0)
        if pending:
            result.append({**candidate, "deducted_usd": money_float(pending), "deducted_nio": 0,
                           "source_row": len(result) + 2})
    return result


def _load(remittance_name):
    document = frappe.get_doc("CN Remittance Allocation", remittance_name)
    document.check_permission("write")
    if document.docstatus == 2:
        frappe.throw(_("El depósito está cancelado."))
    document._assert_open_related_periods()
    if not document.detail_period:
        frappe.throw(_("Seleccione y guarde el período del detalle."))
    period = frappe.get_doc("CN Reconciliation Period", document.detail_period)
    period.check_permission("read")
    if period.status == "Cerrado":
        frappe.throw(_("El período está cerrado."))
    if not document.employer or period.employer != document.employer:
        frappe.throw(_("El período debe pertenecer a la empresa del depósito."))
    return document, period


def _preview(document, period):
    candidates = []
    identity = ("client_number", "employee_number", "client_name", "national_id", "loan_number", "installment_number")
    if period.reconciliation_mode == "Historica":
        parents = frappe.get_list("CN Source Import", filters={
            "source_type": "Movimientos contables", "status": ["in", ["Importado", "Importado con excepciones"]],
        }, pluck="name", limit_page_length=0)
        rows = frappe.get_all("CN Source Row", filters={
            "parent": ["in", parents], "parenttype": "CN Source Import", "parentfield": "rows",
            "historical_period": period.name, "event_type": "Aplicacion", "effective": 1,
            "match_status": "Conciliado", "currency": "USD",
        }, fields=["name", "parent", "amount", "reference", *identity],
            order_by="event_date asc, name asc", limit_page_length=0) if parents else []
        for row in rows:
            candidates.append({**{field: row.get(field) for field in identity},
                "claim_id": "H:" + row.name, "historical_application": row.name,
                "applied_usd": money_float(row.amount), "application_reference": row.reference,
                "application_comment": f"{row.parent} / {row.name}", "row_key": ""})
    else:
        for row in period.collection_rows:
            if money(row.applied_usd) <= 0:
                continue
            if not row.row_key:
                frappe.throw(_("Una fila aplicada no tiene Fila ID. Actualice la conciliación del período."))
            candidates.append({**{field: row.get(field) for field in identity},
                "claim_id": "C:" + row.name, "period": period.name, "row_key": row.row_key,
                "applied_usd": money_float(row.applied_usd),
                "application_reference": row.application_reference,
                "application_comment": f"{period.name} / {row.name}"})
    # Settlement totals must include other deposits even if the operator cannot
    # open them; their details are not exposed in the preview.
    deposits = frappe.get_all("CN Remittance Allocation", filters={"employer": document.employer, "docstatus": 1},
        fields=["name", "docstatus", "allocation_detail"], limit_page_length=0)
    movements = frappe.get_all("CN Complementary Item", filters={
        "category": TOLERANCE_CATEGORY, "docstatus": 1,
        "employer": document.employer, "period": period.name, "status": "Vigente",
    }, fields=["claim_id", "deposit_source_row", "signed_amount_usd", "status"], limit_page_length=0)
    rows = pending_application_rows(candidates, deposits, movements, document.name)
    for row in rows:
        if not (row.get("client_name") or "").strip():
            frappe.throw(_("Complete el nombre del cliente en las aplicaciones del período antes de usarlas como detalle."))
        row["comments"] = _("Generado desde aplicaciones pendientes del período {0}; no es un detalle recibido de la empresa.").format(period.name)
    result = {"period": period.name, "applied_usd": money_float(period.applied_usd),
              "rows": rows, "total_usd": money_float(sum_money(row["deducted_usd"] for row in rows)),
              "deposit_usd": money_float(document.amount_usd), "modified": str(document.modified),
              "replaces_detail": bool(document.detail_rows or document.detail_file)}
    result["fingerprint"] = hashlib.sha256(json.dumps(result, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return result


def preview_application_detail(remittance_name):
    return _preview(*_load(remittance_name))


def _workbook(rows):
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    workbook = load_workbook(io.BytesIO(build_template_xlsx("deposito", rows)))
    sheet = workbook.active
    amount_column = DEPOSIT_HEADERS.index("Deducido US$") + 1
    sheet.cell(1, amount_column).comment = Comment(
        "Aplicación pendiente de cubrir con otros depósitos, expresada en US$. "
        "Generado desde el core; no confirma una deducción informada por la empresa.", "Credinómina")
    for index, row in enumerate(rows, 2):
        sheet.cell(index, amount_column, row["deducted_usd"]).number_format = "0.00"
    # References and client names are text, never executable spreadsheet formulas.
    for row in sheet:
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def select_application_rows(rows, selected_claim_ids):
    """Accept identities only; amounts and eligibility always come from the server."""
    if selected_claim_ids is None:
        selected_claim_ids = [row["claim_id"] for row in rows]
    if isinstance(selected_claim_ids, str):
        try:
            selected_claim_ids = json.loads(selected_claim_ids)
        except (ValueError, TypeError):
            frappe.throw(_("La selección de movimientos no es válida."))
    if not isinstance(selected_claim_ids, list) or any(not isinstance(key, str) for key in selected_claim_ids):
        frappe.throw(_("La selección de movimientos no es válida."))
    if not selected_claim_ids:
        frappe.throw(_("Seleccione al menos un movimiento para generar el detalle."))
    selected = set(selected_claim_ids)
    available = {row["claim_id"] for row in rows}
    if len(selected) != len(selected_claim_ids) or not selected.issubset(available):
        frappe.throw(_("Hay movimientos duplicados o que ya no están disponibles. Abra nuevamente la vista previa."))
    chosen = [row for row in rows if row["claim_id"] in selected]
    return [{**row, "source_row": index} for index, row in enumerate(chosen, 2)]


def use_application_detail(remittance_name, fingerprint, replace_detail=False, selected_claim_ids=None):
    from frappe.utils.file_manager import save_file
    from frappe.utils import cint
    from credinomina_reconciliation.conciliacion_credinomina.doctype.cn_remittance_allocation.cn_remittance_allocation import (
        _apply_remittance_detail,
    )
    frappe.db.sql("select name from `tabCN Remittance Allocation` where name=%s for update", (remittance_name,))
    document, period = _load(remittance_name)
    preview = _preview(document, period)
    if fingerprint != preview["fingerprint"]:
        frappe.throw(_("Los datos cambiaron. Abra nuevamente la vista previa antes de generar el detalle."))
    if not preview["rows"]:
        frappe.throw(_("No hay aplicaciones pendientes para este período."))
    if preview["replaces_detail"] and not cint(replace_detail):
        frappe.throw(_("Confirme el reemplazo del detalle existente."))
    rows = select_application_rows(preview["rows"], selected_claim_ids)
    total_usd = money_float(sum_money(row["deducted_usd"] for row in rows))
    content = _workbook(rows)
    safe_name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", document.name)
    attachment = save_file(f"detalle_aplicaciones_{safe_name}.xlsx", content,
                           document.doctype, document.name, is_private=1, df="detail_file")
    document.detail_file = attachment.file_url
    document.detail_total_usd = total_usd
    document.result = "Pendiente"
    _apply_remittance_detail(document, rows, content, attachment.file_url,
                             origin="Aplicaciones pendientes del período")
    document.add_comment("Comment", _(
        "Se generó detalle desde aplicaciones pendientes de {0}: {1} filas seleccionadas de {3}, US$ {2}. "
        "No es evidencia enviada por la empresa. Se conservaron los destinos, "
        "pero se quitaron sus vínculos al detalle anterior. No se confirmó ni concilió el depósito."
    ).format(period.name, len(rows), total_usd, len(preview["rows"])))
    return {"rows": len(rows), "total_usd": total_usd, "file_url": attachment.file_url}
