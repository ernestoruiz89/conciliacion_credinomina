"""Pure file parsers for the standalone Credinomina reconciliation app.

This module intentionally has no Frappe imports so it can be tested with the
real source files before the app is installed on a site.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable


SOURCE_ACCOUNTING = "Movimientos contables (principal)"
SOURCE_TRANSACTIONS = "Transacciones del core (fallback)"
SOURCE_DEPOSITS = "Detalle de depositos"

SOURCE_TYPES = (SOURCE_ACCOUNTING, SOURCE_TRANSACTIONS, SOURCE_DEPOSITS)


class SourceFileError(ValueError):
    pass


KNOWN_MOJIBAKE = {
    "c\ufffddula": "cedula",
    "cr\ufffddito": "credito",
    "aplicaci\ufffdn": "aplicacion",
    "dep\ufffdsito": "deposito",
}


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    for broken, repaired in KNOWN_MOJIBAKE.items():
        text = text.replace(broken, repaired)
    text = "".join(
        character
        for character in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(character)
    )
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    return text


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().lstrip("'")


def canonical_identifier(value: Any) -> str:
    """Normalize deterministic spreadsheet formatting, never names or fuzzy text."""
    text = clean_text(value)
    if re.fullmatch(r"\d+", text):
        return text.lstrip("0") or "0"
    return text.casefold()


def parse_amount(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return float(value)
    text = clean_text(value)
    text = re.sub(r"[^0-9,().+\-]", "", text)
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    if not text:
        return 0.0
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        last = text.rsplit(",", 1)[-1]
        text = text.replace(",", ".") if len(last) <= 2 else text.replace(",", "")
    try:
        result = float(Decimal(text))
    except (InvalidOperation, ValueError) as exc:
        raise SourceFileError(f"Importe no valido: {value!r}") from exc
    return -result if negative else result


def parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def file_sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def source_key(*parts: Any) -> str:
    payload = "|".join(clean_text(part).casefold() for part in parts)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def read_table(
    file_name: str, content: bytes, *, sheet_name: str | None = None
) -> list[list[Any]]:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise SourceFileError("Se requiere openpyxl para leer archivos .xlsx.") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet_name:
                sheet = next(
                    (item for item in workbook.worksheets
                     if normalize_header(item.title) == normalize_header(sheet_name)),
                    None,
                )
                if sheet is None:
                    if len(workbook.worksheets) == 1:
                        sheet = workbook.worksheets[0]
                    else:
                        raise SourceFileError(f"No se encontró la pestaña {sheet_name}.")
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise SourceFileError("Se requiere xlrd para leer archivos .xls.") from exc
        workbook = xlrd.open_workbook(file_contents=content)
        sheet = workbook.sheet_by_index(0)
        if sheet_name:
            sheet = next(
                (workbook.sheet_by_index(index) for index in range(workbook.nsheets)
                 if normalize_header(workbook.sheet_by_index(index).name)
                 == normalize_header(sheet_name)),
                None,
            )
            if sheet is None:
                if workbook.nsheets == 1:
                    sheet = workbook.sheet_by_index(0)
                else:
                    raise SourceFileError(f"No se encontró la pestaña {sheet_name}.")
        result = []
        for row_index in range(sheet.nrows):
            row = []
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
                else:
                    row.append(cell.value)
            result.append(row)
        return result
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    raise SourceFileError("Use un archivo .xlsx, .xls o .csv.")


COLLECTION_ALIASES = {
    "row_key": ("fila_id", "row_id", "id_fila"),
    "client_number": ("nro_cliente", "numero_cliente", "no_cliente", "customer_id"),
    "client_name": (
        "nombre_y_apellidos_del_cliente",
        "nombre_del_cliente",
        "cliente",
        "customer_name",
    ),
    "national_id": ("nro_cedula", "numero_cedula", "cedula", "identificacion"),
    "loan_number": (
        "nro_credito",
        "numero_credito",
        "no_credito",
        "credito",
        "loan",
    ),
    "installment_number": ("nro_cuota", "numero_cuota", "no_cuota"),
    "total_installments": (
        "nro_de_cuotas_totales",
        "numero_de_cuotas_totales",
        "cuotas_totales",
    ),
    "expected_usd": ("monto_de_la_cuota_en_us", "monto_de_la_cuota_en_usd", "cuota_usd"),
    "expected_nio": ("monto_de_la_cuota_en_c", "monto_de_la_cuota_en_cs", "monto_de_la_cuota_en_nio", "cuota_nio"),
    "comments": ("comentarios", "comentario", "observacion", "observaciones"),
    "application_reference": ("referencia_de_aplicacion", "referencia_aplicacion"),
    "application_comment": ("comentario_de_aplicacion", "comentario_aplicacion"),
    "deducted_nio": ("deducido_c", "deducido_cs", "deducido_nio", "monto_deducido_nio"),
    "deducted_usd": ("deducido_us", "deducido_usd", "monto_deducido_usd"),
}


def header_mapping(row: Iterable[Any], aliases: dict[str, tuple[str, ...]]) -> dict[str, int]:
    normalized = [normalize_header(value) for value in row]
    result = {}
    for fieldname, candidates in aliases.items():
        for candidate in candidates:
            if candidate in normalized:
                result[fieldname] = normalized.index(candidate)
                break
    return result


def _value(row: list[Any], mapping: dict[str, int], fieldname: str) -> Any:
    index = mapping.get(fieldname)
    return row[index] if index is not None and index < len(row) else None


def parse_collection_file(
    file_name: str, content: bytes, *, require_deduction: bool = False
) -> list[dict[str, Any]]:
    rows = read_table(file_name, content)
    mapping: dict[str, int] | None = None
    parsed: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows, start=1):
        candidate = header_mapping(row, COLLECTION_ALIASES)
        if "loan_number" in candidate and (
            "client_number" in candidate or "national_id" in candidate
        ):
            mapping = candidate
            continue
        if not mapping:
            continue
        loan_number = clean_text(_value(row, mapping, "loan_number"))
        client_number = clean_text(_value(row, mapping, "client_number"))
        national_id = clean_text(_value(row, mapping, "national_id"))
        if not any((loan_number, client_number, national_id)):
            continue
        expected_usd = parse_amount(_value(row, mapping, "expected_usd"))
        expected_nio = parse_amount(_value(row, mapping, "expected_nio"))
        deducted_usd = parse_amount(_value(row, mapping, "deducted_usd"))
        deducted_nio = parse_amount(_value(row, mapping, "deducted_nio"))
        if not any((expected_usd, expected_nio, deducted_usd, deducted_nio)):
            continue
        if require_deduction and not (
            "deducted_usd" in mapping or "deducted_nio" in mapping
        ):
            raise SourceFileError(
                "El detalle de la empresa debe contener Deducido C$ o Deducido US$."
            )
        parsed.append(
            {
                "source_row": row_number,
                "row_key": clean_text(_value(row, mapping, "row_key")),
                "client_number": client_number,
                "client_name": clean_text(_value(row, mapping, "client_name")),
                "national_id": national_id,
                "loan_number": loan_number,
                "installment_number": clean_text(_value(row, mapping, "installment_number")),
                "total_installments": clean_text(_value(row, mapping, "total_installments")),
                "expected_usd": expected_usd,
                "expected_nio": expected_nio,
                "comments": clean_text(_value(row, mapping, "comments")),
                "application_reference": clean_text(
                    _value(row, mapping, "application_reference")
                ),
                "application_comment": clean_text(
                    _value(row, mapping, "application_comment")
                ),
                "deducted_usd": deducted_usd,
                "deducted_nio": deducted_nio,
            }
        )
    if not parsed:
        raise SourceFileError("No se encontraron filas de cobranza reconocibles.")
    return parsed


def _records_from_header(rows: list[list[Any]], required_header: str) -> list[tuple[int, dict[str, Any]]]:
    mapping = None
    records = []
    for row_number, row in enumerate(rows, start=1):
        normalized = [normalize_header(value) for value in row]
        if required_header in normalized:
            mapping = {header: index for index, header in enumerate(normalized) if header}
            continue
        if not mapping or not any(value not in (None, "") for value in row):
            continue
        records.append(
            (
                row_number,
                {
                    header: row[index] if index < len(row) else None
                    for header, index in mapping.items()
                },
            )
        )
    return records


def _extract_reference(description: str, fallback: Any = None) -> str:
    match = re.search(r"\|\s*REF\s*:\s*([^|]+)\|", description, re.IGNORECASE)
    if match:
        return clean_text(match.group(1))
    match = re.search(r"\|\s*([^|]+)\s*\|", description)
    return clean_text(match.group(1) if match else fallback)


def _extract_employer(description: str, fallback: Any = None) -> str:
    match = re.search(
        r"CONVENIO\s+(.+?)(?:\s*\(|\s+EN\s+LA\s+CUENTA|\s+-|\s+\||$)",
        description,
        re.IGNORECASE,
    )
    return clean_text(match.group(1) if match else fallback)


def _native_deposit_amount(description: str, fallback: Any) -> tuple[str, float]:
    match = re.search(r"\b(U\$|US\$|C\$)\s*([0-9.,]+)", description, re.IGNORECASE)
    if match:
        currency = "USD" if "U" in match.group(1).upper() else "NIO"
        return currency, parse_amount(match.group(2))
    return "USD", parse_amount(fallback)


def parse_accounting_movements(file_name: str, content: bytes) -> list[dict[str, Any]]:
    records = _records_from_header(read_table(file_name, content), "cuenta_contable")
    parsed = []
    for row_number, record in records:
        description = clean_text(record.get("descripcion"))
        if not description:
            continue
        upper = description.upper()
        event_type = None
        loan_number = ""
        reference = ""
        employer = ""
        if "NOTA AL PRESTAMO" in upper:
            event_type = "Aplicacion"
            match = re.search(r"NOTA\s+AL\s+PRESTAMO\s+0*([0-9]+)", description, re.IGNORECASE)
            loan_number = clean_text(match.group(1) if match else "")
            reference = _extract_reference(description, record.get("no_ref"))
            account_name = clean_text(record.get("descripcion_cta_contable")).upper()
            currency = "NIO" if "M.N" in account_name else "USD"
            amount = parse_amount(record.get("debito_del_mes"))
            employer = _extract_employer(description)
        elif "DEPOSITO POR" in upper:
            event_type = "Deposito"
            reference = clean_text(record.get("no_ref"))
            currency, amount = _native_deposit_amount(
                description, record.get("credito_del_mes")
            )
            employer = _extract_employer(description)
        else:
            continue
        if amount <= 0:
            continue
        event_date = parse_date(record.get("fecha_aplica"))
        equivalent_currency = ""
        equivalent_amount = 0.0
        fx_basis = ""
        if event_type == "Deposito":
            account_name = clean_text(record.get("descripcion_cta_contable")).upper()
            accounting_currency = (
                "USD" if "M.E." in account_name else "NIO" if "M.N." in account_name else ""
            )
            accounting_amount = parse_amount(record.get("credito_del_mes"))
            if accounting_currency and accounting_currency != currency and accounting_amount > 0:
                equivalent_currency = accounting_currency
                equivalent_amount = accounting_amount
                fx_basis = "Importe del movimiento contable"
        parsed.append(
            _source_record(
                row_number=row_number,
                event_type=event_type,
                event_date=event_date,
                reference=reference,
                voucher=record.get("no_cmpte"),
                employer=employer,
                loan_number=loan_number,
                currency=currency,
                amount=amount,
                description=description,
                equivalent_currency=equivalent_currency,
                equivalent_amount=equivalent_amount,
                fx_basis=fx_basis,
            )
        )
    if not parsed:
        raise SourceFileError("No se encontraron aplicaciones o depositos contables.")
    return parsed


def parse_transactions(file_name: str, content: bytes) -> list[dict[str, Any]]:
    records = _records_from_header(read_table(file_name, content), "cod_trans")
    parsed = []
    for row_number, record in records:
        transaction = clean_text(record.get("desc_transaccion"))
        if not transaction:
            continue
        event_type = "Ajuste" if "DISPENSA" in transaction.upper() else "Aplicacion"
        if event_type == "Aplicacion" and "DEPOSITO" not in transaction.upper():
            continue
        currency_text = clean_text(record.get("moneda")).upper()
        currency = "NIO" if "CORD" in currency_text else "USD"
        amount = parse_amount(record.get("total"))
        if amount <= 0:
            continue
        transaction_rate = parse_amount(record.get("tipo_cambio"))
        equivalent_currency = "NIO" if currency == "USD" else "USD"
        equivalent_amount = (
            amount * transaction_rate
            if currency == "USD"
            else amount / transaction_rate if transaction_rate > 0 else 0
        )
        reference = clean_text(record.get("referencia")) or _extract_reference(
            clean_text(record.get("concepto"))
        )
        parsed.append(
            _source_record(
                row_number=row_number,
                event_type=event_type,
                event_date=parse_date(record.get("fecha")),
                reference=reference,
                voucher=record.get("nro_comprobante"),
                employer=record.get("convenio"),
                client_number=record.get("nrocliente"),
                client_name=record.get("cliente"),
                national_id=record.get("identificacion"),
                loan_number=record.get("nrocredito"),
                installment_number=record.get("nrocuota"),
                currency=currency,
                amount=amount,
                description=record.get("concepto") or transaction,
                equivalent_currency=equivalent_currency if transaction_rate > 0 else "",
                equivalent_amount=equivalent_amount,
                fx_basis="TIPO_CAMBIO de Transacciones" if transaction_rate > 0 else "",
            )
        )
    if not parsed:
        raise SourceFileError("No se encontraron aplicaciones en Transacciones.")
    return parsed


def parse_deposit_detail(file_name: str, content: bytes) -> list[dict[str, Any]]:
    sheet_name = "Depósito" if Path(file_name).suffix.lower() in {".xlsx", ".xls"} else None
    records = _records_from_header(
        read_table(file_name, content, sheet_name=sheet_name), "referencia"
    )
    parsed = []
    for row_number, record in records:
        nio_amount = parse_amount(record.get("valorc"))
        usd_amount = parse_amount(record.get("valoru"))
        if nio_amount <= 0 and usd_amount <= 0:
            continue
        description = clean_text(record.get("descripcion"))
        client_name = clean_text(record.get("cliente"))
        native_amounts = [("NIO", nio_amount), ("USD", usd_amount)]
        for currency, amount in native_amounts:
            if amount <= 0:
                continue
            reported_usd = parse_amount(record.get("monto_u"))
            has_separate_currency_amounts = nio_amount > 0 and usd_amount > 0
            parsed.append(
                _source_record(
                    row_number=row_number,
                    event_type="Deposito",
                    event_date=parse_date(record.get("fecha")),
                    reference=record.get("referencia"),
                    voucher=record.get("banco"),
                    employer=_extract_employer(client_name, client_name),
                    client_name=client_name,
                    loan_number=record.get("no_credito"),
                    currency=currency,
                    amount=amount,
                    description=description,
                    equivalent_currency="USD" if currency == "NIO" and reported_usd > 0 and not has_separate_currency_amounts else "",
                    equivalent_amount=reported_usd if currency == "NIO" and not has_separate_currency_amounts else 0,
                    fx_basis="MONTO U$ del detalle de deposito" if currency == "NIO" and reported_usd > 0 and not has_separate_currency_amounts else "",
                )
            )
    if not parsed:
        raise SourceFileError("No se encontraron depositos reconocibles.")
    return parsed


def _source_record(
    *,
    row_number: int,
    event_type: str,
    event_date: date | None,
    reference: Any,
    voucher: Any = None,
    employer: Any = None,
    client_number: Any = None,
    client_name: Any = None,
    national_id: Any = None,
    loan_number: Any = None,
    installment_number: Any = None,
    currency: str,
    amount: float,
    description: Any = None,
    equivalent_currency: str = "",
    equivalent_amount: float = 0,
    fx_basis: str = "",
) -> dict[str, Any]:
    cleaned = {
        "source_row": row_number,
        "event_type": event_type,
        "event_date": event_date,
        "reference": clean_text(reference),
        "voucher": clean_text(voucher),
        "employer_text": clean_text(employer),
        "client_number": clean_text(client_number),
        "client_name": clean_text(client_name),
        "national_id": clean_text(national_id),
        "loan_number": clean_text(loan_number),
        "installment_number": clean_text(installment_number),
        "currency": currency,
        "amount": round(float(amount), 4),
        "description": clean_text(description),
        "equivalent_currency": equivalent_currency,
        "equivalent_amount": round(float(equivalent_amount or 0), 4),
        "fx_basis": fx_basis,
    }
    cleaned["fx_rate"] = 0.0
    if cleaned["amount"] > 0 and cleaned["equivalent_amount"] > 0:
        if currency == "NIO" and equivalent_currency == "USD":
            cleaned["fx_rate"] = round(cleaned["amount"] / cleaned["equivalent_amount"], 8)
        elif currency == "USD" and equivalent_currency == "NIO":
            cleaned["fx_rate"] = round(cleaned["equivalent_amount"] / cleaned["amount"], 8)
    cleaned["amount_usd"] = cleaned["amount"] if currency == "USD" else 0
    cleaned["amount_nio"] = cleaned["amount"] if currency == "NIO" else 0
    cleaned["source_key"] = source_key(
        event_type,
        event_date,
        cleaned["reference"],
        cleaned["loan_number"],
        currency,
        cleaned["amount"],
    )
    return cleaned


def parse_source_file(source_type: str, file_name: str, content: bytes) -> list[dict[str, Any]]:
    if source_type == SOURCE_ACCOUNTING:
        return parse_accounting_movements(file_name, content)
    if source_type == SOURCE_TRANSACTIONS:
        return parse_transactions(file_name, content)
    if source_type == SOURCE_DEPOSITS:
        return parse_deposit_detail(file_name, content)
    raise SourceFileError(f"Tipo de fuente no soportado: {source_type}")
