"""Build the Spanish operating manual from its maintained Markdown source.

Design reference: Loan Manager's build_branch_manager_manual.py and
build_payroll_nicaragua_manual.py (cover, Calibri, Letter, routes, numbered
procedures, comparison tables and review checklists). No Loan Manager imports
or dependencies are required. Run with python-docx in the document runtime.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "procedimiento_operativo.md"
OUTPUT = ROOT / "docs" / "Manual_Operativo_Conciliacion_Credinomina.docx"
BLACK, NAVY, PALE, BORDER = "000000", "17365D", "F3F7FB", "D9D9D9"
WIDTH = 9878


def node(tag, **attrs):
    element = OxmlElement("w:" + tag)
    for key, value in attrs.items():
        element.set(qn("w:" + key), str(value))
    return element


def font(run, *, size=11, bold=False, color=BLACK):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    props = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        props.rFonts.set(qn("w:" + key), "Calibri")


def inline(paragraph, text, *, size=11, color=BLACK):
    # Keep the manual editable; emphasis and code labels remain native runs.
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if bold else part[1:-1] if part.startswith("`") else part
        font(paragraph.add_run(value), size=size, bold=bold, color=color)


def numbering(doc):
    root = doc.part.numbering_part.element
    abstract_id = max([int(e.get(qn("w:abstractNumId"))) for e in root.findall(qn("w:abstractNum"))] + [-1]) + 1
    num_id = max([int(e.get(qn("w:numId"))) for e in root.findall(qn("w:num"))] + [0]) + 1
    abstract = node("abstractNum", abstractNumId=abstract_id)
    abstract.append(node("nsid", val=f"{0x434E0000 + abstract_id:08X}"))
    abstract.append(node("multiLevelType", val="singleLevel"))
    level = node("lvl", ilvl=0)
    for tag, val in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1."), ("lvlJc", "left")):
        level.append(node(tag, val=val))
    props = node("pPr")
    props.append(node("ind", left=540, hanging=270))
    level.append(props)
    abstract.append(level)
    # OOXML requires all abstract definitions before concrete num records.
    # Interleaving them makes Word repair IDs and join unrelated lists.
    first_num = root.find(qn("w:num"))
    root.insert(list(root).index(first_num) if first_num is not None else len(root), abstract)
    num = node("num", numId=num_id)
    num.append(node("abstractNumId", val=abstract_id))
    restart = node("lvlOverride", ilvl=0)
    restart.append(node("startOverride", val=1))
    num.append(restart)
    root.append(num)
    return num_id


def numbered_paragraph(doc, text, num_id):
    paragraph = doc.add_paragraph()
    props = paragraph._p.get_or_add_pPr()
    num = node("numPr")
    num.append(node("ilvl", val=0))
    num.append(node("numId", val=num_id))
    props.append(num)
    paragraph.paragraph_format.space_after = Pt(5)
    inline(paragraph, text)


def table(doc, records):
    count = len(records[0])
    fractions = [0.26, 0.74] if count == 2 else [0.24, 0.38, 0.38] if count == 3 else [1 / count] * count
    widths = [round(WIDTH * f) for f in fractions]
    widths[-1] += WIDTH - sum(widths)
    result = doc.add_table(rows=0, cols=count)
    result.autofit = False
    props = result._tbl.tblPr
    props.append(node("tblW", w=WIDTH, type="dxa"))
    borders = node("tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(node(edge, val="single", sz=4, color=BORDER))
    props.append(borders)
    grid = result._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid.append(node("gridCol", w=width))
    for index, values in enumerate(records):
        row = result.add_row()
        tr_props = row._tr.get_or_add_trPr()
        tr_props.append(node("cantSplit"))
        if index == 0:
            tr_props.append(node("tblHeader", val="true"))
        for col, value in enumerate(values):
            cell = row.cells[col]
            cell.width = Inches(widths[col] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_props = cell._tc.get_or_add_tcPr()
            cell_props.append(node("shd", fill=NAVY if index == 0 else PALE if index % 2 else "FFFFFF"))
            margins = node("tcMar")
            for side, amount in (("top", 95), ("bottom", 95), ("start", 120), ("end", 120)):
                margins.append(node(side, w=amount, type="dxa"))
            cell_props.append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
            if index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                font(paragraph.add_run(value), size=10.5, bold=True, color="FFFFFF")
            else:
                inline(paragraph, value, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return result


def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.8), Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.82)
    section.header_distance = section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True
    for name, size in (("Normal", 11), ("Title", 27), ("Subtitle", 14), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5), ("List Bullet", 11)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.18
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 9)
            style.paragraph_format.keep_with_next = True
    # The runtime's default template has a blue Title border. Strip style
    # residue, including inherited borders, rather than merely hiding a run.
    for style in doc.styles:
        ppr = style.element.find(qn("w:pPr"))
        if ppr is not None:
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
    header = section.header.paragraphs[0]
    font(header.add_run("MIDESA  |  MANUAL OPERATIVO DE CONCILIACIÓN CREDINÓMINA"), size=8.5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run("Página "), size=9)
    run = footer.add_run()
    run._r.append(node("fldChar", fldCharType="begin"))
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGE "
    run._r.append(instruction)
    run._r.append(node("fldChar", fldCharType="end"))
    doc.core_properties.title = "Manual operativo de conciliación Credinómina"
    doc.core_properties.subject = "Aplicaciones, depósitos, ajustes, saldos a favor y control de cuentas por cobrar"
    doc.core_properties.author = "MIDESA"
    doc.core_properties.language = "es-NI"


def cover(doc, source_text):
    metadata = re.search(
        r"Versión del manual: ([^.]+\.[^.]+)\. Fecha de revisión: ([^.]+)\. "
        r"Referencia de plataforma: ([^.]+(?:\.[0-9]+)*)\.", source_text,
    )
    if not metadata:
        raise ValueError("Faltan versión, fecha o plataforma en el procedimiento")
    version, reviewed, platform = metadata.groups()
    paragraph = doc.add_paragraph("MIDESA")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(85)
    paragraph.paragraph_format.space_after = Pt(18)
    paragraph.runs[0].bold = True
    title = doc.add_paragraph("Manual operativo de conciliación Credinómina", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    subtitle = doc.add_paragraph("Aplicaciones y depósitos\nAjustes y saldos a favor\nControl de cuentas por cobrar", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(32)
    table(doc, [["Control del documento", "Referencia"], ["Versión", version], ["Fecha de revisión", reviewed], ["Usuarios", "Operadores y supervisores de conciliación"], ["Plataforma", platform], ["Alcance", "App independiente y core externo de la IMF"]])
    doc.add_paragraph("Uso interno. Conserve archivos y datos personales con acceso restringido. Las autorizaciones institucionales deben obtenerse por el procedimiento de la IMF.")
    doc.add_page_break()


def contents(doc, headings):
    p = doc.add_paragraph("Contenido", style="Title")
    p.runs[0].font.size = Pt(20)
    p.paragraph_format.space_after = Pt(14)
    doc.add_paragraph("Seleccione un tema para ir al procedimiento. Los encabezados también están disponibles en el panel de navegación de Word.")
    for index, text in enumerate(headings, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), f"chapter_{index}")
        run = OxmlElement("w:r")
        props = node("rPr")
        props.append(node("color", val=BLACK))
        props.append(node("sz", val=22))
        run.append(props)
        value = node("t")
        value.text = text
        run.append(value)
        link.append(run)
        p._p.append(link)
    doc.add_page_break()


def list_item(lines, index, prefix):
    """Keep Markdown continuation lines inside the same native list item."""
    parts = [re.sub(prefix, "", lines[index].strip())]
    index += 1
    while index < len(lines):
        next_line = lines[index].strip()
        if not next_line or re.match(r"^(#+ |\||- |\d+\. )", next_line):
            break
        parts.append(next_line)
        index += 1
    return " ".join(parts), index


def build():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    headings = [line[3:] for line in lines if line.startswith("## ")]
    doc = Document()
    configure(doc)
    cover(doc, text)
    contents(doc, headings)
    index, chapter, ordered_id = 1, 0, None
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("## ") or line.startswith("### "):
            level = 1 if line.startswith("## ") else 2
            p = doc.add_paragraph(line[level + 2:], style=f"Heading {level}")
            if level == 1:
                chapter += 1
                start = node("bookmarkStart", id=chapter, name=f"chapter_{chapter}")
                p._p.insert(0, start)
                p._p.append(node("bookmarkEnd", id=chapter))
            ordered_id = None
            index += 1
        elif line.startswith("|"):
            records = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = [value.strip() for value in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", value.replace(" ", "")) for value in row):
                    records.append(row)
                index += 1
            table(doc, records)
            ordered_id = None
        elif re.match(r"\d+\. ", line):
            if line.startswith("1. ") or ordered_id is None:
                ordered_id = numbering(doc)
            value, index = list_item(lines, index, r"^\d+\. ")
            numbered_paragraph(doc, value, ordered_id)
        elif line.startswith("- "):
            value, index = list_item(lines, index, r"^- ")
            p = doc.add_paragraph(style="List Bullet")
            inline(p, value)
            ordered_id = None
        else:
            parts = [line]
            index += 1
            while index < len(lines) and lines[index].strip() and not re.match(r"^(#+ |\||- |\d+\. )", lines[index]):
                parts.append(lines[index].strip())
                index += 1
            inline(doc.add_paragraph(), " ".join(parts))
            ordered_id = None
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Saved {OUTPUT}; chapters={len(headings)}; words={len(text.split())}")


if __name__ == "__main__":
    build()
