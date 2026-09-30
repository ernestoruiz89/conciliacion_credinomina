import io
import json
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch

import frappe
from openpyxl import load_workbook

from credinomina_reconciliation import application_deposit_detail as module
from credinomina_reconciliation.parsers import parse_collection_file


class ApplicationDepositDetailTests(unittest.TestCase):
    def test_pending_excludes_other_payments_but_not_current_deposit(self):
        candidates = [dict(historical_application="A", claim_id="H:A", applied_usd=100),
                      dict(period="P", row_key="R", claim_id="C:C", applied_usd=50)]
        def deposit(name, status, history=0, collection=0):
            return dict(name=name, docstatus=status, allocation_detail=json.dumps([
                {"aplicacion_id": "A", "importe_usd": history},
                {"periodo": "P", "fila_id": "R", "importe_usd": collection}]))
        rows = module.pending_application_rows(candidates, [
            deposit("OTHER", 1, 40, 20), deposit("CURRENT", 1, 30, 10),
            deposit("DRAFT", 0, 99, 99), deposit("CANCELLED", 2, 99, 99),
        ], [], "CURRENT")
        self.assertEqual([r["deducted_usd"] for r in rows], [60, 30])

    def test_rounding_shortage_settles_cents_and_reversals_do_not(self):
        candidate = dict(historical_application="A", claim_id="H:A", applied_usd=46.53)
        deposits = [dict(name="D", docstatus=1, allocation_detail=json.dumps([
            {"aplicacion_id": "A", "importe_usd": 46.52}]))]
        movement = dict(status="Vigente", claim_id="H:A", deposit_source_row="D", signed_amount_usd=-0.01)
        self.assertEqual(module.pending_application_rows([candidate], deposits, [movement], "CURRENT"), [])
        movement["status"] = "Revertido"
        self.assertEqual(module.pending_application_rows([candidate], deposits, [movement], "CURRENT")[0]["deducted_usd"], 0.01)

    def test_workbook_roundtrips_names_ids_and_pending_usd_without_formulas(self):
        data = [dict(client_name="=Nombre no ejecutable", client_number="001", loan_number="090-1",
                     deducted_usd=22.52, row_key="ROW", comments="Origen: aplicaciones pendientes")]
        content = module._workbook(data)
        sheet = load_workbook(io.BytesIO(content)).active
        self.assertEqual(sheet.cell(2, 3).data_type, "s")
        rows = parse_collection_file("detail.xlsx", content, require_deduction=True, require_name=True)
        self.assertEqual(rows[0]["client_name"], "=Nombre no ejecutable")
        self.assertEqual(rows[0]["client_number"], "001")
        self.assertEqual(rows[0]["deducted_usd"], 22.52)
        self.assertEqual(rows[0]["row_key"], "ROW")

    def document(self):
        return SimpleNamespace(name="D", doctype="CN Remittance Allocation", employer="E", detail_period="P",
            docstatus=1, check_permission=Mock(), _assert_open_related_periods=Mock(),
            modified="2026-09-30", amount_usd=100, detail_rows=[], detail_file=None, add_comment=Mock())

    def test_load_checks_permission_company_closed_period_and_cancellation(self):
        document = self.document()
        period = SimpleNamespace(name="P", employer="E", status="Abierto", check_permission=Mock())
        with patch.object(module.frappe, "get_doc", side_effect=[document, period]):
            self.assertEqual(module._load("D"), (document, period))
        document.check_permission.assert_called_with("write")
        period.check_permission.assert_called_with("read")
        for status, company, cancelled in [("Cerrado", "E", False), ("Abierto", "OTHER", False), ("Abierto", "E", True)]:
            period.status, period.employer = status, company
            document.docstatus = 2 if cancelled else 1
            with patch.object(module.frappe, "get_doc", side_effect=[document, period]), \
                 patch.object(module, "_", side_effect=lambda v: v), \
                 patch.object(module.frappe, "throw", side_effect=ValueError):
                with self.assertRaises(ValueError):
                    module._load("D")

    def test_preview_historical_uses_readable_imports_and_excludes_duplicates(self):
        document = self.document()
        period = SimpleNamespace(name="P", reconciliation_mode="Historica", applied_usd=100)
        def get_all(doctype, **kwargs):
            if doctype == "CN Source Row":
                filters = kwargs["filters"]
                self.assertEqual(filters["parent"], ["in", ["READABLE"]])
                self.assertEqual(filters["effective"], 1)
                self.assertEqual(filters["match_status"], "Conciliado")
                self.assertEqual(filters["historical_period"], "P")
                return [frappe._dict(name="A", parent="READABLE", amount=100, client_name="Ana", reference="REF")]
            return []
        with patch.object(module.frappe, "get_list", return_value=["READABLE"]), \
             patch.object(module.frappe, "get_all", side_effect=get_all), \
             patch.object(module, "_", side_effect=lambda v: v):
            result = module._preview(document, period)
        self.assertEqual(result["total_usd"], 100)
        self.assertIn("no es un detalle recibido", result["rows"][0]["comments"])

    def test_operative_uses_applied_not_requested_or_deducted_amount(self):
        document = self.document()
        period = SimpleNamespace(name="P", reconciliation_mode="Operativa", applied_usd=30,
            collection_rows=[frappe._dict(name="C", row_key="R", client_name="Ana", applied_usd=30,
                                         expected_usd=100, deducted_usd=50)])
        with patch.object(module.frappe, "get_all", return_value=[]), patch.object(module, "_", side_effect=lambda v: v):
            result = module._preview(document, period)
        self.assertEqual(result["rows"][0]["deducted_usd"], 30)
        self.assertEqual(result["rows"][0]["row_key"], "R")

    def test_generation_requires_fresh_preview_and_explicit_replacement(self):
        document = self.document()
        preview = {"fingerprint": "F", "rows": [{"client_name": "Ana", "deducted_usd": 10}],
                   "replaces_detail": True, "total_usd": 10}
        with patch.object(module.frappe, "db", SimpleNamespace(sql=Mock())), \
             patch.object(module, "_load", return_value=(document, SimpleNamespace(name="P"))), \
             patch.object(module, "_preview", return_value=preview), \
             patch.object(module, "_", side_effect=lambda v: v), \
             patch.object(module.frappe, "throw", side_effect=ValueError), \
             patch("frappe.utils.file_manager.save_file", return_value=SimpleNamespace(file_url="/private/files/generated.xlsx")) as save_file, \
             patch("credinomina_reconciliation.conciliacion_credinomina.doctype.cn_remittance_allocation.cn_remittance_allocation._apply_remittance_detail") as apply:
            with self.assertRaises(ValueError):
                module.use_application_detail("D", "STALE", True)
            with self.assertRaises(ValueError):
                module.use_application_detail("D", "F", False)
            save_file.assert_not_called()
            result = module.use_application_detail("D", "F", True)
        self.assertEqual(result["rows"], 1)
        self.assertEqual(document.result, "Pendiente")
        self.assertEqual(document.detail_file, "/private/files/generated.xlsx")
        self.assertEqual(save_file.call_args.kwargs["is_private"], 1)
        self.assertEqual(apply.call_args.kwargs["origin"], "Aplicaciones pendientes del período")
        document.add_comment.assert_called_once()


if __name__ == "__main__":
    unittest.main()
