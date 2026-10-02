import unittest

from credinomina_reconciliation.paying_employers import allowed_employers, reconciliation_companies, choose_detail_client
from credinomina_reconciliation.allocation import allocate_cash
from credinomina_reconciliation.remittance_detail import suggest_detail_targets, manual_detail_targets
from credinomina_reconciliation.application_deposit_detail import pending_application_rows
from credinomina_reconciliation.parsers import parse_collection_file
from credinomina_reconciliation.templates import build_template_xlsx


class PayingEmployersTests(unittest.TestCase):
    def test_direct_authorization_not_symmetric_or_transitive(self):
        mapping = {"INDENICSA": {"CBC"}, "CBC": {"OTRA"}}
        self.assertEqual(allowed_employers("INDENICSA", mapping), {"INDENICSA", "CBC"})
        self.assertEqual(allowed_employers("CBC", mapping), {"CBC", "OTRA"})
        self.assertEqual(reconciliation_companies("CBC", mapping), ["CBC", "INDENICSA", "OTRA"])

    def test_one_deposit_two_companies_keeps_cash_and_balances_separate(self):
        deposits = [{"id": "D", "group": "INDENICSA", "allowed_groups": ["INDENICSA", "CBC"], "amount_usd": 1000}]
        claims = [{"id": "A", "group": "INDENICSA", "amount_usd": 700}, {"id": "B", "group": "CBC", "amount_usd": 500}]
        instructions = [{"id": "1", "deposit_id": "D", "claim_id": "A", "amount_usd": 700},
                        {"id": "2", "deposit_id": "D", "claim_id": "B", "amount_usd": 300}]
        result = allocate_cash(deposits, claims, instructions)
        self.assertEqual(result["deposit_remaining"], {"D": 0})
        self.assertEqual(result["claim_remaining"], {"A": 0, "B": 200})
        self.assertEqual(sum(row["amount_usd"] for row in result["allocations"]), 1000)

    def test_forged_destination_is_rejected_even_in_positive_manual_engine(self):
        result = allocate_cash([{"id": "D", "group": "A", "allowed_groups": ["A", "B"], "amount_usd": 100}],
            [{"id": "C", "group": "C", "amount_usd": 100}],
            [{"id": "T", "deposit_id": "D", "claim_id": "C", "amount_usd": 100}])
        self.assertEqual(result["allocations"], [])
        self.assertEqual(result["instruction_results"]["T"], "Empresa no autorizada")

    def test_identification_is_scoped_and_ambiguous_names_need_company(self):
        clients = [{"name": "1", "client_name": "Ana Ruiz", "employer": "A", "client_number": "1"},
                   {"name": "2", "client_name": "Ruiz Ana", "employer": "B", "client_number": "2"},
                   {"name": "3", "client_name": "Otro", "employer": "C", "client_number": "3"}]
        client, reason = choose_detail_client({"client_name": "Ana Ruiz"}, clients, "A", {"A", "B"})
        self.assertIsNone(client)
        self.assertIn("ambiguo", reason)
        client, _ = choose_detail_client({"client_name": "Ana Ruiz", "employer": "B"}, clients, "A", {"A", "B"})
        self.assertEqual(client["name"], "2")
        client, _ = choose_detail_client({"client_number": "3"}, clients, "A", {"A", "B"})
        self.assertIsNone(client)

    def test_detail_can_match_beneficiary_but_never_combine_ambiguous_companies(self):
        row = {"client_name": "Ana Ruiz"}
        claims = [{"id": "H:1", "group": "B", "kind": "H", "amount_usd": 100, "client_name": "Ana Ruiz"}]
        targets, _ = suggest_detail_targets(row, claims, 100, "A", allowed_groups=["A", "B"])
        self.assertEqual(targets[0]["claim_id"], "H:1")
        manual = [{"id": "T", "claim_id": "H:1", "amount_usd": 100}]
        self.assertTrue(manual_detail_targets(row, claims, manual, 100, "A", allowed_groups=["A", "B"])[0])
        claims.append({**claims[0], "id": "H:2", "group": "A"})
        self.assertFalse(suggest_detail_targets(row, claims, 200, "A", allowed_groups=["A", "B"])[0])
        self.assertTrue(suggest_detail_targets({**row, "employer": "B"}, claims, 100, "A", allowed_groups=["A", "B"])[0])

    def test_pending_consumes_payments_from_any_payer(self):
        rows = pending_application_rows([{"claim_id": "H:A", "historical_application": "A", "applied_usd": 500}],
            [{"name": "D", "employer": "OTHER", "docstatus": 1,
              "allocation_detail": '[{"aplicacion_id":"A", "importe_usd":300}]'}], [], "NEW")
        self.assertEqual(rows[0]["deducted_usd"], 200)

    def test_template_and_import_accept_optional_beneficiary(self):
        from openpyxl import load_workbook
        from io import BytesIO
        content = build_template_xlsx("deposito", [{"client_name": "Ana", "employer": "CBC"}])
        workbook = load_workbook(BytesIO(content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        sheet.cell(2, headers.index("Deducido US$") + 1, 100)
        output = BytesIO()
        workbook.save(output)
        self.assertEqual(parse_collection_file("detail.xlsx", output.getvalue(), require_deduction=True)[0]["employer"], "CBC")
