"""Templates must be directly importable after the operator fills their rows."""

import io
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from credinomina_reconciliation.parsers import parse_collection_file
from credinomina_reconciliation.template_download import download_import_template
from credinomina_reconciliation.templates import (
    COLLECTION_HEADERS,
    DETAIL_HEADERS,
    build_template_xlsx,
)


def filled_content(content, values):
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    sheet.append([values.get(header) for header in headers])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class TemplateTests(unittest.TestCase):
    def test_collection_template_has_importable_headers_and_no_fake_clients(self):
        content = build_template_xlsx("cobranza")
        workbook = load_workbook(io.BytesIO(content))
        self.assertEqual(workbook.sheetnames, ["Cobranza"])
        sheet = workbook.active
        self.assertEqual(tuple(cell.value for cell in sheet[1]), COLLECTION_HEADERS)
        self.assertEqual(sheet.max_row, 1)
        self.assertIsNotNone(sheet["C1"].comment)
        self.assertEqual(sheet.column_dimensions["M"].number_format, "@")
        content = filled_content(content, {
            "Nro. Cliente": "001", "Nombre y Apellidos del Cliente": "ANA PÉREZ",
            "Nro. Crédito": "0900", "Monto de la cuota en US$": 46.52,
        })
        rows = parse_collection_file("cobranza.xlsx", content, require_name=True)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["client_number"], "001")
        self.assertEqual(rows[0]["loan_number"], "0900")
        self.assertEqual(rows[0]["expected_usd"], 46.52)

    def test_company_template_prefills_collection_identity_not_deductions(self):
        content = build_template_xlsx("empresa", [{
            "client_number": "001", "client_name": "ANA PÉREZ",
            "loan_number": "0900", "expected_usd": 46.52,
            "row_key": "FILA-1",
        }])
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook.active
        self.assertEqual(tuple(cell.value for cell in sheet[1]), DETAIL_HEADERS)
        self.assertEqual(sheet["C2"].value, "ANA PÉREZ")
        self.assertEqual(sheet["O2"].value, "FILA-1")
        self.assertIsNone(sheet["M2"].value)
        self.assertIsNone(sheet["N2"].value)
        sheet["N2"] = 46.52
        stream = io.BytesIO()
        workbook.save(stream)
        rows = parse_collection_file(
            "detalle_empresa.xlsx", stream.getvalue(),
            require_deduction=True, keep_zero_rows=True, require_name=True,
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["row_key"], "FILA-1")
        self.assertEqual(rows[0]["deducted_usd"], 46.52)

    def test_deposit_template_accepts_name_only_identity(self):
        content = build_template_xlsx("deposito")
        content = filled_content(content, {
            "Nombre y Apellidos del Cliente": "LUIS GÓMEZ",
            "Deducido C$": 3650,
        })
        rows = parse_collection_file(
            "detalle_deposito.xlsx", content,
            require_deduction=True, keep_zero_rows=True, require_name=True,
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["client_name"], "LUIS GÓMEZ")
        self.assertEqual(rows[0]["deducted_nio"], 3650)

    def test_invalid_template_type_is_rejected(self):
        with self.assertRaises(ValueError):
            build_template_xlsx("transacciones")

    def test_buttons_exist_on_both_forms(self):
        root = Path(__file__).resolve().parents[1]
        doctype = root / "credinomina_reconciliation" / "conciliacion_credinomina" / "doctype"
        period = (doctype / "cn_reconciliation_period" / "cn_reconciliation_period.js").read_text(encoding="utf-8")
        remittance = (doctype / "cn_remittance_allocation" / "cn_remittance_allocation.js").read_text(encoding="utf-8")
        self.assertIn("Plantilla de cobranza", period)
        self.assertIn("Plantilla de detalle empresa", period)
        self.assertIn("Plantilla de detalle del depósito", remittance)
        self.assertIn("template_download.download_import_template", period)
        self.assertIn("template_download.download_import_template", remittance)

    def test_download_requires_permission_and_checks_prefilled_period(self):
        with (
            patch("frappe.has_permission", return_value=False),
            patch("frappe.throw", side_effect=PermissionError("denied")),
        ):
            with self.assertRaises(PermissionError):
                download_import_template("deposito")

        response = SimpleNamespace()
        period = SimpleNamespace(
            collection_rows=[SimpleNamespace(as_dict=lambda: {
                "client_name": "ANA PÉREZ", "loan_number": "0900", "row_key": "FILA-1",
            })],
            check_permission=Mock(),
        )
        with (
            patch("frappe.has_permission", return_value=True),
            patch("frappe.get_doc", return_value=period) as get_doc,
            patch("frappe.local", new=SimpleNamespace(response=response)),
        ):
            download_import_template("deposito", "CN-PER-1")
        get_doc.assert_called_once_with("CN Reconciliation Period", "CN-PER-1")
        period.check_permission.assert_called_once_with("read")
        self.assertEqual(response.type, "download")
        self.assertEqual(response.filename, "plantilla_detalle_deposito.xlsx")
        sheet = load_workbook(io.BytesIO(response.filecontent)).active
        self.assertEqual(sheet["O2"].value, "FILA-1")


if __name__ == "__main__":
    unittest.main()
