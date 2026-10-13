"""The follow-up export is complete, filtered, and safe to open in Excel."""
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch

import frappe
from openpyxl import load_workbook

from credinomina_reconciliation.conciliacion_credinomina.page.control_credinomina import control_credinomina as page


class ControlWorkExportTests(TestCase):
    def test_export_uses_queue_filters_and_contains_all_matching_cases(self):
        response = frappe._dict()
        queue = [
            {"priority": 0, "kind": "credit_management", "employer": "A", "employer_name": "AIRTEC S A",
             "period_label": "Sin período", "summary": "Saldo a favor del cliente por gestionar",
             "client_name": "=1+1", "client_number": "123", "loan_number": "109-1", "amount_usd": 19.25,
             "amount_kind": "Pendiente de devolver/aplicar", "event_date": "2025-04-01", "due_date": "2025-05-01",
             "responsible": "ana@example.com", "next_action": "Documentar devolución", "target_doctype": "CN Complementary Item",
             "target_name": "CN-COMP-1"},
            {"priority": 2, "kind": "open_exception", "employer": "A", "employer_name": "AIRTEC S A",
             "period_label": "2025-04", "summary": "Excepción abierta", "client_name": "Luis",
             "amount_usd": 5, "event_date": "2025-04-02", "target_doctype": "CN Reconciliation Exception",
             "target_name": "CN-EXC-1"},
        ]
        with patch.object(page, "_work_consultation", return_value=({"work_items": queue}, {"consultation": "snapshot"})) as consult, \
             patch.object(page, "now_datetime", return_value=datetime(2026, 10, 12, 10, 30)), \
             patch.object(frappe, "local", SimpleNamespace(response=response, flags=frappe._dict(in_test=False))):
            page.export_work_items_excel(employer="A", work_kind="credits", responsible="ana",
                due="", search="", min_amount="10", sort_by="amount", consultation="snapshot")

        consult.assert_called_once_with("snapshot", "A")
        self.assertEqual(response.type, "download")
        self.assertIn("AIRTEC-S-A", response.filename)
        book = load_workbook(BytesIO(response.filecontent), data_only=False)
        sheet = book["Qué falta hacer"]
        self.assertEqual(sheet.max_row, 6)
        self.assertEqual(sheet.cell(6, 2).value, "AIRTEC S A")
        self.assertEqual(sheet.cell(6, 6).value, "'=1+1")
        self.assertEqual(sheet.cell(6, 9).value, 19.25)
        self.assertEqual(sheet.cell(6, 15).value, "CN-COMP-1")
        self.assertEqual(sheet.freeze_panes, "A6")
        self.assertEqual(sheet.auto_filter.ref, "A5:Q6")
