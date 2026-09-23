import io
import unittest

from openpyxl import Workbook

from credinomina_reconciliation.allocation import allocate_cash, can_document_surplus
from credinomina_reconciliation.parsers import (
    SOURCE_ACCOUNTING,
    SOURCE_DEPOSITS,
    SourceFileError,
    parse_collection_file,
    parse_source_file,
)
from credinomina_reconciliation.reconciliation import (
    classify_deduction,
    complementary_matches_collection,
    deposit_pair_result,
    match_collection_record,
    matching_exception_notes,
    settlement_result,
)


def workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class CollectionParserTest(unittest.TestCase):
    def test_reads_repeated_spanish_header_blocks(self):
        headers = [
            "Nro. Cliente",
            "Nombre y Apellidos del Cliente",
            "Nro Cédula",
            "Nro. Crédito",
            "Nro. cuota",
            "Nro. de cuotas totales",
            "Monto de la cuota en US$",
            "Monto de la cuota en C$",
            "Comentarios",
            "Referencia de Aplicación",
            "Comentario de Aplicación",
            "Deducido C$",
            "Deducido US$",
            "Fila ID",
        ]
        content = workbook_bytes(
            [
                ["DETALLE DE COBRO"],
                headers,
                [100, "ANA", "001A", 900, 3, 12, 10, 366, "", "R1", "", 366, 10, "F-1"],
                ["OTRA EMPRESA"],
                headers,
                [101, "LUIS", "002B", 901, 1, 24, 20, 732, "", "R2", "", 0, 0, "F-2"],
            ]
        )
        rows = parse_collection_file("cobranza.xlsx", content, require_deduction=True)
        self.assertEqual(2, len(rows))
        self.assertEqual("900", rows[0]["loan_number"])
        self.assertEqual(366, rows[0]["deducted_nio"])
        self.assertEqual("F-2", rows[1]["row_key"])

    def test_remittance_detail_keeps_a_zero_deduction_for_review(self):
        content = workbook_bytes([
            ["Nro. Cliente", "Nro. Crédito", "Deducido US$"],
            [100, 900, 0],
        ])
        rows = parse_collection_file(
            "detalle.xlsx", content, require_deduction=True, keep_zero_rows=True,
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["deducted_usd"], 0)

    def test_exact_matching_never_uses_name_similarity(self):
        response = {
            "client_number": "100",
            "loan_number": "900",
            "installment_number": "3",
            "client_name": "NOMBRE DIFERENTE",
        }
        candidates = [
            {
                "name": "ROW-1",
                "client_number": "100",
                "loan_number": "900",
                "installment_number": "3",
                "client_name": "ANA",
            }
        ]
        match, reason = match_collection_record(response, candidates)
        self.assertEqual("ROW-1", match["name"])
        self.assertIn("cliente", reason)

    def test_ambiguous_loan_is_not_auto_assigned(self):
        response = {"loan_number": "900"}
        candidates = [
            {"name": "A", "loan_number": "900"},
            {"name": "B", "loan_number": "900"},
        ]
        match, reason = match_collection_record(response, candidates)
        self.assertIsNone(match)
        self.assertIn("ambigua", reason)

    def test_numeric_ids_match_when_excel_drops_leading_zeroes(self):
        response = {
            "client_number": "00100",
            "loan_number": "00900",
            "installment_number": "03",
        }
        candidates = [
            {
                "name": "ROW-1",
                "client_number": "100",
                "loan_number": "900",
                "installment_number": "3",
            }
        ]
        match, _reason = match_collection_record(response, candidates)
        self.assertEqual("ROW-1", match["name"])


class SourceParserTest(unittest.TestCase):
    def test_accounting_parser_keeps_native_deposit_currency(self):
        content = workbook_bytes(
            [
                ["MIDESA"],
                ["MOVIMIENTOS"],
                [
                    "CUENTA_CONTABLE",
                    "DESCRIPCION_CTA_CONTABLE",
                    "FECHA_APLICA",
                    "NO_CMPTE",
                    "NO_REF",
                    "DESCRIPCION",
                    "DEBITO_DEL_MES",
                    "CREDITO_DEL_MES",
                ],
                [
                    "1602",
                    "Creditos M.E.",
                    "02/01/2026",
                    "001",
                    "321001591",
                    "C$1,711.00 DEPOSITO POR: CONVENIO ACME EN LA CUENTA BAC",
                    0,
                    46.72,
                ],
                [
                    "1602",
                    "Creditos M.E.",
                    "02/01/2026",
                    "002",
                    "002",
                    "NOTA AL PRESTAMO 109367-1 NO. DOCUM 7541 CLIENTE ANA PEREZ N.C. CONVENIO ACME |REF:321001591|",
                    46.72,
                    0,
                ],
            ]
        )
        rows = parse_source_file(SOURCE_ACCOUNTING, "movimientos.xlsx", content)
        self.assertEqual(2, len(rows))
        self.assertEqual("NIO", rows[0]["currency"])
        self.assertEqual(1711, rows[0]["amount"])
        self.assertEqual("USD", rows[0]["equivalent_currency"])
        self.assertEqual(46.72, rows[0]["equivalent_amount"])
        self.assertAlmostEqual(1711 / 46.72, rows[0]["fx_rate"], places=6)
        self.assertEqual("USD", rows[1]["currency"])
        self.assertEqual("109367", rows[1]["loan_number"])
        self.assertEqual("321001591", rows[1]["reference"])
        self.assertEqual("ANA PEREZ", rows[1]["client_name"])
        self.assertEqual("002", rows[1]["accounting_entry"])
        self.assertEqual("7541", rows[1]["receipt"])

    def test_bank_detail_preserves_nio_payment_and_usd_equivalent(self):
        content = workbook_bytes(
            [
                ["Fecha", "REFERENCIA", "ValorC", "ValorU", "MONTO U$"],
                ["30/04/2026", "R-1", 3660, 0, 100],
            ]
        )
        rows = parse_source_file(SOURCE_DEPOSITS, "depositos.xlsx", content)
        self.assertEqual(1, len(rows))
        self.assertEqual("NIO", rows[0]["currency"])
        self.assertEqual(3660, rows[0]["amount"])
        self.assertEqual(100, rows[0]["equivalent_amount"])
        self.assertEqual(36.6, rows[0]["fx_rate"])

    def test_monthly_bank_file_uses_deposito_tab_and_keeps_mixed_rows(self):
        workbook = Workbook()
        workbook.active.title = "BD"
        workbook.active.append(["REFERENCIA", "ValorC", "ValorU"])
        workbook.active.append(["WRONG", 9999, 0])
        sheet = workbook.create_sheet("Depósito")
        sheet.append([
            "Fecha", "REFERENCIA ", "DESCRIPCION", "BANCO", " ValorC ",
            " ValorU ", "No.Credito", "CLIENTE", "MONTO U$",
        ])
        sheet.append(["15/04/2025", "R-1", "Remesa", "BAC C$", 3660, 0, None,
                      "CONVENIO ACME", 100])
        sheet.append(["16/04/2025", "R-2", "Pago personal", "BAC U$", 0, 25,
                      "1", "Cliente personal", 25])
        stream = io.BytesIO()
        workbook.save(stream)

        rows = parse_source_file(SOURCE_DEPOSITS, "4_ABRIL 2025.xlsx", stream.getvalue())
        self.assertEqual(2, len(rows))
        self.assertEqual(["R-1", "R-2"], [row["reference"] for row in rows])
        self.assertEqual([2, 3], [row["source_row"] for row in rows])
        self.assertEqual("ACME", rows[0]["employer_text"])
        self.assertEqual("NIO", rows[0]["currency"])
        self.assertEqual("USD", rows[1]["currency"])

    def test_multisheet_bank_file_without_deposito_tab_is_rejected(self):
        workbook = Workbook()
        workbook.active.title = "BD"
        workbook.create_sheet("Otra")
        stream = io.BytesIO()
        workbook.save(stream)
        with self.assertRaisesRegex(SourceFileError, "pestaña Depósito"):
            parse_source_file(SOURCE_DEPOSITS, "depositos.xlsx", stream.getvalue())


class ReconciliationTest(unittest.TestCase):
    def test_application_exception_is_reused_for_equal_payment_shortfall(self):
        notes = matching_exception_notes(
            expected_usd=100, deducted_usd=100, applied_usd=90,
            complementary_usd=0, remitted_usd=90,
            application_notes=[{"comment": "Subsidio informado", "exception_id": "EX-1"}],
        )
        self.assertEqual(1, len(notes))
        self.assertEqual("Cobranza vs aplicacion", notes[0]["origin"])
        self.assertEqual(10, notes[0]["gap_usd"])
        self.assertEqual("EX-1", notes[0]["exception_id"])

    def test_deduction_exception_is_reused_for_equal_payment_shortfall(self):
        notes = matching_exception_notes(
            expected_usd=100, deducted_usd=90, applied_usd=90,
            complementary_usd=0, remitted_usd=90,
            deduction_notes=[{"comment": "Ingreso insuficiente"}],
        )
        self.assertEqual("Cobranza vs deduccion", notes[0]["origin"])
        self.assertEqual(10, notes[0]["gap_usd"])

    def test_different_or_absent_payment_does_not_inherit_comment(self):
        common = dict(
            expected_usd=100, deducted_usd=100, applied_usd=90,
            complementary_usd=0,
            application_notes=[{"comment": "Aplicacion parcial"}],
        )
        self.assertEqual([], matching_exception_notes(**common, remitted_usd=80))
        self.assertEqual([], matching_exception_notes(**common, remitted_usd=0))

    def test_complementary_item_and_duplicate_comments_do_not_create_new_gap(self):
        common = dict(
            expected_usd=100, deducted_usd=90, applied_usd=90,
            remitted_usd=90,
            application_notes=[{"comment": "Deduccion parcial"}],
            deduction_notes=[{"comment": "Deduccion parcial"}],
        )
        self.assertEqual(1, len(matching_exception_notes(**common, complementary_usd=0)))
        self.assertEqual(1, len(matching_exception_notes(**common, complementary_usd=10)))
        self.assertEqual(
            [], matching_exception_notes(
                expected_usd=100, deducted_usd=100, applied_usd=90,
                complementary_usd=10, remitted_usd=100,
                application_notes=[{"comment": "Administrativo"}],
            ),
        )

    def test_partial_and_full_deductions(self):
        self.assertEqual(
            "Deduccion total",
            classify_deduction(expected_usd=10, deducted_usd=10),
        )
        self.assertEqual(
            "Deduccion parcial",
            classify_deduction(expected_nio=1000, deducted_nio=600),
        )
        self.assertEqual(
            "No deducido",
            classify_deduction(expected_nio=1000, deducted_nio=0),
        )

    def test_nio_deposit_settles_two_usd_applications(self):
        accounting = {
            "reference": "REF-1",
            "currency": "NIO",
            "amount": 3660,
            "equivalent_currency": "USD",
            "equivalent_amount": 100,
            "fx_basis": "Movimiento contable",
        }
        bank = {
            "reference": "REF-1",
            "currency": "NIO",
            "amount": 3660,
            "equivalent_currency": "USD",
            "equivalent_amount": 100,
            "fx_basis": "Detalle de deposito",
        }
        self.assertTrue(deposit_pair_result(accounting, bank)[0])
        applications = [
            {"currency": "USD", "amount": 40},
            {"currency": "USD", "amount": 60},
        ]
        result = settlement_result(applications, [(accounting, bank)])
        self.assertTrue(result["matched"])
        self.assertEqual("USD", result["currency"])
        self.assertEqual(100, result["deposited"])

    def test_different_deposit_currency_uses_documented_rate(self):
        accounting = {
            "reference": "REF-2",
            "currency": "NIO",
            "amount": 3660,
            "equivalent_currency": "USD",
            "equivalent_amount": 100,
            "fx_basis": "Movimiento contable",
        }
        bank = {"reference": "REF-2", "currency": "USD", "amount": 100}
        self.assertTrue(deposit_pair_result(accounting, bank)[0])
        self.assertTrue(
            settlement_result([{"currency": "USD", "amount": 100}], [(accounting, bank)])["matched"]
        )

    def test_missing_rate_or_amount_difference_stays_unreconciled(self):
        accounting = {"reference": "REF-3", "currency": "NIO", "amount": 3660}
        bank = {"reference": "REF-3", "currency": "USD", "amount": 100}
        self.assertFalse(deposit_pair_result(accounting, bank)[0])
        accounting["manual_fx_rate"] = 36.6
        accounting["manual_fx_evidence"] = "Constancia bancaria"
        self.assertTrue(deposit_pair_result(accounting, bank)[0])
        result = settlement_result(
            [{"currency": "USD", "amount": 98}], [(accounting, bank)]
        )
        self.assertFalse(result["matched"])
        self.assertEqual(2, result["difference"])

    def test_separate_administrative_entry_completes_deposit_not_loan_payment(self):
        accounting = {"reference": "R-100", "currency": "USD", "amount": 100}
        bank = {"reference": "R-100", "currency": "USD", "amount": 100}
        applications = [{"reference": "R-100", "currency": "USD", "amount": 90}]
        self.assertFalse(settlement_result(applications, [(accounting, bank)])["matched"])
        result = settlement_result(
            applications,
            [(accounting, bank)],
            [{"reference": "R-100", "amount_usd": 10}],
        )
        self.assertTrue(result["matched"])
        self.assertEqual(90, result["applied"])
        self.assertEqual(10, result["complementary"])

    def test_cent_difference_needs_explicit_reconciliation_movement(self):
        for applied, deposited in ((46.52, 46.53), (46.53, 46.52)):
            with self.subTest(applied=applied, deposited=deposited):
                accounting = {"reference": "R-CENT", "currency": "USD", "amount": deposited}
                bank = dict(accounting)
                result = settlement_result(
                    [{"currency": "USD", "amount": applied}],
                    [(accounting, bank)],
                )
                self.assertFalse(result["matched"])
                self.assertEqual(round(deposited - applied, 4), result["difference"])

    def test_complementary_entry_needs_explicit_unique_loan_allocation(self):
        target = {
            "parent": "PER-APR",
            "employer": "ACME",
            "client_number": "100",
            "loan_number": "900",
            "installment_number": "3",
        }
        self.assertFalse(
            complementary_matches_collection(
                {"reference": "R-100", "amount_usd": 10}, target, "R-100"
            )
        )
        self.assertTrue(
            complementary_matches_collection(
                {"reference": "R-100", "loan_number": "900", "period": "PER-APR"},
                target,
                "R-100",
            )
        )
        self.assertFalse(
            complementary_matches_collection(
                {"reference": "OTHER", "loan_number": "900"}, target, "R-100"
            )
        )


class CashAllocationTest(unittest.TestCase):
    def test_one_deposit_pays_collection_partially(self):
        result = allocate_cash(
            [{"id": "D1", "reference": "R1", "amount_usd": 40}],
            [{"id": "C1", "references": ["R1"], "amount_usd": 100}],
        )
        self.assertEqual(40, result["allocations"][0]["amount_usd"])
        self.assertEqual(60, result["claim_remaining"]["C1"])
        self.assertEqual(0, result["deposit_remaining"]["D1"])

    def test_one_deposit_pays_several_collections(self):
        result = allocate_cash(
            [{"id": "D1", "reference": "R1", "amount_usd": 100}],
            [
                {"id": "C1", "references": ["R1"], "amount_usd": 40},
                {"id": "C2", "references": ["R1"], "amount_usd": 60},
            ],
        )
        self.assertEqual(2, len(result["allocations"]))
        self.assertEqual(0, result["deposit_remaining"]["D1"])

    def test_same_reference_across_employers_does_not_auto_split(self):
        result = allocate_cash(
            [{"id": "D1", "reference": "R1", "amount_usd": 100}],
            [
                {"id": "C1", "references": ["R1"], "amount_usd": 40, "group": "ACME"},
                {"id": "C2", "references": ["R1"], "amount_usd": 60, "group": "BETA"},
            ],
        )
        self.assertEqual([], result["allocations"])

    def test_known_deposit_employer_must_match_collection_employer(self):
        result = allocate_cash(
            [{"id": "D1", "reference": "R1", "amount_usd": 40, "group": "ACME"}],
            [{"id": "C1", "references": ["R1"], "amount_usd": 40, "group": "BETA"}],
        )
        self.assertEqual([], result["allocations"])

    def test_several_deposits_pay_one_collection(self):
        result = allocate_cash(
            [
                {"id": "D1", "reference": "R1", "amount_usd": 40},
                {"id": "D2", "reference": "R1", "amount_usd": 60},
            ],
            [{"id": "C1", "references": ["R1"], "amount_usd": 100}],
        )
        self.assertEqual(2, len(result["allocations"]))
        self.assertEqual(0, result["claim_remaining"]["C1"])

    def test_one_deposit_can_pay_several_collections_partially_with_exact_application_hints(self):
        result = allocate_cash(
            [{"id": "D1", "reference": "R1", "amount_usd": 50}],
            [
                {"id": "C1", "references": ["R1"], "amount_usd": 100, "hints": {"R1": 20}},
                {"id": "C2", "references": ["R1"], "amount_usd": 100, "hints": {"R1": 30}},
            ],
        )
        self.assertEqual([20, 30], [row["amount_usd"] for row in result["allocations"]])
        self.assertEqual(80, result["claim_remaining"]["C1"])
        self.assertEqual(70, result["claim_remaining"]["C2"])

    def test_different_deposit_references_can_fund_one_collection_when_approved(self):
        result = allocate_cash(
            [
                {"id": "D1", "reference": "R1", "amount_usd": 40},
                {"id": "D2", "reference": "R2", "amount_usd": 60},
            ],
            [{"id": "C1", "references": [], "amount_usd": 100}],
            [
                {"id": "A1", "deposit_id": "D1", "claim_id": "C1", "amount_usd": 40},
                {"id": "A2", "deposit_id": "D2", "claim_id": "C1", "amount_usd": 60},
            ],
        )
        self.assertEqual(0, result["claim_remaining"]["C1"])
        self.assertEqual(2, len(result["allocations"]))

    def test_ambiguous_split_requires_manual_instruction(self):
        deposits = [{"id": "D1", "reference": "R1", "amount_usd": 50}]
        claims = [
            {"id": "C1", "references": ["R1"], "amount_usd": 50},
            {"id": "C2", "references": ["R1"], "amount_usd": 50},
        ]
        self.assertEqual([], allocate_cash(deposits, claims)["allocations"])
        result = allocate_cash(
            deposits, claims,
            [{"id": "A1", "deposit_id": "D1", "claim_id": "C2", "amount_usd": 50}],
        )
        self.assertEqual("Aplicada", result["instruction_results"]["A1"])
        self.assertEqual("C2", result["allocations"][0]["claim_id"])

    def test_allocation_cannot_exceed_deposit_or_collection(self):
        result = allocate_cash(
            [{"id": "D1", "reference": "R1", "amount_usd": 50}],
            [{"id": "C1", "references": ["R1"], "amount_usd": 40}],
            [{"id": "A1", "deposit_id": "D1", "claim_id": "C1", "amount_usd": 50}],
        )
        self.assertEqual("Excede cobranza", result["instruction_results"]["A1"])
        self.assertEqual([], result["allocations"])

    def test_company_overpayment_remains_unallocated(self):
        result = allocate_cash(
            [{"id": "D1", "reference": "R1", "amount_usd": 120}],
            [{"id": "C1", "references": ["R1"], "amount_usd": 100}],
        )
        self.assertEqual(100, result["allocations"][0]["amount_usd"])
        self.assertEqual(20, result["deposit_remaining"]["D1"])
        self.assertTrue(can_document_surplus(20, 0, 20))
        self.assertFalse(can_document_surplus(20, 15, 10))


if __name__ == "__main__":
    unittest.main()
