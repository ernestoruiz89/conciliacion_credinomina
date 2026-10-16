"""The Excel summary shares the company report, including indeterminate balances."""
import io
from datetime import datetime
from unittest import TestCase
from unittest.mock import patch

from openpyxl import load_workbook

from credinomina_reconciliation.company_statement import METRICS
from credinomina_reconciliation.control_export import build_control_workbook, MONEY_FORMAT
from credinomina_reconciliation.conciliacion_credinomina.report.estado_de_cuenta_por_empresa import estado_de_cuenta_por_empresa as report


class CompanySummaryExportTests(TestCase):
    def test_report_values_and_columns_are_preserved(self):
        # These positions need no reconciliation period to appear in the summary.
        positions = []
        for field, value in zip((field for field, _ in METRICS),
                                [4474.07, -7.35, -0.18, -154.28, -3173.81, 0.01]):
            positions.append(dict(employer='REPSA', balance_usd=value,
                                  **{key: value if key == field else 0 for key, _ in METRICS}))
        positions.append(dict(employer='Sin importe', balance_usd=None,
                              **{key: None if key == 'core_pending_usd' else 0 for key, _ in METRICS}))
        with patch.object(report, 'load_detail', return_value=positions), patch.object(report, 'load_core_detail', return_value=[]):
            columns, rows, message = report.execute({'view_mode': 'Resumen'})
        payload = build_control_workbook(
            {'year': 'Todos', 'periods': [],
             'company_statement': {'columns': columns, 'rows': rows, 'message': message}},
            exceptions=[], actions=[], employer_label='Todas las empresas',
            generated_at=datetime(2026, 10, 4))
        book = load_workbook(io.BytesIO(payload))
        self.assertEqual(book.sheetnames[:2], ['Resumen', 'Resumen mensual'])
        sheet = book['Resumen']
        self.assertEqual([cell.value for cell in sheet[4]], [column['label'] for column in columns])
        by_field = {column['fieldname']: index for index, column in enumerate(columns, 1)}
        for field, expected in {'core_balance_usd': 0, 'client_credit_usd': 154.28,
                                'company_credit_usd': .18, 'company_receivable_usd': .01,
                                'application_pending_usd': 4474.07, 'deposit_pending_usd': 3173.81}.items():
            self.assertEqual(sheet.cell(5, by_field[field]).value, expected)
            self.assertEqual(sheet.cell(5, by_field[field]).number_format, MONEY_FORMAT)
        self.assertEqual(sheet.cell(6, by_field['core_pending_usd']).value, 'N/D')
        self.assertEqual(sheet.cell(6, by_field['core_balance_usd']).value, 0)
        self.assertNotIn('balance_usd', by_field)
        from openpyxl.utils import get_column_letter
        self.assertEqual(sheet.auto_filter.ref, f'A4:{get_column_letter(len(columns))}6')
        self.assertEqual(sheet['A9'].value, message)

    def test_empty_summary_retains_report_headers(self):
        payload = build_control_workbook(
            {'year': 2025, 'periods': []}, exceptions=[], actions=[],
            employer_label='Todas las empresas', generated_at=datetime(2026, 10, 4))
        sheet = load_workbook(io.BytesIO(payload))['Resumen']
        self.assertEqual([cell.value for cell in sheet[4]],
                         [column['label'] for column in report.get_columns()])
        from openpyxl.utils import get_column_letter
        self.assertEqual(sheet.auto_filter.ref, f'A4:{get_column_letter(len(report.get_columns()))}4')
