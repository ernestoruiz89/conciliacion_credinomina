"""Deposit templates and stored details omit collection-only amount columns."""

import json
import io
import unittest
from pathlib import Path

from openpyxl import load_workbook

from credinomina_reconciliation.templates import (
    DEPOSIT_HEADERS,
    DETAIL_HEADERS,
    build_template_xlsx,
)


ROOT = Path(__file__).resolve().parents[1]


class DepositDetailColumnTests(unittest.TestCase):
    def test_deposit_template_omits_total_installments_and_installment_amounts(self):
        omitted = {
            "Nro. cuota",
            "Referencia de Aplicación",
            "Comentario de Aplicación",
            "Nro. de cuotas totales",
            "Monto de la cuota en US$",
            "Monto de la cuota en C$",
        }
        self.assertFalse(omitted & set(DEPOSIT_HEADERS))
        self.assertTrue(omitted <= set(DETAIL_HEADERS))

        sheet = load_workbook(
            io.BytesIO(build_template_xlsx("deposito")), read_only=True
        ).active
        self.assertEqual(tuple(cell.value for cell in sheet[1]), DEPOSIT_HEADERS)

    def test_deposit_template_prefills_identity_without_reintroducing_omitted_fields(self):
        sheet = load_workbook(io.BytesIO(build_template_xlsx("deposito", [{
            "client_number": "001",
            "client_name": "ANA PÉREZ",
            "loan_number": "0900",
            "total_installments": "12",
            "expected_usd": 46.52,
            "expected_nio": 1700,
            "installment_number": "3",
            "application_reference": "REF-OMITIDA",
            "application_comment": "Comentario omitido",
            "comments": "Observación del depósito",
            "row_key": "FILA-1",
        }])), read_only=True).active
        self.assertEqual("ANA PÉREZ", sheet["C2"].value)
        self.assertEqual(sheet.max_column, 10)
        self.assertEqual(sheet["F2"].value, "Observación del depósito")
        self.assertIsNone(sheet["G2"].value)
        self.assertIsNone(sheet["H2"].value)
        self.assertEqual("FILA-1", sheet["I2"].value)

    def test_deposit_template_keeps_widths_formats_and_filter_aligned(self):
        sheet = load_workbook(io.BytesIO(build_template_xlsx("deposito"))).active
        self.assertEqual(sheet.auto_filter.ref, "A1:J2")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.column_dimensions["F"].width, 32)
        self.assertEqual(sheet.column_dimensions["I"].number_format, "@")
        for column in ("G", "H"):
            self.assertNotEqual(sheet.column_dimensions[column].number_format, "@")

    def test_detail_rows_hide_legacy_expected_amounts_and_import_no_longer_assigns_them(self):
        doctype_dir = (
            ROOT / "credinomina_reconciliation" / "conciliacion_credinomina"
            / "doctype"
        )
        detail = json.loads(
            (doctype_dir / "cn_remittance_detail" / "cn_remittance_detail.json")
            .read_text(encoding="utf-8")
        )
        fields = {field["fieldname"]: field for field in detail["fields"]}
        self.assertTrue(fields["expected_usd"]["hidden"])
        self.assertTrue(fields["expected_nio"]["hidden"])
        self.assertNotIn("total_installments", fields)

        source = (
            doctype_dir / "cn_remittance_allocation" / "cn_remittance_allocation.py"
        ).read_text(encoding="utf-8")
        self.assertNotIn('"expected_usd", "expected_nio"', source)


if __name__ == "__main__":
    unittest.main()
