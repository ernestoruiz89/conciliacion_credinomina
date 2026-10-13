"""Monetary evidence must agree across report counts, percentages and detail."""
import io
from datetime import datetime
from unittest.mock import patch

import frappe
import pytest
from openpyxl import load_workbook

from credinomina_reconciliation.company_statement import build_detail, summarize
from credinomina_reconciliation.complementary_balances import financial_balance
from credinomina_reconciliation.core_item_position import signed_pending
from credinomina_reconciliation.control_export import build_control_workbook
from credinomina_reconciliation.company_transactions import application_progress, application_progress_index, application_month_amounts
from credinomina_reconciliation.conciliacion_credinomina.report.transacciones_por_empresa import transacciones_por_empresa as report


def core_item(**values):
    return dict(name='CORE', employer='A', docstatus=1, category='Otros ingresos',
                amount_usd=100, source_debit=100, source_credit=0, source_currency='USD',
                accounting_source_key='ledger', source_date='2025-04-30',
                source_voucher='AS-1', source_description='Movimiento original') | values


@pytest.mark.parametrize('debit,credit', [(100, 0), (0, 100)])
@pytest.mark.parametrize('used', [110, -10])
def test_invalid_complement_distribution_cannot_turn_into_credit_or_receivable(debit, credit, used):
    item = core_item(source_debit=debit, source_credit=credit)
    balance = financial_balance(item, [dict(amount_usd=used)])
    item.update(_balance=balance, _signed_pending_usd=signed_pending(item, balance))
    assert balance['financial_status'] == 'Revisar distribución'
    assert item['_signed_pending_usd'] is None
    rows = build_detail([dict(employer='A', amount_usd=50, source_import='I'),
                         dict(employer='B', amount_usd=30, source_import='OTHER')], [item], [], [])
    observed = next(row for row in rows if row['position_type'] == 'Partida del core')
    assert observed['core_pending_usd'] is None
    assert observed['status'] == 'Revisar distribución'
    assert observed['observed_usd'] == 10
    assert 'saldo no determinable' in observed['observation']
    totals = {row['employer']: row for row in summarize(rows)}
    assert totals['A']['application_pending_usd'] == 50
    assert totals['A']['core_pending_usd'] is None
    assert totals['A']['balance_usd'] is None
    assert totals['B']['balance_usd'] == 30
    assert item['amount_usd'] == 100 and balance['used_usd'] == used


def test_invalid_core_remainder_is_blank_in_actionable_excel_not_a_negative_credit():
    item = core_item(accounting_status='Importada del core')
    item['_balance'] = financial_balance(item, [dict(amount_usd=110)])
    item['_signed_pending_usd'] = signed_pending(item, item['_balance'])
    data = dict(year=2025, periods=[], totals={}, core_complementary_items=[item])
    book = load_workbook(io.BytesIO(build_control_workbook(data, exceptions=[], actions=[],
        employer_label='A', generated_at=datetime(2026, 10, 12))), data_only=True)
    row = next(row for row in book['Partidas y excepciones'].iter_rows(min_row=5, values_only=True)
               if row[6] == item['name'])
    assert row[7] is None
    assert row[8] == 'Revisar distribución'
    assert item['_balance']['pending_usd'] == -10  # Original audit evidence is retained.


def application(**values):
    return frappe._dict(dict(name='ROW', parent='IMPORT', idx=1, event_type='Aplicacion',
        event_date='2025-04-30', currency='USD', amount=100, effective=1,
        match_status='Conciliado', historical_period='P', historical_remitted_usd=80,
        application_adjustment_usd=0, deposit_match_status='Depósito conciliado',
        client_number='1', client_name='Ana', loan_number='109277-1') | values)


@pytest.mark.parametrize('paid,adjustment,saved,state,covered,pending', [
    (0, 0, 'Depósito conciliado', 'Pendiente', 0, 100),
    (80, 0, 'Depósito conciliado', 'Parcial', 80, 20),
    (100, 0, 'Pendiente', 'Conciliado', 100, 0),
    (0, 100, 'Pendiente', 'Conciliado', 100, 0),
    (50, 20, 'Conciliada: depósito + ajuste', 'Parcial', 70, 30),
    (80, 20, 'Pendiente', 'Conciliado', 100, 0),
    (110, 0, 'Depósito conciliado', 'Revisar distribución', None, 0),
])
def test_report_filters_counts_percentage_and_detail_share_real_coverage(paid, adjustment, saved, state, covered, pending):
    parent = frappe._dict(name='IMPORT', employer='A', status='Importado', historical_period='P')
    source = application(historical_remitted_usd=paid, application_adjustment_usd=adjustment,
                         deposit_match_status=saved)
    with patch.object(frappe, 'has_permission', return_value=True), \
            patch.object(frappe, 'get_list', return_value=[parent]), \
            patch.object(frappe, 'get_all', return_value=[source]):
        _, matrix, _, _, counts = report.execute(dict(year=2025, employer='A'))
        detail = report.get_month_detail(2025, 4, 'A', state=state)
        other = report.get_month_detail(2025, 4, 'A', state='Parcial' if state != 'Parcial' else 'Conciliado')
    assert matrix[0]['m04_state'] == state
    assert matrix[0]['m04_covered_usd'] == covered
    assert matrix[0]['m04_percentage'] == covered
    assert sum(card['value'] for card in counts) == 1
    assert next(card['value'] for card in counts if card['label'] == {
        'Pendiente': 'Pendientes', 'Parcial': 'Parcialmente conciliadas',
        'Conciliado': 'Conciliadas', 'Revisar distribución': 'Distribuciones inconsistentes'}[state]) == 1
    assert detail['filtered_count'] == 1 and other['filtered_count'] == 0
    assert detail['rows'][0]['state'] == state
    assert detail['rows'][0]['pending_usd'] == detail['totals']['pending_usd'] == pending
    assert detail['rows'][0]['assigned_usd'] == paid
    assert detail['totals']['net_usd'] == 100 - adjustment
    assert source.deposit_match_status == saved


def test_operative_saved_settled_label_does_not_skip_reading_collection_evidence():
    parent = frappe._dict(name='IMPORT', employer='A', status='Importado')
    source = application(historical_period='', historical_remitted_usd=0,
                         collection_row_id='C', processing_route='Operativa')
    collection = frappe._dict(name='C', parent='P', applied_usd=100,
                             remittance_detail='[{"importe_usd":80}]')
    def listed(doctype, **kwargs):
        return [parent] if doctype == 'CN Accounting Import' else ['P']
    def children(doctype, **kwargs):
        return [source] if doctype == 'CN Source Row' else [collection]
    with patch.object(frappe, 'has_permission', return_value=True), \
            patch.object(frappe, 'get_list', side_effect=listed), \
            patch.object(frappe, 'get_all', side_effect=children) as read:
        matrix = report.execute(dict(year=2025, employer='A'))[1][0]
        detail = report.get_month_detail(2025, 4, 'A', state='Parcial')
    assert any(call.args[0] == 'CN Collection Row' for call in read.call_args_list)
    assert matrix['m04_state'] == 'Parcial' and matrix['m04_percentage'] == 80
    assert detail['rows'][0]['pending_usd'] == 20


def test_nio_coverage_is_in_usd_and_drafts_do_not_consume_it():
    source = application(currency='NIO', amount=3662.43, manual_fx_rate=36.6243)
    parent = dict(status='Importado', employer='A')
    result = application_progress(source, parent)
    assert (result['original_usd'], result['assigned_usd'], result['pending_usd']) == (100, 80, 20)
    assert result['state'] == 'Parcial'
    source.update(manual_fx_rate=0)
    assert application_progress(source, parent)['covered_usd'] is None
    source.update(manual_fx_rate=36.6243)
    draft = application_progress(source, dict(parent, status='Borrador'))
    assert (draft['state'], draft['assigned_usd'], draft['pending_usd']) == ('Pendiente', 0, 100)
    assert application_month_amounts([source], {'IMPORT': dict(parent, status='Borrador')}, {})[('A','2025-04')]['covered_usd'] == 0


def test_wholly_paid_shared_claim_is_attributable_without_using_saved_states():
    parent = dict(status='Importado', employer='A')
    first = application(historical_period='', historical_remitted_usd=0, collection_row_id='C',
                        deposit_match_status='Pendiente', processing_route='Operativa')
    second = application(**dict(first, name='OTHER', event_date='2025-05-15'))
    sources = [first, second]
    collections = {'C': dict(applied_usd=200, remittance_detail='[{"importe_usd":200}]')}
    progress = application_progress_index(sources, {'IMPORT': parent}, collections)
    for result in progress.values():
        assert (result['state'], result['assigned_usd'], result['pending_usd']) == ('Conciliado', 100, 0)
    amounts = application_month_amounts(sources, {'IMPORT': parent}, collections, progress)
    assert amounts[('A', '2025-04')] == amounts[('A', '2025-05')] == dict(total_usd=100, covered_usd=100)
    assert sum(result['assigned_usd'] for result in progress.values()) == 200


def test_overlinked_claim_cannot_be_counted_twice_even_in_different_months():
    parent = dict(status='Importado', employer='A')
    first = application(historical_period='', historical_remitted_usd=0, collection_row_id='C', processing_route='Operativa')
    second = application(**dict(first, name='OTHER', event_date='2025-05-15'))
    collections = {'C': dict(applied_usd=100, remittance_detail='[{"importe_usd":100}]')}
    progress = application_progress_index([first, second], {'IMPORT': parent}, collections)
    assert all(result['state'] == 'Revisar distribución' and result['covered_usd'] is None for result in progress.values())
    amounts = application_month_amounts([first, second], {'IMPORT': parent}, collections, progress)
    assert amounts[('A', '2025-04')]['covered_usd'] is None
    assert amounts[('A', '2025-05')]['covered_usd'] is None


def test_month_detail_keeps_the_same_shared_claim_context_as_the_calendar():
    parent = frappe._dict(name='IMPORT', status='Importado', employer='A')
    first = application(historical_period='', historical_remitted_usd=0, collection_row_id='C', processing_route='Operativa')
    second = application(**dict(first, name='OTHER', event_date='2025-05-15'))
    collection = frappe._dict(name='C', parent='P', applied_usd=100, remittance_detail='[{"importe_usd":100}]')
    def listed(doctype, **kwargs):
        return [parent] if doctype == 'CN Accounting Import' else ['P']
    def children(doctype, **kwargs):
        return [first, second] if doctype == 'CN Source Row' else [collection]
    with patch.object(frappe, 'has_permission', return_value=True), \
            patch.object(frappe, 'get_list', side_effect=listed), \
            patch.object(frappe, 'get_all', side_effect=children):
        calendar = report.execute(dict(year=2025, employer='A'))[1][0]
        april = report.get_month_detail(2025, 4, 'A', state='Revisar distribución')
        may = report.get_month_detail(2025, 5, 'A', state='Revisar distribución')
    assert calendar['m04_state'] == calendar['m05_state'] == 'Revisar distribución'
    assert april['filtered_count'] == may['filtered_count'] == 1
    assert april['rows'][0]['name'] == first.name and may['rows'][0]['name'] == second.name
    assert april['totals']['assigned_usd'] is may['totals']['assigned_usd'] is None


def test_partial_shared_cash_stays_unattributed_but_same_month_aggregate_is_preserved():
    parent = dict(status='Importado', employer='A')
    first = application(historical_period='', historical_remitted_usd=0, collection_row_id='C', processing_route='Operativa')
    second = application(**dict(first, name='OTHER'))
    collections = {'C': dict(applied_usd=200, remittance_detail='[{"importe_usd":60}]')}
    progress = application_progress_index([first, second], {'IMPORT': parent}, collections)
    assert all(result['assigned_usd'] is None for result in progress.values())
    assert application_month_amounts([first, second], {'IMPORT': parent}, collections, progress)[('A', '2025-04')] == dict(total_usd=200, covered_usd=60)


def test_unreadable_distribution_is_flagged_without_relying_on_settled_label():
    parent = dict(status='Importado', employer='A')
    source = application(historical_detail='not json')
    result = application_progress(source, parent)
    assert result['state'] == 'Revisar distribución' and result['pending_usd'] is None
    assert 'inválido' in result['observations'][0]
    assert application_month_amounts([source], {'IMPORT': parent}, {})[('A', '2025-04')]['covered_usd'] is None
