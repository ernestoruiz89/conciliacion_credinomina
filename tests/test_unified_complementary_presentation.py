import io
import unittest
from datetime import datetime

from openpyxl import load_workbook

from credinomina_reconciliation.control_export import build_control_workbook
from credinomina_reconciliation.client_position import build_position
from credinomina_reconciliation.complementary_balances import financial_balance
from credinomina_reconciliation.conciliacion_credinomina.report.estado_de_cuenta_operativo.estado_de_cuenta_operativo import get_columns
from credinomina_reconciliation.conciliacion_credinomina.report.resumen_de_conciliacion.resumen_de_conciliacion import get_columns as summary_columns


class UnifiedComplementaryPresentationTests(unittest.TestCase):
    def workbook(self, delta):
        cash = max(delta, 0)
        movement = dict(name='AUTO-1', signed_amount_usd=delta, deposit_reference='REF', posting_date='2025-05-01')
        period = dict(name='P', employer='E', month='2025-04', complementary_usd=10,
            rounding_adjustment_usd=delta, rows=[], historical_rows=[], rounding_movements=[movement, movement])
        data = dict(year=2025, periods=[period], totals={}, cash_deposits=[dict(name='D', employer='E',
            date='2025-05-01', month='2025-05', total_usd=100 + cash, credits_usd=90, other_usd=10,
            adjustments_usd=cash, credit_balance_usd=0, unclassified_usd=0, review_usd=0,
            fx_rate=36.6243)])
        return load_workbook(io.BytesIO(build_control_workbook(data, exceptions=[], actions=[],
            employer_label='E', generated_at=datetime(2026, 10, 4))))

    def row(self, sheet, number=5):
        return {header.value: cell.value for header, cell in zip(sheet[4], sheet[number])}

    def test_removed_columns_and_cash_totals_for_both_signs(self):
        for delta in [.01, -.01]:
            with self.subTest(delta=delta):
                book = self.workbook(delta)
                for sheet in book:
                    for cell in sheet[4]:
                        self.assertNotIn(cell.value, {'Tolerancia US$', 'Ajuste de redondeo USD', 'Efectivo de ajustes USD', 'Ajuste USD'})
                period = self.row(book['Períodos'])
                self.assertAlmostEqual(period['Partidas complementarias USD'], 10 + delta)
                self.assertEqual(book['Períodos']['M5'].number_format, '#,##0.00;[Red](#,##0.00);"-"')
                deposit = self.row(book['Depósitos'])
                self.assertAlmostEqual(deposit['A partidas complementarias USD'], 10 + max(delta, 0))
                self.assertEqual(deposit['Depositado completo USD'], deposit['A créditos USD'] + deposit['A partidas complementarias USD'])
                self.assertEqual(book['Depósitos']['P5'].number_format, '0.00000000')
                entries = list(book['Partidas y excepciones'].iter_rows(min_row=5, values_only=True))
                automatic = [r for r in entries if r[0] == 'Partida complementaria']
                self.assertEqual(len(automatic), 1)
                self.assertEqual(automatic[0][6:10], ('AUTO-1', delta, 'Vigente', 'Diferencia por tolerancia'))
                self.assertEqual(automatic[0][1].date().isoformat(), '2025-05-01')

    def test_automatic_is_a_single_complementary_row_not_an_application_column(self):
        item = dict(name='AUTO', employer='E', category='Diferencia por tolerancia', docstatus=1,
            status='Vigente', amount_usd=-.01, accounting_status='No requiere registro')
        application = dict(source_import='I', applied_usd=46.53, paid_usd=46.52,
            adjustment_usd=-.01, amount_usd=0)
        rows = build_position([], [application], [item], {'AUTO': financial_balance(item)}, [], {})
        self.assertEqual(len(rows), 2)
        comp = next(row for row in rows if row['position_type'] == 'Partida complementaria')
        self.assertEqual(comp['complementary_usd'], -.01)
        self.assertEqual(comp['source_document'], 'AUTO')
        self.assertEqual(next(row for row in rows if row['position_type'] == 'Aplicación')['applied_pending_usd'], 0)
        for columns in [get_columns({}), get_columns({'position_type': 'Aplicación'}), summary_columns()]:
            self.assertNotIn('rounding_adjustment_usd', {col['fieldname'] for col in columns})
