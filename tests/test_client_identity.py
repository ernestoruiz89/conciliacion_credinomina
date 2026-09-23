import io
import unittest

from openpyxl import Workbook

from credinomina_reconciliation.client_identity import (
    choose_client, matching_name, name_key,
)
from credinomina_reconciliation.parsers import parse_collection_file
from credinomina_reconciliation.parsers import SourceFileError
from credinomina_reconciliation.reconciliation import (
    application_matches_collection, match_collection_record,
)
from credinomina_reconciliation.remittance_detail import suggest_detail_targets


class ClientIdentityTests(unittest.TestCase):
    def test_name_order_and_accents_are_normalized_without_fuzzy_spelling(self):
        self.assertEqual(name_key("Pérez Ana María"), name_key("ANA MARIA PEREZ"))
        self.assertFalse(matching_name("ANA M. PEREZ", {
            "client_name": "Ana María Pérez", "client_aliases": [],
        }))
        self.assertTrue(matching_name("ANA M. PEREZ", {
            "client_name": "Ana María Pérez", "client_aliases": ["ANA M. PEREZ"],
        }))

    def test_exact_id_creates_alias_but_conflicting_ids_never_merge(self):
        clients = [
            {"name": "A", "client_name": "ANA PEREZ", "client_number": "001", "national_id": "CED-A", "client_aliases": []},
            {"name": "B", "client_name": "LUIS RUIZ", "client_number": "002", "national_id": "CED-B", "client_aliases": []},
        ]
        found, reason = choose_client({"client_number": "1", "national_id": "CED-A"}, clients)
        self.assertEqual(found["name"], "A")
        found, reason = choose_client({"client_number": "001", "national_id": "CED-B"}, clients)
        self.assertIsNone(found)
        self.assertIn("Conflicto", reason)

    def test_name_and_employee_number_are_scoped_to_employer(self):
        clients = [
            {"name": "A", "employer": "Empresa A", "client_name": "ANA PEREZ",
             "client_number": "100", "employee_number": "7", "national_id": "CED-A"},
            {"name": "B", "employer": "Empresa B", "client_name": "PEREZ ANA",
             "client_number": "200", "employee_number": "7", "national_id": "CED-B"},
        ]
        found, _ = choose_client({"client_name": "ANA PEREZ"}, clients, "Empresa A")
        self.assertEqual(found["name"], "A")
        found, _ = choose_client({"employee_number": "007"}, clients, "Empresa B")
        self.assertEqual(found["name"], "B")
        found, reason = choose_client({"client_number": "100"}, clients, "Empresa B")
        self.assertIsNone(found)
        self.assertIn("otra empresa", reason)

    def test_name_only_response_matches_unique_alias(self):
        rows = [
            {"name": "ROW-1", "client_name": "ANA MARIA PEREZ", "client_aliases": ["ANA M. PEREZ"], "loan_number": "111"},
            {"name": "ROW-2", "client_name": "LUIS RUIZ", "loan_number": "222"},
        ]
        match, reason = match_collection_record({"client_name": "PEREZ ANA MARIA"}, rows)
        self.assertEqual(match["name"], "ROW-1")
        self.assertIn("nombre", reason)
        match, _reason = match_collection_record({"client_name": "ANA M. PEREZ"}, rows)
        self.assertEqual(match["name"], "ROW-1")

    def test_same_name_for_two_clients_is_ambiguous(self):
        rows = [
            {"name": "ROW-1", "client_name": "ANA PEREZ", "client_number": "1"},
            {"name": "ROW-2", "client_name": "PEREZ ANA", "client_number": "2"},
        ]
        match, reason = match_collection_record({"client_name": "ANA PEREZ"}, rows)
        self.assertIsNone(match)
        self.assertIn("ambigua", reason)

    def test_conflicting_identifiers_do_not_fall_back_to_loan_or_name(self):
        rows = [{
            "name": "ROW-1", "client_name": "ANA PEREZ",
            "client_number": "1", "national_id": "CED-A",
            "loan_number": "111",
        }]
        match, _reason = match_collection_record({
            "client_name": "ANA PEREZ", "client_number": "1",
            "national_id": "CED-OTRA", "loan_number": "111",
        }, rows)
        self.assertIsNone(match)

    def test_employee_number_matches_collection_and_rejects_conflicts(self):
        rows = [{
            "name": "ROW-1", "client_name": "ANA PEREZ", "employee_number": "007",
            "loan_number": "111",
        }]
        match, _ = match_collection_record({"employee_number": "7"}, rows)
        self.assertEqual(match["name"], "ROW-1")
        match, _ = match_collection_record({
            "employee_number": "008", "loan_number": "111",
        }, rows)
        self.assertIsNone(match)
        self.assertTrue(application_matches_collection(
            {"employee_number": "7"}, rows[0],
        ))
        self.assertFalse(application_matches_collection(
            {"employee_number": "8"}, rows[0],
        ))

    def test_row_key_does_not_override_conflicting_client_identity(self):
        rows = [{
            "name": "ROW-1", "row_key": "K-1", "client_name": "ANA PEREZ",
            "client_number": "1", "national_id": "CED-A", "loan_number": "111",
        }]
        match, reason = match_collection_record({
            "row_key": "K-1", "client_name": "LUIS RUIZ",
            "client_number": "2", "loan_number": "111",
        }, rows)
        self.assertIsNone(match)
        self.assertIn("contradice", reason)

    def test_name_only_deposit_detail_matches_unique_person(self):
        claims = [{
            "id": "H:1", "kind": "H", "amount_usd": 50,
            "group": "Empresa A", "period": "P1", "client_name": "ANA MARIA PEREZ",
            "client_names": ["ANA M. PEREZ"], "client_number": "1",
            "loan_number": "111", "references": [],
        }]
        matches, _reason = suggest_detail_targets(
            {"client_name": "ANA M. PEREZ"}, claims, 50, "Empresa A"
        )
        self.assertEqual(matches, [{"claim_id": "H:1", "amount_usd": 50}])
        matches, _reason = suggest_detail_targets(
            {"client_name": "ANA M. PEREZ", "client": "B"},
            [{**claims[0], "client": "A"}], 50, "Empresa A",
        )
        self.assertEqual(matches, [])

    def test_application_uses_name_only_without_any_identifier(self):
        collection = {
            "client_name": "ANA MARIA PEREZ", "client_number": "1",
            "national_id": "CED-A", "loan_number": "111",
        }
        self.assertTrue(application_matches_collection(
            {"client_name": "PEREZ ANA MARIA"}, collection,
        ))
        self.assertTrue(application_matches_collection(
            {"client_name": "ANA M. PEREZ"}, collection, ["ANA M. PEREZ"],
        ))
        self.assertFalse(application_matches_collection(
            {"client_name": "ANA M. PEREZ"}, collection,
        ))
        self.assertFalse(application_matches_collection(
            {"client_name": "PEREZ ANA MARIA", "client_number": "2"}, collection,
        ))

    def test_parser_accepts_name_only_company_detail(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Nombre y Apellidos del Cliente", "Deducido US$"])
        sheet.append(["ANA PEREZ", 50])
        buffer = io.BytesIO()
        book.save(buffer)
        rows = parse_collection_file(
            "empresa.xlsx", buffer.getvalue(),
            require_deduction=True, require_name=True,
        )
        self.assertEqual(rows[0]["client_name"], "ANA PEREZ")
        self.assertEqual(rows[0]["loan_number"], "")

    def test_parser_reads_employee_number_separately_from_client_number(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Nro. Cliente", "Nro. Empleado", "Nombre y Apellidos del Cliente", "Deducido US$"])
        sheet.append(["C-100", "E-7", "ANA PEREZ", 50])
        buffer = io.BytesIO()
        book.save(buffer)
        rows = parse_collection_file(
            "empresa.xlsx", buffer.getvalue(), require_deduction=True,
            require_name=True,
        )
        self.assertEqual(rows[0]["client_number"], "C-100")
        self.assertEqual(rows[0]["employee_number"], "E-7")

    def test_parser_rejects_amount_without_client_name(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Nombre y Apellidos del Cliente", "Deducido US$"])
        sheet.append(["", 50])
        buffer = io.BytesIO()
        book.save(buffer)
        with self.assertRaises(SourceFileError):
            parse_collection_file(
                "empresa.xlsx", buffer.getvalue(),
                require_deduction=True, require_name=True,
            )


if __name__ == "__main__":
    unittest.main()
